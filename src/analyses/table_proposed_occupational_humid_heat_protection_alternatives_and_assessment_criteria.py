#!/usr/bin/env python3
"""Generate proposed occupational humid-heat protection alternatives table.

Plan: Prespecify the six action alternatives, their district-selection rules,
capacity treatment, assessment criteria, distributional implications,
implementation assumptions, and evidence limits before scenario comparison.

Framework: AnaSOP Section 5 defines A0-A5 and the district-count capacity
contract. Section 6 defines the priority components, selection rules, coverage
criteria, and noncausal interpretation boundary. Section 7 Step 6 freezes these
alternatives before their consequences are calculated. This table is
constructed entirely from that approved research contract and does not read or
write the research database.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageChops

ROOT = Path(__file__).resolve().parents[2]
TITLE = "Proposed Occupational Humid-Heat Protection Alternatives and Assessment Criteria"
DEFAULT_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_proposed_occupational_humid_heat_protection_alternatives_and_assessment_criteria.xlsx"
)
DEFAULT_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_proposed_occupational_humid_heat_protection_alternatives_and_assessment_criteria.png"
)

FULL_HEADERS = [
    "Strategy",
    "Action Alternative",
    "District-Selection Rule",
    "Capacity Treatment",
    "Priority Components",
    "Primary Assessment Criteria",
    "Distributional Implication",
    "Implementation Assumption",
    "Evidence Limit",
]

COMPACT_HEADERS = [
    "Strategy",
    "Selection Basis",
    "Optimized Objective",
    "Capacity",
    "Evidence Boundary",
]

NAVY = "17324D"
PALE_TEAL = "DCEBEA"
PALE_BLUE = "EEF4F7"
PALE_GOLD = "FFF2CC"
WHITE = "FFFFFF"
TEXT = "23313D"
MUTED = "5C6B73"
LIGHT_BORDER = "B8C7CE"
MEDIUM_BORDER = "6F8791"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--review-output", type=Path, default=DEFAULT_REVIEW_OUTPUT)
    return parser.parse_args()


def resolve_path(path: Path) -> Path:
    return path if path.is_absolute() else ROOT / path


def set_cell_style(
    cell: object,
    *,
    fill: str | None = None,
    color: str = TEXT,
    bold: bool = False,
    size: int = 9,
    align: str = "center",
    wrap: bool = False,
) -> None:
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)


def render_review_png(workbook_path: Path, png_path: Path) -> None:
    soffice = Path("/opt/homebrew/bin/soffice")
    if not soffice.exists():
        located = shutil.which("soffice")
        if not located:
            raise RuntimeError("LibreOffice/soffice is required for workbook rendering")
        soffice = Path(located)
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("pdftoppm is required for PNG rendering")

    png_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix="mj02b-alternatives-table-review-", dir="/private/tmp"
    ) as temporary:
        temporary_path = Path(temporary)
        profile_uri = (temporary_path / "lo-profile").as_uri()
        conversion = subprocess.run(
            [
                str(soffice),
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temporary_path),
                str(workbook_path),
            ],
            check=False,
            capture_output=True,
            text=True,
        )
        if conversion.returncode != 0:
            detail = (conversion.stderr or conversion.stdout).strip()
            raise RuntimeError(
                f"LibreOffice PDF conversion failed with exit {conversion.returncode}: {detail}"
            )
        pdf_path = temporary_path / f"{workbook_path.stem}.pdf"
        if not pdf_path.exists():
            raise RuntimeError("LibreOffice did not create the expected PDF")
        prefix = png_path.with_suffix("")
        subprocess.run(
            [
                pdftoppm,
                "-png",
                "-r",
                "180",
                "-f",
                "1",
                "-singlefile",
                str(pdf_path),
                str(prefix),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
    if not png_path.exists() or png_path.stat().st_size == 0:
        raise RuntimeError("PNG review render is missing or empty")
    with Image.open(png_path) as image:
        rgb = image.convert("RGB")
        background = Image.new("RGB", rgb.size, "white")
        content_box = ImageChops.difference(rgb, background).getbbox()
        if content_box:
            margin = 32
            left = max(0, content_box[0] - margin)
            top = max(0, content_box[1] - margin)
            right = min(rgb.width, content_box[2] + margin)
            bottom = min(rgb.height, content_box[3] + margin)
            rgb.crop((left, top, right, bottom)).save(png_path)


def build_table() -> pd.DataFrame:
    rows = [
        {
            "Strategy": "A0",
            "Action Alternative": "No additional targeting",
            "District-Selection Rule": "Select no District; retain the current no-program counterfactual.",
            "Capacity Treatment": "No implementation capacity deployed; repeated as a fixed lower reference at each tier.",
            "Priority Components": "None",
            "Primary Assessment Criteria": "Zero Humid-Heat, Vulnerable-Worker, and Working-Age Coverage by construction.",
            "Distributional Implication": "No District or worker group receives additional targeted protection.",
            "Implementation Assumption": "Represents continuation without the proposed district-targeted package.",
            "Evidence Limit": "Reference condition only; it does not estimate current informal protection or realized harm.",
        },
        {
            "Strategy": "A1",
            "Action Alternative": "Uniform national provision",
            "District-Selection Rule": "Include all 197 current Districts without priority ranking.",
            "Capacity Treatment": "Outside the 20, 40, and 60 District constraints; repeated as a fixed upper reference.",
            "Priority Components": "Universal geographic inclusion",
            "Primary Assessment Criteria": "Full geographic and modeled coverage by construction.",
            "Distributional Implication": "Equal formal geographic eligibility, without extra intensity for higher-need Districts.",
            "Implementation Assumption": "Requires capacity to operate in every current District.",
            "Evidence Limit": "Upper benchmark, not a demonstrated feasible option under limited resources.",
        },
        {
            "Strategy": "A2",
            "Action Alternative": "Humid-heat-burden-first targeting",
            "District-Selection Rule": "Select the top-ranked Districts by working-age-weighted threshold humid-heat burden; replicate with working-age-weighted nonnegative wet-bulb excess above 26 C.",
            "Capacity Treatment": "Select 20, 40, or 60 Districts; repeat at 10%, 20%, and 30% worker-share targets.",
            "Priority Components": "District Threshold Humid-Heat Burden; 26 C continuous-excess burden replication",
            "Primary Assessment Criteria": "Humid-Heat Coverage, with Working-Age and Vulnerable-Worker Coverage reported jointly.",
            "Distributional Implication": "Maximizes observed working-age-weighted heat-burden coverage at fixed District capacity.",
            "Implementation Assumption": "Observed historical relative burden is relevant to near-term screening.",
            "Evidence Limit": "Burden is District-level and historical; it is not individual exposure or a future projection.",
        },
        {
            "Strategy": "A3",
            "Action Alternative": "Vulnerable-worker-first targeting",
            "District-Selection Rule": "Select the top-ranked Districts by District Vulnerable-Worker Scale.",
            "Capacity Treatment": "Select 20, 40, or 60 Districts; repeat at 10%, 20%, and 30% worker-share targets.",
            "Priority Components": "District Vulnerable-Worker Scale",
            "Primary Assessment Criteria": "Vulnerable-Worker Coverage, with Humid-Heat and Working-Age Coverage reported jointly.",
            "Distributional Implication": "Maximizes low-education working-age coverage at fixed District capacity.",
            "Implementation Assumption": "Low Education is a relevant communication and adjustment-capacity marker.",
            "Evidence Limit": "Education is not a complete vulnerability index and its role is not interpreted causally.",
        },
        {
            "Strategy": "A4",
            "Action Alternative": "Worker-scale-first targeting",
            "District-Selection Rule": "Select the top-ranked Districts by Weighted Working-Age Scale.",
            "Capacity Treatment": "Select 20, 40, or 60 Districts; repeat at 10%, 20%, and 30% worker-share targets.",
            "Priority Components": "Weighted Working-Age Scale",
            "Primary Assessment Criteria": "Working-Age Coverage, with Humid-Heat and Vulnerable-Worker Coverage reported jointly.",
            "Distributional Implication": "Maximizes potential population reach but may favor populous, lower-hazard Districts.",
            "Implementation Assumption": "Survey-weighted relative scale approximates the population potentially reached.",
            "Evidence Limit": "Relative survey-weighted scale is not an official current population projection.",
        },
        {
            "Strategy": "A5",
            "Action Alternative": "Combined targeting",
            "District-Selection Rule": "Select by the equal-weight percentile-rank score; test six pure or emphasis weight alternatives.",
            "Capacity Treatment": "Select the highest 20, 40, or 60 Districts; repeat under worker-share capacity.",
            "Priority Components": "District Humid-Heat Burden, Weighted Working-Age Scale, and District Vulnerable-Worker Scale",
            "Primary Assessment Criteria": "Joint Humid-Heat, Vulnerable-Worker, and Working-Age Coverage plus inclusion stability.",
            "Distributional Implication": "Seeks a transparent coverage-equity compromise rather than maximizing one objective.",
            "Implementation Assumption": "Equal weights are a transparent compromise convention; alternative values are disclosed as sensitivities.",
            "Evidence Limit": "Data Support Grade stays outside the score; ranking does not estimate causal program benefit.",
        },
    ]
    table = pd.DataFrame(rows, columns=FULL_HEADERS)
    validate_table(table)
    return table


def build_compact_table() -> pd.DataFrame:
    """Return the manuscript-facing, Stargazer-style policy summary."""
    compact = pd.DataFrame(
        [
            ["A0  No targeting", "Select no district", "Lower reference", "None", "Does not measure existing informal protection"],
            ["A1  Uniform national", "Include all 197 districts", "Upper reference", "National", "Outside the limited-capacity comparison"],
            ["A2  Heat burden", "Rank historical humid-heat burden", "Humid-heat coverage", "20 / 40 / 60 districts", "District-level historical burden; not individual exposure"],
            ["A3  Vulnerable workers", "Rank low-education worker scale", "Vulnerable-worker coverage", "20 / 40 / 60 districts", "Education is an incomplete, noncausal vulnerability marker"],
            ["A4  Worker scale", "Rank working-age scale", "Working-age coverage", "20 / 40 / 60 districts", "Survey-weighted scale is not an official population projection"],
            ["A5  Combined", "Equal-rank average of A2-A4 inputs", "Coverage compromise", "20 / 40 / 60 districts", "Weights are transparent conventions; no causal benefit estimate"],
        ],
        columns=COMPACT_HEADERS,
    )
    if compact.shape != (6, 5) or compact.isna().any().any():
        raise AssertionError("Compact alternatives display must be complete and 6 x 5")
    return compact


def validate_table(table: pd.DataFrame) -> None:
    if table.shape != (6, 9):
        raise AssertionError(f"Expected a 6 x 9 table, observed {table.shape}")
    if table["Strategy"].tolist() != ["A0", "A1", "A2", "A3", "A4", "A5"]:
        raise AssertionError("Strategy order must be A0-A5")
    combined = table.loc[table["Strategy"].eq("A5")].iloc[0]
    if "Data Support Grade stays outside the score" not in combined["Evidence Limit"]:
        raise AssertionError("The approved Data Support Grade boundary is missing")
    a2 = table.loc[table["Strategy"].eq("A2")].iloc[0]
    if "nonnegative wet-bulb excess above 26 C" not in a2["District-Selection Rule"]:
        raise AssertionError("The corrected continuous-burden replication is missing")
    if table.isna().any().any():
        raise AssertionError("The alternatives table must not contain missing cells")


def build_workbook(table: pd.DataFrame, compact: pd.DataFrame, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Alternatives"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COMPACT_HEADERS))
    title_cell = sheet.cell(1, 1, TITLE)
    set_cell_style(title_cell, color=NAVY, bold=True, size=14, align="left")
    sheet.row_dimensions[1].height = 31

    for column, header in enumerate(COMPACT_HEADERS, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, bold=True, size=9, wrap=True)
    sheet.row_dimensions[2].height = 30

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    for offset, row in enumerate(compact.itertuples(index=False, name=None)):
        row_number = 3 + offset
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                bold=column == 1,
                size=9,
                align="left",
                wrap=True,
            )
            cell.border = Border(top=group_top if offset in {0, 2, 5} else Side(), bottom=thin)
        sheet.row_dimensions[row_number].height = 42

    notes = [
        "Notes: A0 and A1 are fixed references. A2-A5 share the same 20, 40, and 60 district capacities; matched worker-share capacities are retained in the audit sheet.",
        "A2 uses threshold-day burden as primary and nonnegative wet-bulb excess above 26 C as a sensitivity. Data Support Grade remains outside every priority score.",
        "Coverage is descriptive allocation reach, not intervention effectiveness or avoided labor loss. Full assumptions and definitions are retained in Specification Audit.",
    ]
    notes_start = 10
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=len(COMPACT_HEADERS),
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 25

    widths = [24, 34, 28, 25, 48]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = "A2:E8"
    sheet.print_area = f"A1:E{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.15
    sheet.page_margins.right = 0.15
    sheet.page_margins.top = 0.20
    sheet.page_margins.bottom = 0.20

    audit = workbook.create_sheet("Specification Audit")
    audit.sheet_view.showGridLines = False
    for column, header in enumerate(FULL_HEADERS, start=1):
        cell = audit.cell(1, column, header)
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=8, wrap=True)
    for row_number, row in enumerate(table.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            cell = audit.cell(row_number, column, value)
            set_cell_style(cell, size=8, align="left", wrap=True)
            cell.border = Border(bottom=thin)
        audit.row_dimensions[row_number].height = 78
    for column, width in enumerate([10, 24, 34, 31, 32, 35, 35, 34, 34], start=1):
        audit.column_dimensions[get_column_letter(column)].width = width
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit.print_area = audit.dimensions
    audit.page_setup.orientation = "landscape"
    audit.page_setup.paperSize = audit.PAPERSIZE_A3
    audit.page_setup.fitToWidth = 1
    audit.page_setup.fitToHeight = 1
    audit.sheet_properties.pageSetUpPr.fitToPage = True

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    metadata_rows = [
        ("Field", "Value"),
        ("Table", TITLE),
        ("Research type", "Applied ex ante spatial social impact assessment"),
        ("Decision unit", "Current District or Municipality"),
        ("Capacity definitions", "20, 40, and 60 Districts; 10%, 20%, and 30% worker-share targets"),
        ("Priority inputs", "Humid-heat burden, working-age scale, and vulnerable-worker scale"),
        ("Uncertainty input", "Data Support Grade, kept outside the score"),
        ("Database access", "None required"),
        ("Database writes", "None"),
    ]
    for row_number, values in enumerate(metadata_rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = metadata.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=NAVY if row_number == 1 else None,
                color=WHITE if row_number == 1 else TEXT,
                bold=row_number == 1 or column == 1,
                align="left",
                wrap=True,
            )
    metadata.column_dimensions["A"].width = 26
    metadata.column_dimensions["B"].width = 78
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = metadata.dimensions
    metadata.print_area = metadata.dimensions
    metadata.page_setup.fitToWidth = 1
    metadata.page_setup.fitToHeight = 1
    metadata.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.save(output)


def validate_outputs(
    table: pd.DataFrame, compact: pd.DataFrame, output: Path, review: Path
) -> None:
    validate_table(table)
    if compact.shape != (6, 5):
        raise AssertionError("Unexpected compact alternatives dimensions")
    workbook = load_workbook(output, data_only=False, read_only=False)
    if workbook.sheetnames != ["Alternatives", "Specification Audit", "Metadata"]:
        raise AssertionError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Alternatives"]
    if sheet["A1"].value != TITLE:
        raise AssertionError("Workbook title is missing")
    if sheet["A2"].value != "Strategy" or sheet.freeze_panes != "A3":
        raise AssertionError("The title must be followed immediately by column headers")
    if sheet.max_row != 12 or sheet.max_column != 5:
        raise AssertionError("Unexpected worksheet dimensions")
    if workbook["Specification Audit"].max_row != 7:
        raise AssertionError("The complete six-strategy specification audit is missing")
    if any(item.print_area is None for item in workbook.worksheets):
        raise AssertionError("Every worksheet must have a print area")
    workbook.close()
    if not review.exists() or review.stat().st_size == 0:
        raise AssertionError("PNG review copy is missing")


def main() -> None:
    args = parse_args()
    output = resolve_path(args.output)
    review = resolve_path(args.review_output)
    table = build_table()
    compact = build_compact_table()
    build_workbook(table, compact, output)
    render_review_png(output, review)
    validate_outputs(table, compact, output, review)
    print(f"table={output.relative_to(ROOT)}")
    print(f"review={review.relative_to(ROOT)}")
    print(f"display_rows={len(compact)} display_columns={len(compact.columns)}")
    print(f"audit_rows={len(table)} audit_columns={len(table.columns)}")
    print("database_access=none database_writes=none")


if __name__ == "__main__":
    main()
