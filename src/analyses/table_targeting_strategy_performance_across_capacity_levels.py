#!/usr/bin/env python3
"""Generate targeting-strategy performance across capacity levels.

Plan: Compare the district reach, humid-heat burden coverage, vulnerable-worker
coverage, working-age coverage, relative capture, and data-support profile of
six prespecified alternatives at three capacity levels.

Framework: AnaSOP Section 5 defines A0-A5 and treats A0/A1 as fixed lower and
upper references. Section 6 defines district selection, the 20/40/60 District
capacity levels, and the three coverage formulas. Section 7 Step 7 requires all
targeted strategies to be compared at identical capacity. Authoritative inputs
are rebuilt from PostgreSQL under default_transaction_read_only=on; this script
performs no database writes.
"""

from __future__ import annotations

import argparse
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, Side
from openpyxl.utils import get_column_letter

from build_impact_assessment_indicators import (
    build_district_indicators,
    connection_args,
    load_inputs,
    verify_read_only,
)
from table_proposed_occupational_humid_heat_protection_alternatives_and_assessment_criteria import (
    LIGHT_BORDER,
    MEDIUM_BORDER,
    MUTED,
    NAVY,
    PALE_BLUE,
    PALE_GOLD,
    PALE_TEAL,
    TEXT,
    WHITE,
    render_review_png,
    resolve_path,
    set_cell_style,
)


ROOT = Path(__file__).resolve().parents[2]
TITLE = "Targeting Strategy Performance across Capacity Levels"
DEFAULT_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_targeting_strategy_performance_across_capacity_levels.xlsx"
)
DEFAULT_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_targeting_strategy_performance_across_capacity_levels.png"
)
DEFAULT_WORKER_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_targeting_strategy_performance_across_capacity_levels_worker_share.png"
)

CAPACITIES = [20, 40, 60]
WORKER_SHARE_TARGETS = [0.10, 0.20, 0.30]
STRATEGIES = [
    ("A0", "No additional targeting"),
    ("A1", "Uniform national provision"),
    ("A2", "Humid-heat-burden-first targeting"),
    ("A3", "Vulnerable-worker-first targeting"),
    ("A4", "Worker-scale-first targeting"),
    ("A5", "Combined targeting"),
]
HEADERS = [
    "Strategy",
    "Capacity",
    "Selected Districts",
    "Humid-Heat Coverage",
    "Vulnerable-Worker Coverage",
    "Working-Age Coverage",
    "Heat Capture Ratio",
    "Vulnerable Capture Ratio",
    "Data Support (H/M/L)",
    "Decision Interpretation",
]

CENTRAL_CAPACITY = 40
CENTRAL_HEADERS = [
    "Metric",
    "A0\nNo targeting",
    "A1\nNational",
    "A2\nHeat burden",
    "A3\nVulnerable workers",
    "A4\nWorker scale",
    "A5\nCombined",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--review-output", type=Path, default=DEFAULT_REVIEW_OUTPUT)
    parser.add_argument(
        "--worker-review-output", type=Path, default=DEFAULT_WORKER_REVIEW_OUTPUT
    )
    return parser.parse_args()


def strategy_score(district: pd.DataFrame, strategy: str) -> pd.Series:
    if strategy == "A2":
        return district["threshold_burden_percentile"]
    if strategy == "A3":
        return district["vulnerable_worker_percentile"]
    if strategy == "A4":
        return district["working_age_percentile"]
    if strategy == "A5":
        return (
            district["threshold_burden_percentile"]
            + district["working_age_percentile"]
            + district["vulnerable_worker_percentile"]
        ) / 3.0
    raise ValueError(f"No targeting score is defined for {strategy}")


def selected_districts(
    district: pd.DataFrame, strategy: str, capacity: int
) -> pd.DataFrame:
    if strategy == "A0":
        return district.iloc[0:0].copy()
    if strategy == "A1":
        return district.copy()
    ranked = district.assign(_score=strategy_score(district, strategy)).sort_values(
        ["_score", "admin2_code"], ascending=[False, True], kind="mergesort"
    )
    return ranked.head(capacity).drop(columns="_score")


def selected_by_worker_share(
    district: pd.DataFrame, strategy: str, target: float
) -> pd.DataFrame:
    if strategy == "A0":
        return district.iloc[0:0].copy()
    if strategy == "A1":
        return district.copy()
    ranked = district.assign(_score=strategy_score(district, strategy)).sort_values(
        ["_score", "admin2_code"], ascending=[False, True], kind="mergesort"
    )
    cumulative = ranked["weighted_working_age_scale"].cumsum()
    reached = np.flatnonzero(cumulative.to_numpy() >= target)
    if len(reached) == 0:
        return ranked.drop(columns="_score")
    return ranked.iloc[: int(reached[0]) + 1].drop(columns="_score")


def coverage_metrics(
    selected: pd.DataFrame,
    heat_denominator: float,
    vulnerable_denominator: float,
) -> tuple[float, float, float]:
    heat_numerator = (
        selected["weighted_working_age_scale"]
        * selected["mean_cumulative_wb26_days"]
    ).sum()
    vulnerable_numerator = (
        selected["weighted_working_age_scale"]
        * selected["district_low_education_share"]
    ).sum()
    working_age = selected["weighted_working_age_scale"].sum()
    return (
        float(heat_numerator / heat_denominator),
        float(vulnerable_numerator / vulnerable_denominator),
        float(working_age),
    )


def support_profile(selected: pd.DataFrame) -> str:
    counts = selected["data_support_grade"].value_counts()
    return f"{int(counts.get('High', 0))}/{int(counts.get('Medium', 0))}/{int(counts.get('Limited', 0))}"


def decision_interpretation(
    strategy: str,
    heat_is_highest: bool,
    vulnerable_is_highest: bool,
    working_age_is_highest: bool,
) -> str:
    if strategy == "A0":
        return "Lower reference: no additional targeted reach."
    if strategy == "A1":
        return "Upper reference: full reach, outside limited capacity."
    achievements = []
    if heat_is_highest:
        achievements.append("highest heat coverage")
    if vulnerable_is_highest:
        achievements.append("highest vulnerable-worker coverage")
    if working_age_is_highest:
        achievements.append("highest working-age coverage")
    if achievements:
        return "Targeted-strategy leader: " + "; ".join(achievements) + "."
    labels = {
        "A2": "Humid-heat-burden benchmark; compare population and equity trade-offs.",
        "A3": "Vulnerable-worker benchmark; compare heat and population trade-offs.",
        "A4": "Working-age-reach benchmark; compare heat and equity trade-offs.",
        "A5": "Equal-weight compromise; retain only if its observed trade-offs are defensible.",
    }
    return labels[strategy]


def build_performance_table(district: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    heat_denominator = float(
        (
            district["weighted_working_age_scale"]
            * district["mean_cumulative_wb26_days"]
        ).sum()
    )
    vulnerable_denominator = float(
        (
            district["weighted_working_age_scale"]
            * district["district_low_education_share"]
        ).sum()
    )
    rows: list[dict[str, object]] = []

    def append_row(
        strategy: str,
        strategy_label: str,
        selected: pd.DataFrame,
        definition: str,
        target: float,
        label: str,
    ) -> None:
        heat, vulnerable, working_age = coverage_metrics(
            selected, heat_denominator, vulnerable_denominator
        )
        if definition == "district_count":
            reference_share = len(selected) / len(district)
        else:
            reference_share = working_age
        rows.append(
            {
                "strategy": strategy,
                "strategy_label": strategy_label,
                "capacity_definition": definition,
                "capacity_target": target,
                "capacity_label": label,
                "selected_districts": len(selected),
                "humid_heat_coverage": heat,
                "vulnerable_worker_coverage": vulnerable,
                "working_age_coverage": working_age,
                "heat_capture_ratio": heat / reference_share if reference_share else np.nan,
                "vulnerable_capture_ratio": vulnerable / reference_share if reference_share else np.nan,
                "data_support_profile": support_profile(selected),
            }
        )

    for capacity in CAPACITIES:
        for strategy, strategy_label in STRATEGIES:
            selected = selected_districts(district, strategy, capacity)
            append_row(
                strategy,
                strategy_label,
                selected,
                "district_count",
                float(capacity),
                f"{capacity} Districts ({capacity / len(district):.0%})",
            )
    for target in WORKER_SHARE_TARGETS:
        for strategy, strategy_label in STRATEGIES:
            selected = selected_by_worker_share(district, strategy, target)
            append_row(
                strategy,
                strategy_label,
                selected,
                "worker_share",
                target,
                f"{target:.0%} worker share",
            )
    audit = pd.DataFrame(rows)

    targeted = audit["strategy"].isin(["A2", "A3", "A4", "A5"])
    groups = audit[["capacity_definition", "capacity_target"]].drop_duplicates()
    for definition, target in groups.itertuples(index=False, name=None):
        mask = (
            targeted
            & audit["capacity_definition"].eq(definition)
            & audit["capacity_target"].eq(target)
        )
        subset = audit.loc[mask]
        maxima = {
            "heat": subset["humid_heat_coverage"].max(),
            "vulnerable": subset["vulnerable_worker_coverage"].max(),
            "working_age": subset["working_age_coverage"].max(),
        }
        group_mask = (
            audit["capacity_definition"].eq(definition)
            & audit["capacity_target"].eq(target)
        )
        for index in audit.index[group_mask]:
            row = audit.loc[index]
            audit.loc[index, "decision_interpretation"] = decision_interpretation(
                str(row["strategy"]),
                bool(
                    row["strategy"] in {"A2", "A3", "A4", "A5"}
                    and np.isclose(row["humid_heat_coverage"], maxima["heat"])
                ),
                bool(
                    row["strategy"] in {"A2", "A3", "A4", "A5"}
                    and np.isclose(
                        row["vulnerable_worker_coverage"], maxima["vulnerable"]
                    )
                ),
                bool(
                    row["strategy"] in {"A2", "A3", "A4", "A5"}
                    and np.isclose(row["working_age_coverage"], maxima["working_age"])
                ),
            )

    display = pd.DataFrame(
        {
            "Strategy": audit["strategy"] + " " + audit["strategy_label"],
            "Capacity": audit["capacity_label"],
            "Selected Districts": audit["selected_districts"],
            "Humid-Heat Coverage": audit["humid_heat_coverage"],
            "Vulnerable-Worker Coverage": audit["vulnerable_worker_coverage"],
            "Working-Age Coverage": audit["working_age_coverage"],
            "Heat Capture Ratio": audit["heat_capture_ratio"],
            "Vulnerable Capture Ratio": audit["vulnerable_capture_ratio"],
            "Data Support (H/M/L)": audit["data_support_profile"],
            "Decision Interpretation": audit["decision_interpretation"],
        },
        columns=HEADERS,
    )
    validate_performance(display, audit, len(district))
    return display, audit


def validate_performance(
    display: pd.DataFrame, audit: pd.DataFrame, district_count: int
) -> None:
    if district_count != 197:
        raise AssertionError(f"Expected 197 Districts, observed {district_count}")
    if display.shape != (36, 10) or audit.shape[0] != 36:
        raise AssertionError("Expected 36 strategy-capacity rows and 10 display columns")
    if audit.groupby("strategy").size().ne(6).any():
        raise AssertionError("Every strategy must appear in all six capacity scenarios")
    district_audit = audit[audit["capacity_definition"].eq("district_count")]
    for strategy in ["A2", "A3", "A4", "A5"]:
        selected = district_audit.loc[
            district_audit["strategy"].eq(strategy), "selected_districts"
        ].tolist()
        if selected != CAPACITIES:
            raise AssertionError(f"{strategy} does not select the planned District counts")
    reference_zero = audit["strategy"].eq("A0")
    if audit.loc[reference_zero, ["humid_heat_coverage", "vulnerable_worker_coverage", "working_age_coverage"]].to_numpy().any():
        raise AssertionError("A0 must have zero coverage")
    reference_full = audit["strategy"].eq("A1")
    if not np.allclose(
        audit.loc[reference_full, ["humid_heat_coverage", "vulnerable_worker_coverage", "working_age_coverage"]],
        1.0,
    ):
        raise AssertionError("A1 must have full coverage")
    coverage_columns = [
        "humid_heat_coverage",
        "vulnerable_worker_coverage",
        "working_age_coverage",
    ]
    if not audit[coverage_columns].apply(lambda column: column.between(0, 1).all()).all():
        raise AssertionError("Coverage metrics must remain within [0, 1]")
    targeted = audit[audit["strategy"].isin(["A2", "A3", "A4", "A5"])]
    for _, group in targeted.groupby(["strategy", "capacity_definition"]):
        if (group.sort_values("capacity_target")[coverage_columns].diff().dropna() < -1e-12).any().any():
            raise AssertionError("Coverage must not fall as capacity rises")
    worker = targeted[targeted["capacity_definition"].eq("worker_share")]
    if (worker["working_age_coverage"] + 1e-12 < worker["capacity_target"]).any():
        raise AssertionError("Worker-share selections must reach their stated target")
    for capacity in CAPACITIES:
        group = district_audit[
            district_audit["capacity_target"].eq(float(capacity))
            & district_audit["strategy"].isin(["A2", "A3", "A4", "A5"])
        ].set_index("strategy")
        if not np.isclose(group.loc["A2", "humid_heat_coverage"], group["humid_heat_coverage"].max()):
            raise AssertionError("A2 must maximize Humid-Heat Coverage at fixed District count")
        if not np.isclose(group.loc["A3", "vulnerable_worker_coverage"], group["vulnerable_worker_coverage"].max()):
            raise AssertionError("A3 must maximize Vulnerable-Worker Coverage at fixed District count")
        if not np.isclose(group.loc["A4", "working_age_coverage"], group["working_age_coverage"].max()):
            raise AssertionError("A4 must maximize Working-Age Coverage at fixed District count")


def build_central_matrix(audit: pd.DataFrame) -> pd.DataFrame:
    """Build the manuscript-facing Stargazer-style 40-district comparison."""
    central = audit.loc[
        audit["capacity_definition"].eq("district_count")
        & audit["capacity_target"].eq(float(CENTRAL_CAPACITY))
    ].set_index("strategy")
    if central.index.tolist() != [item[0] for item in STRATEGIES]:
        raise AssertionError("Central-capacity strategy order must be A0-A5")
    matrix = pd.DataFrame(
        {
            "Metric": [
                "Selected districts",
                "Humid-heat coverage",
                "Vulnerable-worker coverage",
                "Working-age coverage",
                "Heat capture ratio",
                "Vulnerable capture ratio",
                "Data support (H/M/L)",
            ],
            **{
                strategy: [
                    central.loc[strategy, "selected_districts"],
                    central.loc[strategy, "humid_heat_coverage"],
                    central.loc[strategy, "vulnerable_worker_coverage"],
                    central.loc[strategy, "working_age_coverage"],
                    central.loc[strategy, "heat_capture_ratio"],
                    central.loc[strategy, "vulnerable_capture_ratio"],
                    central.loc[strategy, "data_support_profile"],
                ]
                for strategy, _ in STRATEGIES
            },
        },
        columns=["Metric"] + [item[0] for item in STRATEGIES],
    )
    if matrix.shape != (7, 7):
        raise AssertionError("Central-capacity display must be 7 x 7")
    return matrix


def write_central_capacity_sheet(sheet: object, matrix: pd.DataFrame) -> None:
    sheet.title = "Central Capacity"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    title_cell = sheet.cell(1, 1, TITLE)
    set_cell_style(title_cell, color=NAVY, bold=True, size=14, align="left")
    sheet.row_dimensions[1].height = 31

    medium = Side(style="medium", color=MEDIUM_BORDER)
    thin = Side(style="thin", color=LIGHT_BORDER)
    for column, header in enumerate(CENTRAL_HEADERS, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, bold=True, size=9, wrap=True)
        cell.border = Border(top=medium, bottom=medium)
    sheet.row_dimensions[2].height = 35

    for offset, row in enumerate(matrix.itertuples(index=False, name=None)):
        row_number = 3 + offset
        metric = str(row[0])
        for column, value in enumerate(row, start=1):
            if pd.isna(value):
                value = "—"
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                color=MUTED if column in {2, 3} else TEXT,
                bold=column == 1,
                size=9,
                align="left" if column == 1 else "center",
                wrap=True,
            )
            cell.border = Border(bottom=medium if offset == len(matrix) - 1 else thin)
        if metric.endswith("coverage"):
            for column in range(2, 8):
                sheet.cell(row_number, column).number_format = "0.0%"
        elif metric.endswith("capture ratio"):
            for column in range(2, 8):
                if isinstance(sheet.cell(row_number, column).value, (int, float)):
                    sheet.cell(row_number, column).number_format = "0.00x"
        sheet.row_dimensions[row_number].height = 28

    # Bold marks the targeted strategy that maximizes its own criterion.
    sheet["D4"].font = Font(name="Aptos", size=9, bold=True, color=NAVY)
    sheet["E5"].font = Font(name="Aptos", size=9, bold=True, color=NAVY)
    sheet["F6"].font = Font(name="Aptos", size=9, bold=True, color=NAVY)

    notes = [
        "Notes: Central comparison at 40 districts. A0 and A1 are fixed lower and upper references; A2-A5 obey the 40-district constraint.",
        "Bold values identify the targeted strategy leading its own objective. H/M/L reports High, Medium, and Limited data-support counts.",
        "Coverage is descriptive allocation reach, not causal program benefit. Complete 20/40/60, worker-share, and scenario results remain in the audit sheets.",
    ]
    notes_start = 11
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=7)
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 24

    for column, width in enumerate([31, 18, 18, 18, 20, 18, 18], start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.print_area = "A1:G13"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25


def write_performance_sheet(
    sheet: object,
    display: pd.DataFrame,
    *,
    sheet_title: str,
    title: str,
    capacity_note: str,
    capture_note: str,
) -> None:
    sheet.title = sheet_title
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    title_cell = sheet.cell(1, 1, title)
    set_cell_style(title_cell, fill=NAVY, color=WHITE, bold=True, size=15, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=8, wrap=True)
    sheet.row_dimensions[2].height = 44

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    for offset, row in enumerate(display.itertuples(index=False, name=None)):
        row_number = 3 + offset
        fill = PALE_BLUE if offset % 2 else WHITE
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=fill,
                bold=column == 1,
                size=8,
                align="left" if column in {1, 10} else ("center" if column == 9 else "right"),
                wrap=True,
            )
            cell.border = Border(
                top=group_top if offset in {0, 6, 12} else Side(),
                bottom=thin,
                left=thin if column in {9, 10} else Side(),
            )
        for column in [4, 5, 6]:
            sheet.cell(row_number, column).number_format = "0.0%"
        for column in [7, 8]:
            sheet.cell(row_number, column).number_format = "0.00x"
        sheet.row_dimensions[row_number].height = 34

    last_data_row = 2 + len(display)
    notes = [
        capacity_note,
        "Humid-Heat Coverage is the selected share of national working-age-weighted cumulative humid-heat burden, constructed by averaging months within District-wave and then observed Survey Waves equally. Vulnerable-Worker Coverage is the selected share of low-education working-age scale. Working-Age Coverage sums relative working-age scale.",
        capture_note,
        "A2, A3, and A4 directly optimize humid-heat burden, vulnerable-worker, and working-age coverage at fixed District count. A5 is an equal-weight compromise and is not presumed preferred.",
        "Data Support reports selected District counts as High/Medium/Limited. Data Support Grade does not enter any priority score. Results use threshold humid-heat burden for the primary exposure definition.",
        "All analytical inputs were reconstructed through read-only PostgreSQL SELECT statements; database writes: none. Coverage is descriptive and is not an estimate of causal program effectiveness or avoided labor loss.",
    ]
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=len(HEADERS),
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 25

    widths = [31, 14, 14, 17, 19, 18, 16, 20, 21, 43]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:J{last_data_row}"
    sheet.print_area = f"A1:J{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.15
    sheet.page_margins.right = 0.15
    sheet.page_margins.top = 0.20
    sheet.page_margins.bottom = 0.20


def build_workbook(display: pd.DataFrame, audit: pd.DataFrame, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    district_mask = audit["capacity_definition"].eq("district_count")
    worker_mask = audit["capacity_definition"].eq("worker_share")
    central_matrix = build_central_matrix(audit)
    write_central_capacity_sheet(workbook.active, central_matrix)
    district_sheet = workbook.create_sheet("District Capacity Audit")
    write_performance_sheet(
        district_sheet,
        display.loc[district_mask].reset_index(drop=True),
        sheet_title="District Capacity Audit",
        title=f"{TITLE} — District-Count Audit",
        capacity_note="Notes: District-count capacity selects 20, 40, or 60 Districts. A0 and A1 are repeated fixed lower and upper references; A2-A5 obey the stated limited-capacity count.",
        capture_note="Capture ratios divide coverage by the selected share of 197 Districts; values above 1 indicate concentration above equal-District allocation. Ratios are undefined for A0.",
    )
    worker_sheet = workbook.create_sheet("Worker-Share Capacity")
    write_performance_sheet(
        worker_sheet,
        display.loc[worker_mask].reset_index(drop=True),
        sheet_title="Worker-Share Capacity",
        title=f"{TITLE} — Worker-Share Sensitivity",
        capacity_note="Notes: Worker-share capacity selects the shortest ranked District prefix reaching 10%, 20%, or 30% of Weighted Working-Age Scale. A0 and A1 remain fixed references.",
        capture_note="Capture ratios divide heat or vulnerable-worker coverage by actual Working-Age Coverage; values above 1 indicate concentration relative to the workforce share reached. Ratios are undefined for A0.",
    )

    audit_sheet = workbook.create_sheet("Scenario Audit")
    audit_sheet.sheet_view.showGridLines = False
    for column, header in enumerate(audit.columns, start=1):
        cell = audit_sheet.cell(1, column, header)
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=8, wrap=True)
    for row_number, row in enumerate(audit.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            cell = audit_sheet.cell(row_number, column, value)
            set_cell_style(cell, size=8, align="left" if column <= 4 else "right", wrap=True)
    audit_sheet.freeze_panes = "A2"
    audit_sheet.auto_filter.ref = audit_sheet.dimensions
    audit_sheet.print_area = audit_sheet.dimensions
    audit_sheet.page_setup.orientation = "landscape"
    audit_sheet.page_setup.paperSize = audit_sheet.PAPERSIZE_A3
    audit_sheet.page_setup.fitToWidth = 1
    audit_sheet.page_setup.fitToHeight = 1
    audit_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    for column in range(1, audit_sheet.max_column + 1):
        audit_sheet.column_dimensions[get_column_letter(column)].width = 20

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    metadata_rows = [
        ("Field", "Value"),
        ("Table", TITLE),
        ("Districts", "197 current Districts or Municipalities"),
        ("Primary objective", "District Threshold Humid-Heat Burden"),
        ("Combined rule", "Equal-weight exact percentile ranks of humid-heat burden, working-age scale, and vulnerable-worker scale"),
        ("Capacities", "20, 40, and 60 Districts; 10%, 20%, and 30% worker-share targets"),
        ("Database session", "PostgreSQL default_transaction_read_only=on"),
        ("Database writes", "None"),
    ]
    for row_number, values in enumerate(metadata_rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = metadata.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=NAVY if row_number == 1 else None,
                color=WHITE if row_number == 1 else TEXT,
                bold=row_number == 1 or column == 1,
                align="left",
                wrap=True,
            )
    metadata.column_dimensions["A"].width = 26
    metadata.column_dimensions["B"].width = 90
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = metadata.dimensions
    metadata.print_area = metadata.dimensions
    metadata.page_setup.fitToWidth = 1
    metadata.page_setup.fitToHeight = 1
    metadata.sheet_properties.pageSetUpPr.fitToPage = True

    workbook.save(output)


def validate_outputs(
    display: pd.DataFrame,
    audit: pd.DataFrame,
    output: Path,
    review: Path,
    worker_review: Path,
) -> None:
    validate_performance(display, audit, 197)
    workbook = load_workbook(output, data_only=False, read_only=False)
    expected = [
        "Central Capacity",
        "District Capacity Audit",
        "Worker-Share Capacity",
        "Scenario Audit",
        "Metadata",
    ]
    if workbook.sheetnames != expected:
        raise AssertionError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Central Capacity"]
    if sheet["A1"].value != TITLE:
        raise AssertionError("Workbook title is missing")
    if sheet["A2"].value != "Metric" or sheet.freeze_panes != "A3":
        raise AssertionError("The title must be followed immediately by column headers")
    if sheet.max_row != 13 or sheet.max_column != 7:
        raise AssertionError("Unexpected central-capacity display dimensions")
    if any(item.print_area is None for item in workbook.worksheets):
        raise AssertionError("Every worksheet must have a print area")
    worker_sheet = workbook["Worker-Share Capacity"]
    if worker_sheet["A2"].value != "Strategy" or worker_sheet.freeze_panes != "A3":
        raise AssertionError("Worker-share title must be followed immediately by headers")
    workbook.close()
    for path in [review, worker_review]:
        if not path.exists() or path.stat().st_size == 0:
            raise AssertionError(f"PNG review copy is missing: {path}")


def render_sheet_review(output: Path, sheet_name: str, review: Path) -> None:
    with tempfile.TemporaryDirectory(
        prefix="mj02b-strategy-performance-review-", dir="/private/tmp"
    ) as temporary:
        temporary_path = Path(temporary) / f"{sheet_name.replace(' ', '_')}.xlsx"
        workbook = load_workbook(output, data_only=False, read_only=False)
        for sheet in list(workbook.worksheets):
            if sheet.title != sheet_name:
                workbook.remove(sheet)
        workbook.active = 0
        workbook.save(temporary_path)
        workbook.close()
        render_review_png(temporary_path, review)


def main() -> None:
    args = parse_args()
    output = resolve_path(args.output)
    review = resolve_path(args.review_output)
    worker_review = resolve_path(args.worker_review_output)
    with psycopg.connect(**connection_args(args)) as connection:
        verify_read_only(connection)
        geography, climate, people = load_inputs(connection, args.schema)
    district, _ = build_district_indicators(geography, climate, people)
    display, audit = build_performance_table(district)
    build_workbook(display, audit, output)
    render_sheet_review(output, "Central Capacity", review)
    render_sheet_review(output, "Worker-Share Capacity", worker_review)
    validate_outputs(display, audit, output, review, worker_review)
    print(f"table={output.relative_to(ROOT)}")
    print(f"review={review.relative_to(ROOT)}")
    print(f"worker_review={worker_review.relative_to(ROOT)}")
    print(f"rows={len(display)} columns={len(display.columns)} districts={len(district)}")
    print("database_read_only=true database_writes=none")
    for capacity in CAPACITIES:
        combined = audit.loc[
            audit["strategy"].eq("A5")
            & audit["capacity_definition"].eq("district_count")
            & audit["capacity_target"].eq(float(capacity))
        ].iloc[0]
        print(
            f"combined_k={capacity} heat={combined['humid_heat_coverage']:.4f} "
            f"vulnerable={combined['vulnerable_worker_coverage']:.4f} "
            f"working_age={combined['working_age_coverage']:.4f}"
        )
    for target in WORKER_SHARE_TARGETS:
        part = audit.loc[
            audit["strategy"].isin(["A2", "A3", "A4", "A5"])
            & audit["capacity_definition"].eq("worker_share")
            & audit["capacity_target"].eq(target)
        ]
        counts = ",".join(
            f"{row.strategy}:{int(row.selected_districts)}"
            for row in part.itertuples(index=False)
        )
        print(f"worker_share_target={target:.0%} selected_districts={counts}")


if __name__ == "__main__":
    main()
