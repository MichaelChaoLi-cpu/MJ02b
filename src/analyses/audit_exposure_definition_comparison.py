#!/usr/bin/env python3
"""Compare dry-heat and humid-heat exposure definitions with read-only data.

This is a temporary diagnostic, not a formal AnaSOP output. Every analytical
input is obtained through SELECT from mda.public and the PostgreSQL session is
forced into read-only mode. The workbook, review image, scientific comparison
figure, CSV extracts, and README are written below data/exp only.
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import psycopg
import seaborn as sns
from linearmodels.iv import AbsorbingLS
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from run_heat_labor_credibility_audit import (
    DEMOGRAPHIC_CONTROLS,
    ROOT,
    connection_args,
    fit_hdfe,
    load_analysis_data,
)


DEFAULT_OUTPUT = Path("data/exp/exposure-definition-audit")
NAVY = "17324D"
PALE_TEAL = "DCEBEA"
PALE_BLUE = "EEF4F7"
WHITE = "FFFFFF"
TEXT = "243746"
MUTED = "5F6B73"
LIGHT_BORDER = "C9D4DA"

EXPOSURES = {
    "wb26_5days": {
        "column": "days_wbmax_ge_26c",
        "label": "WBmax >= 26 C days",
        "short_label": "Humid heat: WBmax >= 26 C",
        "literature_role": "Project-defined humid-heat burden; not a WBGT safety threshold",
    },
    "tmax35_5days": {
        "column": "days_tmax_ge_35c",
        "label": "Tmax >= 35 C days",
        "short_label": "Dry heat: Tmax >= 35 C",
        "literature_role": "Direct dry-heat threshold used in prior labor-productivity research",
    },
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def make_cells(frame: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "survey_wave",
        "survey_year",
        "survey_month",
        "admin2_code",
        "admin1_code",
        "days_wbmax_ge_26c",
        "days_tmax_ge_35c",
    ]
    cells = frame[columns].drop_duplicates(
        ["survey_year", "survey_month", "admin2_code"]
    ).reset_index(drop=True)
    cells["admin2_calendar_month"] = (
        cells["admin2_code"].astype("string")
        + "-m"
        + cells["survey_month"].astype("Int64").astype("string").str.zfill(2)
    )
    return cells


def residualize_exposure(cells: pd.DataFrame, column: str) -> pd.Series:
    working = cells.dropna(
        subset=[column, "admin2_calendar_month", "survey_year"]
    ).copy()
    absorb = pd.DataFrame(
        {
            "admin2_calendar_month": working[
                "admin2_calendar_month"
            ].astype("category"),
            "survey_year": working["survey_year"].astype("Int64").astype("category"),
        },
        index=working.index,
    )
    exog = pd.DataFrame({"constant": 1.0}, index=working.index)
    model = AbsorbingLS(
        dependent=working[column].astype(float),
        exog=exog,
        absorb=absorb,
        drop_absorbed=True,
    ).fit()
    residuals = pd.Series(np.nan, index=cells.index, dtype=float)
    residuals.loc[working.index] = np.asarray(model.resids)
    return residuals


def summarize_exposures(cells: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    support_rows: list[dict[str, object]] = []
    residual_frame = cells[
        ["survey_wave", "survey_year", "survey_month", "admin2_code"]
    ].copy()
    for term, metadata in EXPOSURES.items():
        column = metadata["column"]
        values = cells[column]
        residuals = residualize_exposure(cells, column)
        residual_frame[column] = values
        residual_frame[f"residual_{column}"] = residuals
        observed = values.dropna()
        residual_observed = residuals.dropna()
        quantiles = observed.quantile([0.05, 0.25, 0.50, 0.75, 0.95])
        residual_quantiles = residual_observed.quantile(
            [0.05, 0.25, 0.50, 0.75, 0.95]
        )
        raw_sd = float(observed.std(ddof=0))
        residual_sd = float(residual_observed.std(ddof=0))
        support_rows.append(
            {
                "term": term,
                "exposure_definition": metadata["label"],
                "n_cells": int(len(cells)),
                "observed_cells": int(observed.size),
                "missing_rate": float(values.isna().mean()),
                "mean_days": float(observed.mean()),
                "raw_sd_days": raw_sd,
                "minimum_days": float(observed.min()),
                "p05_days": float(quantiles.loc[0.05]),
                "p25_days": float(quantiles.loc[0.25]),
                "median_days": float(quantiles.loc[0.50]),
                "p75_days": float(quantiles.loc[0.75]),
                "p95_days": float(quantiles.loc[0.95]),
                "maximum_days": float(observed.max()),
                "zero_share": float(observed.eq(0).mean()),
                "nonzero_cells": int(observed.gt(0).sum()),
                "unique_values": int(observed.nunique()),
                "residual_sd_days": residual_sd,
                "residual_to_raw_sd_ratio": (
                    residual_sd / raw_sd if raw_sd > 0 else float("nan")
                ),
                "residual_p05_days": float(residual_quantiles.loc[0.05]),
                "residual_p25_days": float(residual_quantiles.loc[0.25]),
                "residual_median_days": float(residual_quantiles.loc[0.50]),
                "residual_p75_days": float(residual_quantiles.loc[0.75]),
                "residual_p95_days": float(residual_quantiles.loc[0.95]),
            }
        )
    return pd.DataFrame(support_rows), residual_frame


def summarize_by_wave(cells: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for term, metadata in EXPOSURES.items():
        column = metadata["column"]
        for (wave, year), group in cells.groupby(
            ["survey_wave", "survey_year"], sort=True, dropna=False
        ):
            values = group[column].dropna()
            rows.append(
                {
                    "term": term,
                    "exposure_definition": metadata["label"],
                    "survey_wave": wave,
                    "survey_year": int(year),
                    "n_cells": int(len(group)),
                    "mean_days": float(values.mean()),
                    "raw_sd_days": float(values.std(ddof=0)),
                    "p25_days": float(values.quantile(0.25)),
                    "median_days": float(values.median()),
                    "p75_days": float(values.quantile(0.75)),
                    "zero_share": float(values.eq(0).mean()),
                }
            )
    return pd.DataFrame(rows)


def add_model_rows(
    rows: list[dict[str, object]],
    frame: pd.DataFrame,
    *,
    specification: str,
    specification_label: str,
    outcome: str,
    outcome_label: str,
    controls: list[str],
    mask: pd.Series,
) -> None:
    for term in EXPOSURES:
        result, _model = fit_hdfe(
            frame,
            f"{specification}_{term}",
            outcome,
            [term],
            controls,
            sample_label="common sample for both exposure definitions",
            mask=mask,
        )
        for record in result:
            record["model_family"] = specification
            record["specification_label"] = specification_label
            record["outcome_label"] = outcome_label
            record["exposure_definition"] = EXPOSURES[term]["label"]
            rows.append(record)


def run_comparison_models(frame: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    both_exposures = frame[["wb26_5days", "tmax35_5days"]].notna().all(axis=1)
    rain_controls = [*DEMOGRAPHIC_CONTROLS, "precipitation_100mm"]
    specifications = [
        (
            "primary_work_rainfall",
            "Work participation: demographics + rainfall",
            "worked_past_week",
            "Worked in past week",
            rain_controls,
            frame["worked_past_week"].notna() & both_exposures,
        ),
        (
            "work_no_rainfall",
            "Work participation: demographics only",
            "worked_past_week",
            "Worked in past week",
            DEMOGRAPHIC_CONTROLS,
            frame["worked_past_week"].notna() & both_exposures,
        ),
        (
            "work_no_demographics",
            "Work participation: rainfall only",
            "worked_past_week",
            "Worked in past week",
            ["precipitation_100mm"],
            frame["worked_past_week"].notna() & both_exposures,
        ),
        (
            "weekly_hours_rainfall",
            "Weekly hours: demographics + rainfall",
            "weekly_hours_including_zero",
            "Weekly hours including zero",
            rain_controls,
            frame["weekly_hours_including_zero"].notna() & both_exposures,
        ),
        (
            "log_wage_rainfall",
            "Log wage: demographics + rainfall",
            "log_monthly_salary_wages",
            "Log monthly salary/wages",
            rain_controls,
            frame["log_monthly_salary_wages"].notna() & both_exposures,
        ),
    ]
    for specification, label, outcome, outcome_label, controls, mask in specifications:
        add_model_rows(
            rows,
            frame,
            specification=specification,
            specification_label=label,
            outcome=outcome,
            outcome_label=outcome_label,
            controls=controls,
            mask=mask,
        )

    joint_mask = frame["worked_past_week"].notna() & both_exposures
    joint, _model = fit_hdfe(
        frame,
        "joint_primary_work_both_exposures",
        "worked_past_week",
        list(EXPOSURES),
        rain_controls,
        sample_label="common sample; both exposure definitions entered jointly",
        mask=joint_mask,
    )
    for record in joint:
        record["model_family"] = "joint_primary_work"
        record["specification_label"] = (
            "Work participation: both exposures + demographics + rainfall"
        )
        record["outcome_label"] = "Worked in past week"
        record["exposure_definition"] = EXPOSURES[record["term"]]["label"]
        rows.append(record)

    models = pd.DataFrame(rows)
    scale = np.where(models["outcome"].eq("worked_past_week"), 100.0, 1.0)
    models["effect_display"] = models["coefficient"] * scale
    models["se_display"] = models["standard_error"] * scale
    models["ci_lower_display"] = models["ci_lower_95"] * scale
    models["ci_upper_display"] = models["ci_upper_95"] * scale
    models["effect_unit"] = np.select(
        [
            models["outcome"].eq("worked_past_week"),
            models["outcome"].eq("weekly_hours_including_zero"),
            models["outcome"].eq("log_monthly_salary_wages"),
        ],
        [
            "percentage points per 5 exposure days",
            "hours per week per 5 exposure days",
            "log points per 5 exposure days",
        ],
        default="outcome units per 5 exposure days",
    )
    columns = [
        "model_family",
        "specification_label",
        "outcome",
        "outcome_label",
        "term",
        "exposure_definition",
        "sample",
        "coefficient",
        "standard_error",
        "ci_lower_95",
        "ci_upper_95",
        "p_value",
        "effect_display",
        "se_display",
        "ci_lower_display",
        "ci_upper_display",
        "effect_unit",
        "n_obs",
        "n_admin2",
        "n_provinces",
        "cluster_level",
        "weighted_outcome_mean",
        "controls",
        "fixed_effects",
    ]
    return models[columns]


def interpretation_for_primary(effect: float, lower: float, upper: float) -> str:
    if effect < 0 and upper < 0:
        return "Adverse association; 95% interval excludes zero"
    if effect < 0:
        return "Adverse direction; 95% interval includes zero"
    if effect > 0 and lower > 0:
        return "Positive association; 95% interval excludes zero"
    return "Non-adverse direction; 95% interval includes zero"


def make_decision_data(
    support: pd.DataFrame, models: pd.DataFrame
) -> pd.DataFrame:
    primary = models.loc[models["model_family"].eq("primary_work_rainfall")]
    rows: list[dict[str, object]] = []
    for term, metadata in EXPOSURES.items():
        support_row = support.loc[support["term"].eq(term)].iloc[0]
        model_row = primary.loc[primary["term"].eq(term)].iloc[0]
        rows.append(
            {
                "term": term,
                "exposure_definition": metadata["label"],
                "literature_role": metadata["literature_role"],
                "mean_days": float(support_row["mean_days"]),
                "zero_share": float(support_row["zero_share"]),
                "residual_sd_days": float(support_row["residual_sd_days"]),
                "primary_effect_pp": float(model_row["effect_display"]),
                "ci_lower_pp": float(model_row["ci_lower_display"]),
                "ci_upper_pp": float(model_row["ci_upper_display"]),
                "p_value": float(model_row["p_value"]),
                "n_obs": int(model_row["n_obs"]),
                "screening_interpretation": interpretation_for_primary(
                    float(model_row["effect_display"]),
                    float(model_row["ci_lower_display"]),
                    float(model_row["ci_upper_display"]),
                ),
            }
        )
    return pd.DataFrame(rows)


def set_cell_style(
    cell: object,
    *,
    fill: str | None = None,
    color: str = TEXT,
    bold: bool = False,
    size: int = 9,
    align: str = "center",
    wrap: bool = False,
) -> None:
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.alignment = Alignment(
        horizontal=align, vertical="center", wrap_text=wrap
    )


def configure_page(sheet: object, print_area: str) -> None:
    sheet.print_area = print_area
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.20
    sheet.page_margins.right = 0.20
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25


def write_frame_sheet(
    workbook: Workbook,
    name: str,
    frame: pd.DataFrame,
    widths: dict[str, float] | None = None,
) -> object:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append(
            [None if pd.isna(value) else value for value in row]
        )
    for cell in sheet[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, wrap=True)
    sheet.row_dimensions[1].height = 34
    thin = Side(style="thin", color=LIGHT_BORDER)
    for row_index in range(2, sheet.max_row + 1):
        fill = PALE_BLUE if row_index % 2 == 0 else WHITE
        for cell in sheet[row_index]:
            set_cell_style(
                cell,
                fill=fill,
                size=8,
                align="left" if isinstance(cell.value, str) else "right",
                wrap=isinstance(cell.value, str),
            )
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[row_index].height = 22
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        if widths and letter in widths:
            width = widths[letter]
        else:
            sample = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 30) + 1)]
            width = min(max(max(map(len, sample), default=8) + 2, 11), 32)
        sheet.column_dimensions[letter].width = width
    configure_page(sheet, sheet.dimensions)
    return sheet


def build_workbook(
    decision: pd.DataFrame,
    support: pd.DataFrame,
    models: pd.DataFrame,
    wave_support: pd.DataFrame,
    output: Path,
    args: argparse.Namespace,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Decision Summary"
    summary.sheet_view.showGridLines = False
    headers = [
        "Exposure Definition",
        "Literature Role",
        "Mean Days",
        "Zero Share",
        "Residual SD",
        "Work Effect (pp / 5 days)",
        "95% CI Lower",
        "95% CI Upper",
        "p-value",
        "Screening Interpretation",
    ]
    summary.merge_cells("A1:J1")
    summary["A1"] = "Exposure Definition Comparison Audit"
    set_cell_style(
        summary["A1"], fill=NAVY, color=WHITE, bold=True, size=15, align="left"
    )
    summary.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        set_cell_style(
            summary.cell(2, column, header),
            fill=PALE_TEAL,
            bold=True,
            size=9,
            wrap=True,
        )
    summary.row_dimensions[2].height = 42

    support_sheet = write_frame_sheet(
        workbook,
        "Exposure Support",
        support,
        widths={"A": 16, "B": 25},
    )
    model_sheet = write_frame_sheet(
        workbook,
        "Model Comparison",
        models,
        widths={"A": 24, "B": 40, "D": 27, "F": 25, "G": 33, "Q": 34, "W": 38, "X": 35},
    )
    write_frame_sheet(
        workbook,
        "Wave Support",
        wave_support,
        widths={"A": 16, "B": 25, "C": 14},
    )

    support_columns = {name: index + 1 for index, name in enumerate(support.columns)}
    model_columns = {name: index + 1 for index, name in enumerate(models.columns)}
    support_rows = {
        support.iloc[index]["term"]: index + 2 for index in range(len(support))
    }
    primary_rows = {
        row["term"]: index + 2
        for index, row in models.reset_index(drop=True).iterrows()
        if row["model_family"] == "primary_work_rainfall"
    }
    thin = Side(style="thin", color=LIGHT_BORDER)
    for offset, record in enumerate(decision.itertuples(index=False), start=3):
        support_row = support_rows[record.term]
        model_row = primary_rows[record.term]
        formulas: list[object] = [
            record.exposure_definition,
            record.literature_role,
            f"='Exposure Support'!{get_column_letter(support_columns['mean_days'])}{support_row}",
            f"='Exposure Support'!{get_column_letter(support_columns['zero_share'])}{support_row}",
            f"='Exposure Support'!{get_column_letter(support_columns['residual_sd_days'])}{support_row}",
            f"='Model Comparison'!{get_column_letter(model_columns['effect_display'])}{model_row}",
            f"='Model Comparison'!{get_column_letter(model_columns['ci_lower_display'])}{model_row}",
            f"='Model Comparison'!{get_column_letter(model_columns['ci_upper_display'])}{model_row}",
            f"='Model Comparison'!{get_column_letter(model_columns['p_value'])}{model_row}",
            record.screening_interpretation,
        ]
        fill = WHITE if offset % 2 else PALE_BLUE
        for column, value in enumerate(formulas, start=1):
            cell = summary.cell(offset, column, value)
            set_cell_style(
                cell,
                fill=fill,
                size=9,
                align="left" if column in [1, 2, 10] else "right",
                wrap=column in [1, 2, 10],
            )
            cell.border = Border(bottom=thin)
        summary.row_dimensions[offset].height = 54
    for column in [3, 5, 6, 7, 8]:
        for row in range(3, 5):
            summary.cell(row, column).number_format = "0.000"
    for row in range(3, 5):
        summary.cell(row, 4).number_format = "0.0%"
        summary.cell(row, 9).number_format = "0.000"

    notes = [
        "Notes: Work effects are percentage points per five additional threshold-exceedance days. All paired estimates use the same observations, controls, weights, fixed effects, and district-clustered inference.",
        "Fixed effects: district-by-calendar-month and survey year. Primary controls: age, age squared, female, rural, and monthly precipitation. Database access: read-only SELECT; writes: none.",
        "This is a temporary screening audit. Tmax >= 35 C is a dry-heat definition; WBmax >= 26 C is a humid-heat burden definition and must not be described as a WBGT occupational-safety threshold.",
    ]
    for index, note in enumerate(notes, start=6):
        summary.merge_cells(start_row=index, start_column=1, end_row=index, end_column=10)
        summary.cell(index, 1, note)
        set_cell_style(summary.cell(index, 1), color=MUTED, size=8, align="left", wrap=True)
        summary.row_dimensions[index].height = 27
    widths = [24, 42, 12, 12, 13, 18, 14, 14, 11, 36]
    for column, width in enumerate(widths, start=1):
        summary.column_dimensions[get_column_letter(column)].width = width
    summary.freeze_panes = "A3"
    configure_page(summary, "A1:J8")
    summary.oddFooter.center.text = "Temporary diagnostic — read-only database input"

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    metadata_rows = [
        ("Field", "Value"),
        ("Generated At", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Status", "Temporary diagnostic; not a formal AnaSOP output"),
        ("Database", f"{args.dbname}.{args.schema}"),
        ("Database Access", "read-only SELECT; database writes: none"),
        ("Database Source", "final_HEAT_LABOR_ANALYTIC joined to dim_admin2_cambodia"),
        ("Exposure Unit", "five additional threshold-exceedance days in matched survey year-month"),
        ("Weights", "positive person analysis weights"),
        ("Fixed Effects", "admin2-by-calendar-month and survey year"),
        ("Inference", "debiased standard errors clustered by admin2"),
        ("Common-Sample Rule", "both exposure definitions observed; identical outcome/control completeness within each paired model"),
        ("Literature Source", "https://www.journals.uchicago.edu/doi/10.1086/713733"),
        ("Interpretation Boundary", "WBmax is wet-bulb temperature, not WBGT; no occupational stop-work threshold is claimed"),
        ("Code", str(Path("src/analyses/audit_exposure_definition_comparison.py"))),
    ]
    for row in metadata_rows:
        metadata.append(list(row))
    for cell in metadata[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True)
    for row in range(2, metadata.max_row + 1):
        set_cell_style(metadata.cell(row, 1), bold=True, align="left")
        set_cell_style(metadata.cell(row, 2), align="left", wrap=True)
        metadata.row_dimensions[row].height = 25
    metadata.column_dimensions["A"].width = 27
    metadata.column_dimensions["B"].width = 100
    configure_page(metadata, metadata.dimensions)

    for sheet in [support_sheet, model_sheet]:
        for row in range(2, sheet.max_row + 1):
            for cell in sheet[row]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)


def make_comparison_figure(
    residuals: pd.DataFrame,
    wave_support: pd.DataFrame,
    models: pd.DataFrame,
    output: Path,
) -> None:
    sns.set_theme(style="whitegrid", context="notebook")
    colors = {"wb26_5days": "#1D6F7A", "tmax35_5days": "#C66A2B"}
    fig, axes = plt.subplots(1, 3, figsize=(15.2, 4.7))

    for term, metadata in EXPOSURES.items():
        values = residuals[metadata["column"]].dropna()
        sns.ecdfplot(
            values,
            ax=axes[0],
            label=metadata["short_label"],
            color=colors[term],
            linewidth=2.2,
        )
    axes[0].set_xlabel("Threshold-exceedance days per month")
    axes[0].set_ylabel("Cumulative share of district-month cells")
    axes[0].text(-0.13, 1.04, "a", transform=axes[0].transAxes, fontweight="bold")

    for term, metadata in EXPOSURES.items():
        subset = wave_support.loc[wave_support["term"].eq(term)].sort_values(
            "survey_year"
        )
        axes[1].plot(
            subset["survey_year"],
            subset["mean_days"],
            marker="o",
            linewidth=2,
            color=colors[term],
            label=metadata["short_label"],
        )
    axes[1].set_xlabel("Survey wave year")
    axes[1].set_ylabel("Mean threshold-exceedance days")
    axes[1].set_xticks(
        sorted(wave_support["survey_year"].dropna().astype(int).unique())
    )
    axes[1].tick_params(axis="x", rotation=45)
    axes[1].text(-0.13, 1.04, "b", transform=axes[1].transAxes, fontweight="bold")

    forest = models.loc[
        models["model_family"].isin(
            ["primary_work_rainfall", "work_no_rainfall", "work_no_demographics"]
        )
    ].copy()
    spec_order = [
        "primary_work_rainfall",
        "work_no_rainfall",
        "work_no_demographics",
    ]
    spec_labels = {
        "primary_work_rainfall": "Demographics + rainfall",
        "work_no_rainfall": "Demographics only",
        "work_no_demographics": "Rainfall only",
    }
    base_positions = np.arange(len(spec_order))[::-1]
    offsets = {"wb26_5days": 0.11, "tmax35_5days": -0.11}
    for term, metadata in EXPOSURES.items():
        subset = forest.loc[forest["term"].eq(term)].set_index("model_family").loc[spec_order]
        y = base_positions + offsets[term]
        estimate = subset["effect_display"].to_numpy()
        lower = subset["ci_lower_display"].to_numpy()
        upper = subset["ci_upper_display"].to_numpy()
        axes[2].errorbar(
            estimate,
            y,
            xerr=np.vstack([estimate - lower, upper - estimate]),
            fmt="o",
            capsize=3,
            color=colors[term],
            label=metadata["short_label"],
        )
    axes[2].axvline(0, color="#4D4D4D", linewidth=1, linestyle="--")
    axes[2].set_yticks(base_positions, [spec_labels[item] for item in spec_order])
    axes[2].set_xlabel("Work participation effect (percentage points / 5 days)")
    axes[2].text(-0.13, 1.04, "c", transform=axes[2].transAxes, fontweight="bold")

    for axis in axes:
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(
        handles,
        labels,
        frameon=False,
        fontsize=9,
        ncol=2,
        loc="lower center",
        bbox_to_anchor=(0.5, -0.015),
    )
    fig.tight_layout(w_pad=2.0, rect=(0, 0.08, 1, 1))
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def render_summary_png(workbook_path: Path, png_path: Path) -> None:
    soffice = Path("/opt/homebrew/bin/soffice")
    if not soffice.exists():
        located = shutil.which("soffice")
        if not located:
            raise RuntimeError("LibreOffice/soffice is required for workbook rendering")
        soffice = Path(located)
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("pdftoppm is required for PNG rendering")
    with tempfile.TemporaryDirectory(prefix="mj02b-exposure-audit-") as temporary:
        temporary_path = Path(temporary)
        profile_uri = (temporary_path / "lo-profile").as_uri()
        subprocess.run(
            [
                str(soffice),
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temporary_path),
                str(workbook_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        pdf_path = temporary_path / f"{workbook_path.stem}.pdf"
        prefix = png_path.with_suffix("")
        subprocess.run(
            [pdftoppm, "-png", "-r", "180", "-f", "1", "-singlefile", str(pdf_path), str(prefix)],
            check=True,
            capture_output=True,
            text=True,
        )
    if not png_path.exists() or png_path.stat().st_size == 0:
        raise RuntimeError("Workbook review PNG was not created")


def validate_outputs(
    decision: pd.DataFrame,
    support: pd.DataFrame,
    models: pd.DataFrame,
    workbook_path: Path,
) -> dict[str, object]:
    if len(support) != 2 or set(support["term"]) != set(EXPOSURES):
        raise AssertionError("Expected support rows for exactly two exposures")
    if not support["n_cells"].eq(3_325).all():
        raise AssertionError("Unexpected district-year-month cell count")
    primary = models.loc[models["model_family"].eq("primary_work_rainfall")]
    if len(primary) != 2 or primary["n_obs"].nunique() != 1:
        raise AssertionError("Primary paired models do not use one common sample")
    for family, group in models.groupby("model_family"):
        if family != "joint_primary_work" and len(group) == 2 and group["n_obs"].nunique() != 1:
            raise AssertionError(f"Paired model sample differs in {family}")
    workbook = load_workbook(workbook_path, data_only=False)
    if workbook.sheetnames != [
        "Decision Summary",
        "Exposure Support",
        "Model Comparison",
        "Wave Support",
        "Metadata",
    ]:
        raise AssertionError("Unexpected workbook sheet order")
    summary = workbook["Decision Summary"]
    formula_count = sum(
        1
        for row in summary.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    if formula_count != 14:
        raise AssertionError(f"Expected 14 summary formulas, observed {formula_count}")
    formula_errors = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
        and any(error in cell.value for error in ["#REF!", "#DIV/0!", "#VALUE!", "#NAME?"])
    ]
    if formula_errors:
        raise AssertionError(f"Workbook contains formula errors: {formula_errors}")
    return {
        "district_month_cells": int(support.iloc[0]["n_cells"]),
        "primary_model_n": int(primary.iloc[0]["n_obs"]),
        "model_rows": int(len(models)),
        "summary_formula_count": formula_count,
        "workbook_sheet_count": len(workbook.sheetnames),
        "decision_rows": int(len(decision)),
    }


def write_readme(
    output: Path,
    decision: pd.DataFrame,
    support: pd.DataFrame,
    models: pd.DataFrame,
) -> None:
    primary = models.loc[models["model_family"].eq("primary_work_rainfall")]
    records = {row.term: row for row in primary.itertuples(index=False)}
    supports = {row.term: row for row in support.itertuples(index=False)}
    recommendation = (
        "The current comparison does not support replacing the humid-heat measure "
        "with Tmax >= 35 C as the sole primary exposure: the humid-heat estimate is "
        "adverse in direction, whereas the dry-heat estimate is positive in direction, "
        "and both intervals include zero. Retain WBmax >= 26 C as the carefully labeled "
        "humid-heat primary measure for now and use Tmax >= 35 C as a dry-heat comparison. "
        "Revisit the primary definition only after the joint model, timing tests, and "
        "locally relative heat metrics are assessed together."
    )
    lines = [
        "# Exposure-definition comparison audit",
        "",
        "Status: temporary diagnostic; not a formal AnaSOP output.",
        "",
        "## Result",
        "",
        recommendation,
        "",
        "Under the common-sample primary work-participation model:",
        "",
        (
            f"- WBmax >= 26 C: {records['wb26_5days'].effect_display:.3f} percentage "
            f"points per five days (95% CI {records['wb26_5days'].ci_lower_display:.3f} "
            f"to {records['wb26_5days'].ci_upper_display:.3f}; "
            f"p={records['wb26_5days'].p_value:.3f})."
        ),
        (
            f"- Tmax >= 35 C: {records['tmax35_5days'].effect_display:.3f} percentage "
            f"points per five days (95% CI {records['tmax35_5days'].ci_lower_display:.3f} "
            f"to {records['tmax35_5days'].ci_upper_display:.3f}; "
            f"p={records['tmax35_5days'].p_value:.3f})."
        ),
        "",
        "Identification support after absorbing district-by-calendar-month and survey-year fixed effects:",
        "",
        (
            f"- WBmax >= 26 C residual SD: {supports['wb26_5days'].residual_sd_days:.3f} days; "
            f"zero share: {supports['wb26_5days'].zero_share:.1%}."
        ),
        (
            f"- Tmax >= 35 C residual SD: {supports['tmax35_5days'].residual_sd_days:.3f} days; "
            f"zero share: {supports['tmax35_5days'].zero_share:.1%}."
        ),
        "",
        "## Design contract",
        "",
        "All paired estimates use the same sample, positive analysis weights, controls, district-by-calendar-month and survey-year fixed effects, and district-clustered debiased standard errors. The database was queried with read-only SELECT and was not modified.",
        "",
        "WBmax denotes wet-bulb temperature, not WBGT. The 26 C definition is a humid-heat burden measure and is not interpreted as an occupational stop-work threshold. The 35 C Tmax definition has direct precedent in labor-productivity research: https://www.journals.uchicago.edu/doi/10.1086/713733",
        "",
        "## Files",
        "",
        "- `Exposure_definition_comparison_audit.xlsx`: decision summary, exposure support, model results, wave support, and metadata.",
        "- `Exposure_definition_comparison_audit.png`: rendered first-sheet review copy.",
        "- `Exposure_definition_comparison.png`: scientific three-panel comparison.",
        "- `exposure_support.csv`, `wave_support.csv`, `model_comparison.csv`, `decision_summary.csv`, and `cell_exposure_residuals.csv`: machine-readable audit data.",
        "- `audit_manifest.json`: validation and provenance summary.",
    ]
    (output / "README.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    args = parse_args()
    output = args.output if args.output.is_absolute() else ROOT / args.output
    output.mkdir(parents=True, exist_ok=True)
    with psycopg.connect(**connection_args(args)) as connection:
        frame = load_analysis_data(connection, args.schema)
        connection.rollback()
    if not frame.shape[0]:
        raise RuntimeError("No analytical rows were returned")

    cells = make_cells(frame)
    support, residuals = summarize_exposures(cells)
    wave_support = summarize_by_wave(cells)
    models = run_comparison_models(frame)
    decision = make_decision_data(support, models)

    support.to_csv(output / "exposure_support.csv", index=False)
    residuals.to_csv(output / "cell_exposure_residuals.csv", index=False)
    wave_support.to_csv(output / "wave_support.csv", index=False)
    models.to_csv(output / "model_comparison.csv", index=False)
    decision.to_csv(output / "decision_summary.csv", index=False)

    workbook_path = output / "Exposure_definition_comparison_audit.xlsx"
    build_workbook(decision, support, models, wave_support, workbook_path, args)
    make_comparison_figure(
        residuals,
        wave_support,
        models,
        output / "Exposure_definition_comparison.png",
    )
    render_summary_png(
        workbook_path,
        output / "Exposure_definition_comparison_audit.png",
    )
    validation = validate_outputs(decision, support, models, workbook_path)
    write_readme(output, decision, support, models)
    expected_outputs = [
        "Exposure_definition_comparison.png",
        "Exposure_definition_comparison_audit.png",
        "Exposure_definition_comparison_audit.xlsx",
        "README.md",
        "audit_manifest.json",
        "cell_exposure_residuals.csv",
        "decision_summary.csv",
        "exposure_support.csv",
        "model_comparison.csv",
        "wave_support.csv",
    ]
    manifest = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "status": "temporary diagnostic; not a formal AnaSOP output",
        "database_access": "read-only SELECT; writes none",
        "source_table": f"{args.dbname}.{args.schema}.final_HEAT_LABOR_ANALYTIC",
        "code": "src/analyses/audit_exposure_definition_comparison.py",
        "validation": validation,
        "outputs": expected_outputs,
    }
    (output / "audit_manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(manifest, indent=2, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
