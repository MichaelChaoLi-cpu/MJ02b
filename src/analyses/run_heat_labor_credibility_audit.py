#!/usr/bin/env python3
"""Run a read-only credibility audit for the CSES month-level heat-labor result.

Every analytical input is obtained with SELECT from mda.public. The PostgreSQL
session is forced into read-only mode at connection start. Outputs are local
diagnostic artifacts, not formal AnaSOP figures or tables.
"""

from __future__ import annotations

import argparse
import json
import os
from datetime import datetime
from pathlib import Path
from typing import Iterable

ROOT = Path(__file__).resolve().parents[2]
MPL_CACHE = ROOT / "data" / "exp" / ".matplotlib"
MPL_CACHE.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(MPL_CACHE))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import psycopg
import seaborn as sns
from linearmodels.iv import AbsorbingLS
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from psycopg import sql


DEFAULT_OUTPUT = Path("data/exp/experiments/heat-labor-credibility-audit")
DEMOGRAPHIC_CONTROLS = ["age", "age_squared", "female", "rural"]
HEAT_BIN_LABELS = ["0", "1-5", "6-10", "11-15", "16-20", "21-25", "26-31"]
HEAT_BIN_TERMS = [
    "heat_bin_1_5",
    "heat_bin_6_10",
    "heat_bin_11_15",
    "heat_bin_16_20",
    "heat_bin_21_25",
    "heat_bin_26_31",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def connection_args(args: argparse.Namespace) -> dict[str, object]:
    values: dict[str, object] = {
        "host": args.host,
        "port": args.port,
        "dbname": args.dbname,
        "options": "-c default_transaction_read_only=on",
    }
    if args.user:
        values["user"] = args.user
    if args.password:
        values["password"] = args.password
    return values


def fetch_frame(
    connection: psycopg.Connection,
    query: sql.Composed,
) -> pd.DataFrame:
    with connection.cursor() as cursor:
        cursor.execute(query)
        columns = [column.name for column in cursor.description]
        rows = cursor.fetchall()
    return pd.DataFrame(rows, columns=columns)


def load_analysis_data(connection: psycopg.Connection, schema: str) -> pd.DataFrame:
    query = sql.SQL(
        """
        SELECT
            a.survey_wave, a.survey_year, a.survey_month,
            a.person_id, a.household_id, a.admin2_code,
            d.admin1_code, d.admin1_name, d.centroid_lat, d.centroid_lon,
            a.sex, a.age, a.urban_rural, a.analysis_weight,
            a.worked_past_week, a.weekly_hours_including_zero,
            a.log_monthly_salary_wages,
            a.days_wbmax_ge_26c, a.days_wbmax_ge_28c,
            a.days_wbmax_ge_30c, a.wet_bulb_max_mean_c,
            a.wet_bulb_month_max_c, a.temperature_2m_max_mean_c,
            a.temperature_2m_month_max_c, a.days_tmax_ge_35c,
            a.days_tmax_ge_37c, a.apparent_temperature_max_mean_c,
            a.apparent_temperature_month_max_c,
            a.precipitation_month_sum_mm,
            a.lag_days_wbmax_ge_26c, a.lead_days_wbmax_ge_26c,
            a.temperature_location_fallback_used
        FROM {}.{} a
        LEFT JOIN {}.{} d USING (admin2_code)
        WHERE a.analysis_weight > 0
        ORDER BY a.survey_year, a.person_id
        """
    ).format(
        sql.Identifier(schema), sql.Identifier("final_HEAT_LABOR_ANALYTIC"),
        sql.Identifier(schema), sql.Identifier("dim_admin2_cambodia"),
    )
    frame = fetch_frame(connection, query)
    numeric = [
        "survey_year", "survey_month", "sex", "age", "urban_rural",
        "analysis_weight", "worked_past_week", "weekly_hours_including_zero",
        "log_monthly_salary_wages", "days_wbmax_ge_26c",
        "days_wbmax_ge_28c", "days_wbmax_ge_30c",
        "wet_bulb_max_mean_c", "wet_bulb_month_max_c",
        "temperature_2m_max_mean_c", "temperature_2m_month_max_c",
        "days_tmax_ge_35c", "days_tmax_ge_37c",
        "apparent_temperature_max_mean_c",
        "apparent_temperature_month_max_c", "precipitation_month_sum_mm",
        "lag_days_wbmax_ge_26c", "lead_days_wbmax_ge_26c",
        "centroid_lat", "centroid_lon",
    ]
    for column in numeric:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")

    frame["age_squared"] = frame["age"] ** 2
    frame["female"] = frame["sex"].eq(2).where(frame["sex"].isin([1, 2]))
    frame["rural"] = frame["urban_rural"].eq(2).where(
        frame["urban_rural"].isin([1, 2])
    )
    frame["admin2_calendar_month"] = (
        frame["admin2_code"].astype("string")
        + "-m"
        + frame["survey_month"].astype("Int64").astype("string").str.zfill(2)
    )
    frame["wb26_5days"] = frame["days_wbmax_ge_26c"] / 5.0
    frame["wb28_5days"] = frame["days_wbmax_ge_28c"] / 5.0
    frame["tmax35_5days"] = frame["days_tmax_ge_35c"] / 5.0
    frame["tmax37_5days"] = frame["days_tmax_ge_37c"] / 5.0
    frame["lag_wb26_5days"] = frame["lag_days_wbmax_ge_26c"] / 5.0
    frame["lead_wb26_5days"] = frame["lead_days_wbmax_ge_26c"] / 5.0
    frame["precipitation_100mm"] = frame["precipitation_month_sum_mm"] / 100.0
    frame["hours_observed"] = frame["weekly_hours_including_zero"].notna().astype(float)
    frame["wage_observed"] = frame["log_monthly_salary_wages"].notna().astype(float)
    frame["fallback_excluded"] = ~frame[
        "temperature_location_fallback_used"
    ].fillna(False).astype(bool)

    heat_bin = pd.cut(
        frame["days_wbmax_ge_26c"],
        bins=[-0.1, 0, 5, 10, 15, 20, 25, 31],
        labels=HEAT_BIN_LABELS,
        ordered=True,
    )
    frame["heat_bin"] = heat_bin
    for label, term in zip(HEAT_BIN_LABELS[1:], HEAT_BIN_TERMS, strict=True):
        frame[term] = heat_bin.eq(label).astype(float)
    return frame


def weighted_mean(values: pd.Series, weights: pd.Series) -> float:
    valid = values.notna() & weights.notna() & weights.gt(0)
    if not valid.any():
        return float("nan")
    return float(np.average(values[valid], weights=weights[valid]))


def fit_hdfe(
    frame: pd.DataFrame,
    specification: str,
    outcome: str,
    terms: list[str],
    controls: list[str],
    *,
    sample_label: str,
    cluster_columns: list[str] | None = None,
    mask: pd.Series | None = None,
) -> tuple[list[dict[str, object]], object]:
    cluster_columns = cluster_columns or ["admin2_code"]
    required = list(dict.fromkeys(
        [
            outcome, *terms, *controls, "analysis_weight",
            "admin2_calendar_month", "survey_year", *cluster_columns,
        ]
    ))
    selected = frame.loc[mask].copy() if mask is not None else frame.copy()
    selected = selected.dropna(subset=required)
    if selected.empty:
        raise RuntimeError(f"Empty model sample: {specification}")
    if selected["admin2_code"].nunique() < 30:
        raise RuntimeError(f"Too few admin2 areas: {specification}")

    exog = selected[terms + controls].astype(float).copy()
    exog["constant"] = 1.0
    absorb = pd.DataFrame(
        {
            "admin2_calendar_month": selected[
                "admin2_calendar_month"
            ].astype("category"),
            "survey_year": selected["survey_year"].astype("Int64").astype("category"),
        },
        index=selected.index,
    )
    clusters = selected[cluster_columns].copy()
    for column in cluster_columns:
        clusters[column] = clusters[column].astype("category").cat.codes
    model = AbsorbingLS(
        dependent=selected[outcome].astype(float),
        exog=exog,
        absorb=absorb,
        weights=selected["analysis_weight"].astype(float),
        drop_absorbed=True,
    ).fit(cov_type="clustered", clusters=clusters, debiased=True)

    results: list[dict[str, object]] = []
    for term in terms:
        estimate = float(model.params[term])
        standard_error = float(model.std_errors[term])
        results.append(
            {
                "specification": specification,
                "outcome": outcome,
                "term": term,
                "sample": sample_label,
                "coefficient": estimate,
                "standard_error": standard_error,
                "ci_lower_95": estimate - 1.96 * standard_error,
                "ci_upper_95": estimate + 1.96 * standard_error,
                "p_value": float(model.pvalues[term]),
                "n_obs": int(model.nobs),
                "n_admin2": int(selected["admin2_code"].nunique()),
                "n_provinces": int(selected["admin1_code"].nunique()),
                "cluster_level": "+".join(cluster_columns),
                "weighted_outcome_mean": weighted_mean(
                    selected[outcome], selected["analysis_weight"]
                ),
                "controls": "+".join(controls),
                "fixed_effects": "admin2-by-calendar-month + survey year",
            }
        )
    print(
        f"estimated={specification} n={int(model.nobs)} "
        f"admin2={selected['admin2_code'].nunique()}",
        flush=True,
    )
    return results, model


def exposure_cell_frame(frame: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "survey_year", "survey_month", "survey_wave", "admin2_code",
        "admin1_code", "days_wbmax_ge_26c", "days_wbmax_ge_28c",
        "wet_bulb_max_mean_c", "days_tmax_ge_35c",
        "apparent_temperature_max_mean_c", "precipitation_month_sum_mm",
        "heat_bin",
    ]
    return frame[columns].drop_duplicates(
        ["survey_year", "survey_month", "admin2_code"]
    ).reset_index(drop=True)


def exposure_summary(cells: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for variable in [
        "days_wbmax_ge_26c", "days_wbmax_ge_28c", "wet_bulb_max_mean_c",
        "days_tmax_ge_35c", "apparent_temperature_max_mean_c",
        "precipitation_month_sum_mm",
    ]:
        series = cells[variable].dropna()
        quantiles = series.quantile([0.05, 0.25, 0.5, 0.75, 0.95])
        rows.append(
            {
                "variable": variable,
                "n_cells": len(series),
                "mean": series.mean(),
                "std": series.std(),
                "min": series.min(),
                "p05": quantiles.loc[0.05],
                "p25": quantiles.loc[0.25],
                "median": quantiles.loc[0.5],
                "p75": quantiles.loc[0.75],
                "p95": quantiles.loc[0.95],
                "max": series.max(),
                "zero_share": series.eq(0).mean(),
            }
        )
    return pd.DataFrame(rows)


def wave_summary(frame: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for wave, group in frame.groupby("survey_wave", sort=False):
        rows.append(
            {
                "survey_wave": wave,
                "survey_year": int(group["survey_year"].iloc[0]),
                "person_rows": len(group),
                "admin2_areas": group["admin2_code"].nunique(),
                "district_month_cells": group[
                    ["admin2_code", "survey_year", "survey_month"]
                ].drop_duplicates().shape[0],
                "work_observed_rate": group["worked_past_week"].notna().mean(),
                "hours_observed_rate": group["hours_observed"].mean(),
                "wage_observed_rate": group["wage_observed"].mean(),
                "weighted_work_rate": weighted_mean(
                    group["worked_past_week"], group["analysis_weight"]
                ),
                "weighted_hours": weighted_mean(
                    group["weekly_hours_including_zero"], group["analysis_weight"]
                ),
                "mean_wb26_days": group["days_wbmax_ge_26c"].mean(),
                "mean_precipitation_mm": group["precipitation_month_sum_mm"].mean(),
            }
        )
    return pd.DataFrame(rows).sort_values("survey_year", ignore_index=True)


def heat_bin_summary(frame: pd.DataFrame, cells: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for label in HEAT_BIN_LABELS:
        people = frame.loc[frame["heat_bin"].eq(label)]
        cell_group = cells.loc[cells["heat_bin"].eq(label)]
        rows.append(
            {
                "heat_bin": label,
                "district_month_cells": len(cell_group),
                "person_rows": len(people),
                "weighted_work_rate": weighted_mean(
                    people["worked_past_week"], people["analysis_weight"]
                ),
                "hours_observed_rate": people["hours_observed"].mean(),
                "wage_observed_rate": people["wage_observed"].mean(),
                "weighted_mean_age": weighted_mean(
                    people["age"], people["analysis_weight"]
                ),
                "weighted_female_share": weighted_mean(
                    people["female"], people["analysis_weight"]
                ),
                "weighted_rural_share": weighted_mean(
                    people["rural"], people["analysis_weight"]
                ),
            }
        )
    return pd.DataFrame(rows)


def within_fe_exposure_variation(cells: pd.DataFrame) -> pd.DataFrame:
    working = cells.dropna(subset=["days_wbmax_ge_26c"]).copy()
    working["admin2_calendar_month"] = (
        working["admin2_code"].astype("string")
        + "-m"
        + working["survey_month"].astype("Int64").astype("string").str.zfill(2)
    )
    absorb = pd.DataFrame(
        {
            "admin2_calendar_month": working[
                "admin2_calendar_month"
            ].astype("category"),
            "survey_year": working["survey_year"].astype("Int64").astype("category"),
        },
        index=working.index,
    )
    exog = pd.DataFrame({"constant": 1.0}, index=working.index)
    model = AbsorbingLS(
        working["days_wbmax_ge_26c"].astype(float),
        exog,
        absorb=absorb,
        drop_absorbed=True,
    ).fit()
    residuals = pd.Series(model.resids, index=working.index)
    return pd.DataFrame(
        [
            {
                "n_district_month_cells": len(working),
                "raw_standard_deviation_days": working[
                    "days_wbmax_ge_26c"
                ].std(),
                "within_fe_standard_deviation_days": residuals.std(),
                "within_fe_p05_days": residuals.quantile(0.05),
                "within_fe_median_days": residuals.median(),
                "within_fe_p95_days": residuals.quantile(0.95),
            }
        ]
    )


def run_models(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    rows: list[dict[str, object]] = []
    base_mask = frame["worked_past_week"].notna()
    base_controls = DEMOGRAPHIC_CONTROLS
    rain_controls = [*DEMOGRAPHIC_CONTROLS, "precipitation_100mm"]

    specifications = [
        (
            "strict_no_rainfall", "worked_past_week", ["wb26_5days"],
            base_controls, "all working-age observations", ["admin2_code"], base_mask,
        ),
        (
            "strict_rainfall_primary", "worked_past_week", ["wb26_5days"],
            rain_controls, "all working-age observations", ["admin2_code"], base_mask,
        ),
        (
            "strict_rainfall_province_cluster", "worked_past_week", ["wb26_5days"],
            rain_controls, "all; province-cluster inference", ["admin1_code"], base_mask,
        ),
        (
            "strict_rainfall_no_demographics", "worked_past_week", ["wb26_5days"],
            ["precipitation_100mm"], "all; no person controls", ["admin2_code"], base_mask,
        ),
        (
            "strict_rainfall_exclude_coastal_fallback", "worked_past_week",
            ["wb26_5days"], rain_controls, "exclude fallback district",
            ["admin2_code"], base_mask & frame["fallback_excluded"],
        ),
    ]

    complete_waves = frame.groupby("admin2_code")["survey_wave"].nunique()
    balanced_admin2 = complete_waves.index[complete_waves.eq(frame["survey_wave"].nunique())]
    specifications.append(
        (
            "strict_rainfall_balanced_admin2", "worked_past_week", ["wb26_5days"],
            rain_controls, "admin2 observed in all nine waves", ["admin2_code"],
            base_mask & frame["admin2_code"].isin(balanced_admin2),
        )
    )

    for specification in specifications:
        result, _model = fit_hdfe(
            frame,
            specification[0],
            specification[1],
            specification[2],
            specification[3],
            sample_label=specification[4],
            cluster_columns=specification[5],
            mask=specification[6],
        )
        rows.extend(result)

    alternative_specs = [
        ("alternative_wb28_days", ["wb28_5days"]),
        ("alternative_mean_wbmax", ["wet_bulb_max_mean_c"]),
        ("alternative_tmax35_days", ["tmax35_5days"]),
        ("alternative_tmax37_days", ["tmax37_5days"]),
        ("alternative_apparent_temperature", ["apparent_temperature_max_mean_c"]),
    ]
    for name, terms in alternative_specs:
        result, _model = fit_hdfe(
            frame, name, "worked_past_week", terms, rain_controls,
            sample_label="all working-age observations",
            mask=base_mask,
        )
        rows.extend(result)

    for name, terms, mask in [
        (
            "temporal_current_and_lag",
            ["wb26_5days", "lag_wb26_5days"],
            base_mask & frame["lag_wb26_5days"].notna(),
        ),
        (
            "temporal_current_and_lead_placebo",
            ["wb26_5days", "lead_wb26_5days"],
            base_mask & frame["lead_wb26_5days"].notna(),
        ),
    ]:
        result, _model = fit_hdfe(
            frame, name, "worked_past_week", terms, rain_controls,
            sample_label="current and adjacent month available", mask=mask,
        )
        rows.extend(result)

    for name, outcome, mask in [
        (
            "secondary_weekly_hours", "weekly_hours_including_zero",
            frame["weekly_hours_including_zero"].notna(),
        ),
        ("selection_hours_observed", "hours_observed", pd.Series(True, index=frame.index)),
        ("selection_wage_observed", "wage_observed", pd.Series(True, index=frame.index)),
        (
            "exploratory_log_wage", "log_monthly_salary_wages",
            frame["log_monthly_salary_wages"].notna(),
        ),
    ]:
        result, _model = fit_hdfe(
            frame, name, outcome, ["wb26_5days"], rain_controls,
            sample_label=name, mask=mask,
        )
        rows.extend(result)

    high_hours_waves = (
        frame.groupby("survey_wave")["hours_observed"].mean().loc[lambda x: x >= 0.8].index
    )
    result, _model = fit_hdfe(
        frame,
        "secondary_weekly_hours_high_coverage_waves",
        "weekly_hours_including_zero",
        ["wb26_5days"],
        rain_controls,
        sample_label="waves with at least 80 percent hours coverage",
        mask=(
            frame["survey_wave"].isin(high_hours_waves)
            & frame["weekly_hours_including_zero"].notna()
        ),
    )
    rows.extend(result)

    balance_rows: list[dict[str, object]] = []
    balance_specs = [
        ("balance_age", "age", ["female", "rural", "precipitation_100mm"]),
        ("balance_female", "female", ["age", "age_squared", "rural", "precipitation_100mm"]),
        ("balance_rural", "rural", ["age", "age_squared", "female", "precipitation_100mm"]),
    ]
    for name, outcome, controls in balance_specs:
        result, _model = fit_hdfe(
            frame, name, outcome, ["wb26_5days"], controls,
            sample_label="composition balance sample",
        )
        balance_rows.extend(result)

    loo_rows: list[dict[str, object]] = []
    for wave in frame.sort_values("survey_year")["survey_wave"].drop_duplicates():
        result, _model = fit_hdfe(
            frame,
            f"leave_out_{wave}",
            "worked_past_week",
            ["wb26_5days"],
            rain_controls,
            sample_label=f"all waves except {wave}",
            mask=base_mask & frame["survey_wave"].ne(wave),
        )
        result[0]["excluded_wave"] = wave
        loo_rows.extend(result)

    dose_rows, _model = fit_hdfe(
        frame,
        "dose_response_absolute_bins",
        "worked_past_week",
        HEAT_BIN_TERMS,
        rain_controls,
        sample_label="absolute WBmax>=26C day bins; zero days reference",
        mask=base_mask,
    )
    for row, label in zip(dose_rows, HEAT_BIN_LABELS[1:], strict=True):
        row["heat_bin"] = label

    return (
        pd.DataFrame(rows),
        pd.DataFrame(loo_rows),
        pd.concat([pd.DataFrame(balance_rows), pd.DataFrame(dose_rows)], ignore_index=True),
    )


def plot_exposure_distribution(cells: pd.DataFrame, output: Path) -> None:
    sns.set_theme(style="whitegrid", context="notebook")
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8))
    sns.histplot(
        cells, x="days_wbmax_ge_26c", bins=np.arange(-0.5, 32.5, 1),
        color="#2a6f97", ax=axes[0],
    )
    axes[0].set_xlabel("Days with daily WBmax >= 26 C")
    axes[0].set_ylabel("District-month cells")
    sns.boxplot(
        cells, x="survey_year", y="days_wbmax_ge_26c",
        color="#61a5c2", fliersize=1.5, ax=axes[1],
    )
    axes[1].set_xlabel("Survey year")
    axes[1].set_ylabel("Days with daily WBmax >= 26 C")
    for label, axis in zip("ab", axes, strict=True):
        axis.text(
            -0.1, 1.04, label, transform=axis.transAxes,
            fontsize=12, fontweight="bold", va="top",
        )
    fig.tight_layout()
    fig.savefig(output, dpi=150, bbox_inches="tight")
    plt.close(fig)


def plot_forest(estimates: pd.DataFrame, output: Path) -> None:
    order = [
        ("strict_no_rainfall", "No rainfall control"),
        ("strict_rainfall_primary", "Primary + rainfall"),
        ("strict_rainfall_no_demographics", "No demographic controls"),
        ("strict_rainfall_exclude_coastal_fallback", "Exclude coastal fallback"),
        ("strict_rainfall_balanced_admin2", "Balanced districts"),
        ("strict_rainfall_province_cluster", "Province-cluster inference"),
    ]
    selected = estimates.loc[
        estimates["specification"].isin([name for name, _label in order])
        & estimates["term"].eq("wb26_5days")
    ].copy()
    labels = dict(order)
    selected["label"] = selected["specification"].map(labels)
    selected["sort"] = selected["specification"].map(
        {name: index for index, (name, _label) in enumerate(order)}
    )
    selected = selected.sort_values("sort", ascending=False)
    fig, axis = plt.subplots(figsize=(8, 4.8))
    axis.errorbar(
        selected["coefficient"] * 100,
        np.arange(len(selected)),
        xerr=1.96 * selected["standard_error"] * 100,
        fmt="o", color="#1d3557", ecolor="#457b9d", capsize=3,
    )
    axis.axvline(0, color="black", linewidth=0.8, linestyle="--")
    axis.set_yticks(np.arange(len(selected)), selected["label"])
    axis.set_xlabel("Percentage-point change per 5 additional WBmax >= 26 C days")
    axis.set_ylabel("")
    fig.tight_layout()
    fig.savefig(output, dpi=150, bbox_inches="tight")
    plt.close(fig)


def plot_leave_one_wave_out(loo: pd.DataFrame, output: Path) -> None:
    selected = loo.sort_values("excluded_wave", ascending=False).reset_index(drop=True)
    fig, axis = plt.subplots(figsize=(8, 5.5))
    axis.errorbar(
        selected["coefficient"] * 100,
        np.arange(len(selected)),
        xerr=1.96 * selected["standard_error"] * 100,
        fmt="o", color="#6a4c93", ecolor="#9d4edd", capsize=3,
    )
    axis.axvline(0, color="black", linewidth=0.8, linestyle="--")
    axis.set_yticks(
        np.arange(len(selected)),
        [f"Exclude {wave}" for wave in selected["excluded_wave"]],
    )
    axis.set_xlabel("Percentage-point change per 5 additional WBmax >= 26 C days")
    axis.set_ylabel("")
    fig.tight_layout()
    fig.savefig(output, dpi=150, bbox_inches="tight")
    plt.close(fig)


def plot_dose_response(audit_tests: pd.DataFrame, output: Path) -> None:
    dose = audit_tests.loc[
        audit_tests["specification"].eq("dose_response_absolute_bins")
    ].copy()
    dose["heat_bin"] = pd.Categorical(
        dose["heat_bin"], categories=HEAT_BIN_LABELS[1:], ordered=True
    )
    dose = dose.sort_values("heat_bin")
    fig, axis = plt.subplots(figsize=(8, 4.8))
    axis.errorbar(
        np.arange(len(dose)), dose["coefficient"] * 100,
        yerr=1.96 * dose["standard_error"] * 100,
        fmt="o-", color="#d1495b", ecolor="#edae49", capsize=3,
    )
    axis.axhline(0, color="black", linewidth=0.8, linestyle="--")
    axis.set_xticks(np.arange(len(dose)), dose["heat_bin"])
    axis.set_xlabel("Days with daily WBmax >= 26 C (reference: 0 days)")
    axis.set_ylabel("Percentage-point difference in work participation")
    fig.tight_layout()
    fig.savefig(output, dpi=150, bbox_inches="tight")
    plt.close(fig)


def plot_outcome_coverage(waves: pd.DataFrame, output: Path) -> None:
    fig, axis = plt.subplots(figsize=(9, 4.8))
    axis.plot(
        waves["survey_wave"], waves["hours_observed_rate"] * 100,
        marker="o", label="Weekly hours",
    )
    axis.plot(
        waves["survey_wave"], waves["wage_observed_rate"] * 100,
        marker="s", label="Log wage",
    )
    axis.set_xlabel("Survey wave")
    axis.set_ylabel("Outcome observed (%)")
    axis.tick_params(axis="x", rotation=35)
    axis.set_ylim(0, 105)
    axis.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(output, dpi=150, bbox_inches="tight")
    plt.close(fig)


def format_workbook(path: Path) -> None:
    """Apply compact, deterministic review formatting to the audit workbook."""
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1D3557")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = "1:1"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28
        headers = {cell.column: str(cell.value or "") for cell in sheet[1]}
        for column_index in range(1, sheet.max_column + 1):
            values = [
                str(sheet.cell(row=row, column=column_index).value or "")
                for row in range(1, min(sheet.max_row, 250) + 1)
            ]
            width = min(max(max(map(len, values), default=0) + 2, 11), 38)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
            header = headers[column_index]
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row=row, column=column_index)
                if any(token in header for token in ["rate", "share"]):
                    cell.number_format = "0.0%"
                elif header.startswith(("weighted_", "mean_")):
                    cell.number_format = "0.00"
                elif header in {
                    "coefficient", "standard_error", "ci_lower_95",
                    "ci_upper_95", "p_value", "weighted_outcome_mean",
                    "mean", "std", "p05", "p25", "median", "p75", "p95",
                    "min", "max", "raw_standard_deviation_days",
                    "within_fe_standard_deviation_days", "within_fe_p05_days",
                    "within_fe_median_days", "within_fe_p95_days",
                }:
                    cell.number_format = "0.0000"
                elif any(token in header for token in ["rows", "cells", "n_obs", "n_admin"]):
                    cell.number_format = "#,##0"
                cell.alignment = Alignment(vertical="top")
    workbook.save(path)


def save_outputs(
    output: Path,
    frame: pd.DataFrame,
    cells: pd.DataFrame,
    waves: pd.DataFrame,
    exposure: pd.DataFrame,
    variation: pd.DataFrame,
    bins: pd.DataFrame,
    estimates: pd.DataFrame,
    loo: pd.DataFrame,
    audit_tests: pd.DataFrame,
) -> None:
    output.mkdir(parents=True, exist_ok=True)
    tables = {
        "sample_by_wave": waves,
        "exposure_summary": exposure,
        "within_fe_variation": variation,
        "heat_bin_summary": bins,
        "model_estimates": estimates,
        "leave_one_wave_out": loo,
        "audit_tests": audit_tests,
    }
    for name, table in tables.items():
        table.to_csv(output / f"{name}.csv", index=False)
    workbook_path = output / "credibility_audit.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name[:31], index=False)
    format_workbook(workbook_path)

    plot_exposure_distribution(cells, output / "exposure_distribution.png")
    plot_forest(estimates, output / "main_specification_forest.png")
    plot_leave_one_wave_out(loo, output / "leave_one_wave_out.png")
    plot_dose_response(audit_tests, output / "dose_response_bins.png")
    plot_outcome_coverage(waves, output / "outcome_coverage_by_wave.png")

    primary = estimates.loc[
        estimates["specification"].eq("strict_rainfall_primary")
        & estimates["term"].eq("wb26_5days")
    ].iloc[0]
    lead = estimates.loc[
        estimates["specification"].eq("temporal_current_and_lead_placebo")
        & estimates["term"].eq("lead_wb26_5days")
    ].iloc[0]
    hours_selection = estimates.loc[
        estimates["specification"].eq("selection_hours_observed")
    ].iloc[0]
    loo_min = loo["coefficient"].min()
    loo_max = loo["coefficient"].max()
    report = f"""# Heat-labor credibility audit

Status: local diagnostic evidence; not a formal AnaSOP output and not a causal result.

## Read-only contract

- Source: `mda.public.final_HEAT_LABOR_ANALYTIC` joined by SELECT to `dim_admin2_cambodia`.
- PostgreSQL session: `default_transaction_read_only=on`.
- Database writes: none.
- Person rows read: {len(frame):,}; distinct district-month cells: {len(cells):,}.

## Primary month-level association

With admin2-by-calendar-month and survey-year fixed effects, survey weights,
demographic covariates, rainfall control, and admin2-clustered inference, five
additional days with daily WBmax at least 26 C are associated with a
{primary.coefficient * 100:.3f} percentage-point change in work participation
(95% CI {primary.ci_lower_95 * 100:.3f} to {primary.ci_upper_95 * 100:.3f};
p={primary.p_value:.3f}; n={int(primary.n_obs):,}).

## Credibility diagnostics

- Leave-one-wave-out point estimates range from {loo_min * 100:.3f} to
  {loo_max * 100:.3f} percentage points.
- The future-month placebo coefficient is {lead.coefficient * 100:.3f}
  percentage points (p={lead.p_value:.3f}).
- The association between heat and weekly-hours observation is
  {hours_selection.coefficient * 100:.3f} percentage points per five hot days
  (p={hours_selection.p_value:.3f}); outcome-selection sensitivity must be
  considered when interpreting hours.
- Dose-response, alternative exposure, composition-balance, balanced-district,
  coastal-exclusion, and cluster-level results are stored in the workbook.

## Interpretation boundary

Exposure is measured using the nominal survey year and released survey month
because raw visit/interview timing is not yet represented in the database
analytic table. A separate raw-data audit identifies household survey years in
2019/2021 and exact dates in 2004/2019/2021, so this model is not a corrected-date
estimate. Occupational and household-adaptation mechanisms should remain
secondary until timing and cross-wave classifications are resolved.
"""
    (output / "README.md").write_text(report, encoding="utf-8")

    manifest = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "database": "mda",
        "schema": "public",
        "database_access": "read-only SELECT",
        "input_table": "final_HEAT_LABOR_ANALYTIC",
        "geography_table": "dim_admin2_cambodia",
        "person_rows": len(frame),
        "district_month_cells": len(cells),
        "outputs": sorted(path.name for path in output.iterdir()),
    }
    (output / "audit_manifest.json").write_text(
        json.dumps(manifest, indent=2), encoding="utf-8"
    )


def main() -> None:
    args = parse_args()
    output = args.output if args.output.is_absolute() else ROOT / args.output
    with psycopg.connect(**connection_args(args)) as connection:
        frame = load_analysis_data(connection, args.schema)
        connection.rollback()
    if frame.empty:
        raise RuntimeError("Read-only query returned no analytic rows")

    cells = exposure_cell_frame(frame)
    waves = wave_summary(frame)
    exposure = exposure_summary(cells)
    variation = within_fe_exposure_variation(cells)
    bins = heat_bin_summary(frame, cells)
    estimates, loo, audit_tests = run_models(frame)
    save_outputs(
        output, frame, cells, waves, exposure, variation, bins,
        estimates, loo, audit_tests,
    )
    print(f"read_only_input_rows={len(frame)}")
    print(f"district_month_cells={len(cells)}")
    print(f"model_terms={len(estimates) + len(loo) + len(audit_tests)}")
    print(f"output={output}")


if __name__ == "__main__":
    main()
