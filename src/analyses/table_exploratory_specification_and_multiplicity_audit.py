#!/usr/bin/env python3
"""Generate the exploratory specification and multiplicity audit table.

Plan: Preserve every target row from the frozen storyline-search universe with
nominal, family-BH, and global-BH values and an explicit evidence grade.

Framework: AnaSOP Section 5 treats researcher degrees of freedom as an
identification threat. Section 6 requires all frozen, null, mixed, and placebo
results to remain visible. Section 7 Step 9 integrates claims using the full
audit rather than favorable specifications alone. This formatting script reads
the persisted model-output audit generated from read-only PostgreSQL inputs; it
does not query or modify the database.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from table_main_cumulative_humid_heat_and_work_participation_estimates import (
    LIGHT_BORDER,
    MEDIUM_BORDER,
    MUTED,
    NAVY,
    PALE_BLUE,
    PALE_GOLD,
    PALE_TEAL,
    TEXT,
    WHITE,
    resolve_path,
    set_cell_style,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = ROOT / "data/exp/storyline-search/target_tests.csv"
DEFAULT_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_exploratory_specification_and_multiplicity_audit.xlsx"
)
DEFAULT_REVIEW_DIR = ROOT / "data/exp/figure-table-review"
REVIEW_NAMES = {
    "Timing & Average": "Table_exploratory_specification_and_multiplicity_audit_part_1_timing_and_average.png",
    "Modifiers": "Table_exploratory_specification_and_multiplicity_audit_part_2_modifiers.png",
    "Wave Stability": "Table_exploratory_specification_and_multiplicity_audit_part_3_wave_stability.png",
}

DOMAIN_DISPLAY = {
    "average_effect": "Average effect",
    "timing": "Timing and placebo",
    "outcomes": "Secondary outcome",
    "selection": "Outcome coverage",
    "heterogeneity": "Modifier",
    "cumulative_heterogeneity": "Cumulative modifier",
    "occupation": "Occupation modifier",
    "wave_stability": "Leave-one-wave-out",
}

OUTCOME_DISPLAY = {
    "worked_past_week": "Worked in the past week",
    "weekly_hours_including_zero": "Weekly hours including nonworkers",
    "weekly_hours_worked": "Weekly hours among workers",
    "log_monthly_salary_wages": "Log monthly salary or wages",
    "hours_observed": "Weekly-hours outcome observed",
    "wage_observed": "Monthly-wage outcome observed",
}

ROLE_DISPLAY = {
    "primary": "Primary current-month association",
    "control_sensitivity": "Control-set sensitivity",
    "timing_exposure": "Separate timing association",
    "future_placebo": "Future-exposure placebo",
    "distributed_lag": "Joint current-plus-lag coefficient",
    "placebo_joint": "Joint current-plus-future coefficient",
    "secondary_outcome": "Secondary-outcome association",
    "selection_diagnostic": "Outcome-observation diagnostic",
    "effect_modification": "Difference in heat slopes",
    "joint_vulnerability": "Joint female/education slope difference",
    "intersection_incremental": "Incremental three-way difference",
    "stability_diagnostic": "Leave-one-wave-out association",
}

TERM_DISPLAY = {
    "wb26_5days": "Current WB26 threshold days",
    "lag_wb26_5days": "Previous-month WB26 threshold days",
    "lead_wb26_5days": "Future-month WB26 threshold days",
    "wb26_current_lag_total_5days": "Two-month cumulative WB26 threshold days",
    "wbmean_1c": "Current mean maximum wet-bulb temperature",
    "lag_wbmean_1c": "Previous-month mean maximum wet-bulb temperature",
    "lead_wbmean_1c": "Future-month mean maximum wet-bulb temperature",
    "wbmean_current_lag_average_1c": "Two-month mean maximum wet-bulb temperature",
}

MODIFIER_DISPLAY = {
    "female": "Female",
    "low_education": "Low education",
    "female_low_education": "Female with low education",
    "rural": "Rural residence",
    "older_worker": "Age 50-64",
    "electricity_spending_positive": "Positive electricity spending",
    "high_crowding": "High household crowding",
    "female_x_low_education": "Female x low education",
    "agriculture_occupation_candidate": "Agriculture occupation candidate",
    "manual_outdoor_occupation_candidate": "Manual/outdoor occupation candidate",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--review-dir", type=Path, default=DEFAULT_REVIEW_DIR)
    return parser.parse_args()


def readable_term(term: str) -> str:
    if term in TERM_DISPLAY:
        return TERM_DISPLAY[term]
    base, separator, modifier = term.partition("_x_")
    if not separator:
        return term.replace("_", " ")
    base_display = TERM_DISPLAY.get(base, base.replace("_", " "))
    modifier_display = MODIFIER_DISPLAY.get(modifier, modifier.replace("_", " "))
    return f"{base_display} x {modifier_display}"


def effect_unit(outcome: str) -> str:
    if outcome in {"worked_past_week", "hours_observed", "wage_observed"}:
        return "pp"
    if outcome in {"weekly_hours_including_zero", "weekly_hours_worked"}:
        return "hours"
    if outcome == "log_monthly_salary_wages":
        return "log points"
    raise ValueError(f"Unknown outcome: {outcome}")


def evidence_grade(row: pd.Series) -> str:
    if row["analysis_domain"] == "wave_stability":
        return "Stability diagnostic; outside BH"
    if float(row["global_bh_q_value"]) < 0.10:
        return "Global FDR support"
    if float(row["family_bh_q_value"]) < 0.10:
        return "Family-supported; exploratory"
    if float(row["p_value"]) < 0.05:
        return "Nominal only"
    return "No adjusted support"


def prepare_audit(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path)
    if len(frame) != 82 or not frame["is_target_test"].eq(True).all():
        raise AssertionError("Expected exactly 82 frozen target-test rows")
    frame = frame.copy()
    frame.insert(0, "audit_id", [f"T{number:03d}" for number in range(1, len(frame) + 1)])
    frame["analysis_area"] = frame["analysis_domain"].map(DOMAIN_DISPLAY)
    frame["model_family_display"] = frame["test_family"].str.replace("_", " ").str.title()
    frame["outcome_display"] = frame["outcome"].map(OUTCOME_DISPLAY)
    frame["target_display"] = frame["term"].astype(str).map(readable_term)
    frame["estimand_display"] = frame["role"].map(ROLE_DISPLAY)
    frame["unit"] = frame["outcome"].map(effect_unit)
    frame["effect_text"] = [
        f"{float(value):.3f} {unit}"
        for value, unit in zip(frame["effect_display"], frame["unit"], strict=True)
    ]
    frame["ci_text"] = [
        f"[{float(lower):.3f}, {float(upper):.3f}]"
        for lower, upper in zip(
            frame["ci_lower_display"], frame["ci_upper_display"], strict=True
        )
    ]
    frame["evidence_grade"] = frame.apply(evidence_grade, axis=1)
    frame["review_section"] = np.select(
        [
            frame["analysis_domain"].isin(["average_effect", "timing"]),
            frame["analysis_domain"].isin(["outcomes", "selection"]),
            frame["analysis_domain"].isin(
                ["heterogeneity", "cumulative_heterogeneity", "occupation"]
            ),
            frame["analysis_domain"].eq("wave_stability"),
        ],
        ["Timing & Average", "Audit Only", "Modifiers", "Wave Stability"],
        default="Unassigned",
    )
    if frame["review_section"].eq("Unassigned").any():
        raise AssertionError("At least one audit row is not assigned to a review section")
    return frame


def format_q(value: float) -> str:
    return "--" if pd.isna(value) else f"{value:.3f}"


def format_p(value: float) -> str:
    """Return the strongest manuscript-approved p-value threshold."""
    if pd.isna(value) or value >= 0.10:
        return "—"
    if value < 0.01:
        return "p < 0.01"
    if value < 0.05:
        return "p < 0.05"
    return "p < 0.10"


def add_human_sheet(
    workbook: Workbook,
    sheet_name: str,
    title_suffix: str,
    frame: pd.DataFrame,
) -> None:
    sheet = workbook.create_sheet(sheet_name)
    sheet.sheet_view.showGridLines = False
    headers = [
        "Analysis Area",
        "Model Family",
        "Outcome",
        "Exposure / Target",
        "Estimand",
        "Estimate",
        "95% CI",
        "p threshold",
        "Family BH q",
        "Global BH q",
        "Evidence Grade",
    ]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = sheet.cell(
        1,
        1,
        f"Exploratory Specification and Multiplicity Audit — {title_suffix}",
    )
    set_cell_style(title, fill=NAVY, color=WHITE, bold=True, size=14, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=8, wrap=True)
    sheet.row_dimensions[2].height = 42

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    previous_area = None
    for offset, row in enumerate(frame.itertuples(index=False), start=0):
        row_number = 3 + offset
        values = [
            row.analysis_area,
            row.model_family_display,
            row.outcome_display,
            row.target_display,
            row.estimand_display,
            row.effect_text,
            row.ci_text,
            format_p(float(row.p_value)),
            format_q(float(row.family_bh_q_value)),
            format_q(float(row.global_bh_q_value)),
            row.evidence_grade,
        ]
        starts_area = previous_area is not None and row.analysis_area != previous_area
        family_supported = row.evidence_grade == "Family-supported; exploratory"
        fill = PALE_GOLD if family_supported else (PALE_BLUE if offset % 2 else WHITE)
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=fill,
                bold=family_supported and column in [4, 11],
                size=7.5,
                align="left" if column in [1, 2, 3, 4, 5, 6, 7, 11] else "right",
                wrap=column in [1, 2, 3, 4, 5, 6, 7, 11],
            )
            cell.border = Border(top=group_top if starts_area else Side(), bottom=thin)
        sheet.row_dimensions[row_number].height = 28
        previous_area = row.analysis_area

    last_data_row = 2 + len(frame)
    notes = [
        "Notes: The workbook archives 82 audit rows. Global and family BH adjustments apply to 64 substantive/exploratory tests; 18 leave-one-wave-out rows are stability diagnostics outside the BH calculation.",
        "Family-supported rows remain exploratory because no eligible target survives global 10% false-discovery control. Estimate units follow the outcome: percentage points, hours, or log points.",
        "The source audit was generated by the persistent storyline-search script using read-only PostgreSQL inputs. This formatting script performs no database access or writes.",
    ]
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=len(headers)
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=7.5, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 24

    widths = [20, 27, 28, 38, 31, 16, 19, 11, 13, 13, 29]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:K{last_data_row}"
    sheet.print_area = f"A1:K{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.15
    sheet.page_margins.right = 0.15
    sheet.page_margins.top = 0.18
    sheet.page_margins.bottom = 0.18


def add_wave_stability_sheet(workbook: Workbook, frame: pd.DataFrame) -> None:
    """Add a compact paired display with one row per omitted survey wave."""
    if len(frame) != 18:
        raise AssertionError("Expected 18 leave-one-wave-out records")
    sheet = workbook.create_sheet("Wave Stability")
    sheet.sheet_view.showGridLines = False
    headers = [
        "Omitted Wave",
        "Observations",
        "WB26 Estimate (pp)",
        "WB26 95% CI",
        "WB26 p threshold",
        "Mean WBmax Estimate (pp)",
        "Mean WBmax 95% CI",
        "Mean WBmax p threshold",
    ]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = sheet.cell(
        1,
        1,
        "Exploratory Specification and Multiplicity Audit — Leave-One-Wave-Out Stability",
    )
    set_cell_style(title, fill=NAVY, color=WHITE, bold=True, size=14, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=8, wrap=True)
    sheet.row_dimensions[2].height = 42

    wb26 = frame.loc[frame["exposure_family"].eq("wb26")].copy()
    wbmean = frame.loc[frame["exposure_family"].eq("wbmean")].copy()
    wave_order = wb26["excluded_wave"].astype(str).tolist()
    if len(wave_order) != 9 or set(wave_order) != set(wbmean["excluded_wave"].astype(str)):
        raise AssertionError("Leave-one-wave-out exposure definitions do not align by wave")

    thin = Side(style="thin", color=LIGHT_BORDER)
    for offset, wave in enumerate(wave_order):
        threshold = wb26.loc[wb26["excluded_wave"].astype(str).eq(wave)].iloc[0]
        continuous = wbmean.loc[wbmean["excluded_wave"].astype(str).eq(wave)].iloc[0]
        if int(threshold["n_obs"]) != int(continuous["n_obs"]):
            raise AssertionError(f"Sample sizes differ across exposure definitions for {wave}")
        values = [
            wave,
            int(threshold["n_obs"]),
            float(threshold["effect_display"]),
            threshold["ci_text"],
            format_p(float(threshold["p_value"])),
            float(continuous["effect_display"]),
            continuous["ci_text"],
            format_p(float(continuous["p_value"])),
        ]
        row_number = 3 + offset
        fill = PALE_BLUE if offset % 2 else WHITE
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=fill,
                bold=column == 1,
                size=8,
                align="left" if column in [1, 4, 7] else "right",
                wrap=column in [1, 4, 7],
            )
            cell.border = Border(bottom=thin)
        sheet.cell(row_number, 2).number_format = "#,##0"
        for column in [3, 6]:
            sheet.cell(row_number, column).number_format = "0.000"
        sheet.row_dimensions[row_number].height = 27

    last_data_row = 2 + len(wave_order)
    notes = [
        "Notes: Each row excludes one survey wave and pairs the two current-month humid-heat definitions on the same retained sample.",
        "Leave-one-wave-out estimates are stability diagnostics outside the BH calculation; interpretation focuses on sign and range rather than individual significance.",
        "All 18 source records remain in Full Audit. The formatting script performs no database access or writes.",
    ]
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=len(headers)
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=7.5, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 24

    widths = [18, 17, 21, 22, 15, 28, 22, 19]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:H{last_data_row}"
    sheet.print_area = f"A1:H{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.2
    sheet.page_margins.right = 0.2
    sheet.page_margins.top = 0.2
    sheet.page_margins.bottom = 0.2


def build_workbook(frame: pd.DataFrame, output: Path) -> list[str]:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    standard_sections = [
        ("Timing & Average", "Timing and Average Effects"),
        ("Modifiers", "Modifiers and Occupation"),
    ]
    for sheet_name, title_suffix in standard_sections:
        subset = frame.loc[frame["review_section"].eq(sheet_name)].copy()
        add_human_sheet(workbook, sheet_name, title_suffix, subset)
    add_wave_stability_sheet(
        workbook,
        frame.loc[frame["review_section"].eq("Wave Stability")].copy(),
    )

    audit = workbook.create_sheet("Full Audit")
    audit.sheet_view.showGridLines = False
    for column, header in enumerate(frame.columns, start=1):
        display_header = "p_threshold" if header == "p_value" else header
        cell = audit.cell(1, column, display_header)
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=8, wrap=True)
    for row_number, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            if frame.columns[column - 1] == "p_value":
                value = format_p(float(value))
            if pd.isna(value):
                value = None
            cell = audit.cell(row_number, column, value)
            set_cell_style(
                cell,
                size=7,
                align="left" if isinstance(value, str) else "right",
                wrap=isinstance(value, str),
            )
        audit.row_dimensions[row_number].height = 22
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit.print_area = audit.dimensions
    audit.page_setup.orientation = "landscape"
    audit.page_setup.paperSize = audit.PAPERSIZE_A3
    audit.page_setup.fitToWidth = 1
    audit.page_setup.fitToHeight = 0
    audit.sheet_properties.pageSetUpPr.fitToPage = True
    for column in range(1, audit.max_column + 1):
        audit.column_dimensions[get_column_letter(column)].width = 17

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    eligible = frame["global_bh_q_value"].notna()
    metadata_rows = [
        ("Field", "Value"),
        ("Table", "Exploratory Specification and Multiplicity Audit"),
        ("Archived target rows", f"{len(frame)}"),
        ("BH-eligible rows", f"{int(eligible.sum())}"),
        ("Leave-one-wave-out diagnostics", f"{int((~eligible).sum())}"),
        ("Nominal p < 0.05", f"{int(frame['p_value'].lt(0.05).sum())}"),
        ("Family BH q < 0.10", f"{int(frame['family_bh_q_value'].lt(0.10).sum())}"),
        ("Global BH q < 0.10", f"{int(frame['global_bh_q_value'].lt(0.10).sum())}"),
        ("Minimum global BH q", f"{float(frame['global_bh_q_value'].min()):.6f}"),
        ("Source", "data/exp/storyline-search/target_tests.csv"),
        ("Source generation", "Persistent read-only database analysis script"),
        ("Database writes by this script", "None"),
        ("Generated", datetime.now().astimezone().isoformat(timespec="seconds")),
    ]
    for row_number, values in enumerate(metadata_rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = metadata.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=NAVY if row_number == 1 else None,
                color=WHITE if row_number == 1 else TEXT,
                bold=row_number == 1 or column == 1,
                align="left",
                wrap=True,
            )
    metadata.column_dimensions["A"].width = 31
    metadata.column_dimensions["B"].width = 78
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = metadata.dimensions
    metadata.print_area = metadata.dimensions
    workbook.save(output)
    return [item[0] for item in standard_sections] + ["Wave Stability"]


def render_review_pages(
    workbook_path: Path,
    review_dir: Path,
    sheet_names: list[str],
) -> dict[str, Path]:
    soffice = Path("/opt/homebrew/bin/soffice")
    if not soffice.exists():
        located = shutil.which("soffice")
        if not located:
            raise RuntimeError("LibreOffice/soffice is required for workbook rendering")
        soffice = Path(located)
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("pdftoppm is required for PNG rendering")
    review_dir.mkdir(parents=True, exist_ok=True)
    outputs: dict[str, Path] = {}
    for sheet_name in sheet_names:
        with tempfile.TemporaryDirectory(prefix="mj02b-audit-review-") as temporary:
            temporary_path = Path(temporary)
            single = load_workbook(workbook_path, read_only=False, data_only=False)
            for name in list(single.sheetnames):
                if name != sheet_name:
                    single.remove(single[name])
            temporary_xlsx = temporary_path / "review.xlsx"
            single.save(temporary_xlsx)
            single.close()
            profile_uri = (temporary_path / "lo-profile").as_uri()
            subprocess.run(
                [
                    str(soffice),
                    f"-env:UserInstallation={profile_uri}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(temporary_path),
                    str(temporary_xlsx),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
            pdf_path = temporary_path / "review.pdf"
            output = review_dir / REVIEW_NAMES[sheet_name]
            prefix = output.with_suffix("")
            subprocess.run(
                [pdftoppm, "-png", "-r", "180", "-singlefile", str(pdf_path), str(prefix)],
                check=True,
                capture_output=True,
                text=True,
            )
            if not output.exists() or output.stat().st_size == 0:
                raise RuntimeError(f"Review render is missing: {output}")
            outputs[sheet_name] = output
    return outputs


def validate_outputs(
    frame: pd.DataFrame,
    workbook_path: Path,
    review_outputs: dict[str, Path],
) -> None:
    if len(frame) != 82:
        raise AssertionError("Audit row count changed")
    if int(frame["global_bh_q_value"].notna().sum()) != 64:
        raise AssertionError("Expected 64 BH-eligible rows")
    if int(frame["analysis_domain"].eq("wave_stability").sum()) != 18:
        raise AssertionError("Expected 18 wave-stability diagnostics")
    if int(frame["global_bh_q_value"].lt(0.10).sum()) != 0:
        raise AssertionError("A target unexpectedly survives global 10% FDR")
    if not np.isclose(float(frame["global_bh_q_value"].min()), 0.154076825913097):
        raise AssertionError("Unexpected minimum global BH q-value")
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    expected_sheets = [
        "Timing & Average",
        "Modifiers",
        "Wave Stability",
        "Full Audit",
        "Metadata",
    ]
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    for sheet_name in expected_sheets[:3]:
        sheet = workbook[sheet_name]
        if not str(sheet["A1"].value).startswith("Exploratory Specification and Multiplicity Audit"):
            raise AssertionError(f"Missing title on {sheet_name}")
        if sheet["A2"].value != "Analysis Area" or sheet.freeze_panes != "A3":
            if sheet_name != "Wave Stability" or sheet["A2"].value != "Omitted Wave":
                raise AssertionError(f"Title must be followed by headers on {sheet_name}")
    wave_sheet = workbook["Wave Stability"]
    if wave_sheet.max_column != 8 or wave_sheet.max_row != 15:
        raise AssertionError("Unexpected compact wave-stability display dimensions")
    if any(sheet.print_area is None for sheet in workbook.worksheets):
        raise AssertionError("Every worksheet must have a print area")
    workbook.close()
    if set(review_outputs) != set(expected_sheets[:3]):
        raise AssertionError("Review-page set is incomplete")


def main() -> None:
    args = parse_args()
    input_path = resolve_path(args.input)
    output = resolve_path(args.output)
    review_dir = resolve_path(args.review_dir)
    frame = prepare_audit(input_path)
    sheet_names = build_workbook(frame, output)
    review_outputs = render_review_pages(output, review_dir, sheet_names)
    validate_outputs(frame, output, review_outputs)
    print(f"table={output}")
    for sheet_name, path in review_outputs.items():
        print(f"review[{sheet_name}]={path}")
    print(
        f"rows={len(frame)} eligible={frame['global_bh_q_value'].notna().sum()} "
        f"diagnostics={frame['global_bh_q_value'].isna().sum()} "
        f"nominal={frame['p_value'].lt(0.05).sum()} "
        f"family_q_lt_10={frame['family_bh_q_value'].lt(0.10).sum()} "
        f"global_q_lt_10={frame['global_bh_q_value'].lt(0.10).sum()}"
    )


if __name__ == "__main__":
    main()
