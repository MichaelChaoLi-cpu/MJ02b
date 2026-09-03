#!/usr/bin/env python3
"""Generate boundary outcomes and unsupported modifier tests.

Plan: Preserve cumulative hours and wage results, outcome-observation checks,
and non-education modifier tests that do not support the central story.

Framework: AnaSOP Section 6 treats hours and wages as secondary outcomes and
sex, age, rural residence, occupation, crowding, and electricity spending as
boundary modifiers. Section 7 Step 8 requires these null, mixed, or incomplete
patterns to remain visible. Models are re-estimated from PostgreSQL under a
forced read-only session; the persisted 82-row audit supplies only the frozen
family and global BH values and is checked against the new estimates.
"""

from __future__ import annotations

import argparse
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from run_heat_labor_credibility_audit import fit_hdfe
from run_heat_labor_extensions import BASE_CONTROLS, connection_args, load_data
from search_heat_labor_storylines import (
    EXPOSURES,
    build_cell_frame,
    calculate_term_support,
    prepare_frame,
    result_scale,
)
from table_main_cumulative_humid_heat_and_work_participation_estimates import (
    LIGHT_BORDER,
    MEDIUM_BORDER,
    MUTED,
    NAVY,
    PALE_BLUE,
    PALE_GOLD,
    PALE_TEAL,
    TEXT,
    WHITE,
    render_review_png,
    resolve_path,
    set_cell_style,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_boundary_outcomes_and_unsupported_modifiers.xlsx"
)
DEFAULT_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_boundary_outcomes_and_unsupported_modifiers.png"
)
DEFAULT_MODIFIER_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_boundary_outcomes_and_unsupported_modifiers_modifiers.png"
)
DEFAULT_MULTIPLICITY_AUDIT = ROOT / "data/exp/storyline-search/target_tests.csv"

OUTCOME_DISPLAY = {
    "weekly_hours_including_zero": ("Weekly hours including nonworkers", "Hours"),
    "weekly_hours_worked": ("Weekly hours among current workers", "Hours"),
    "log_monthly_salary_wages": ("Log monthly salary or wages", "Log points"),
    "hours_observed": ("Weekly-hours outcome observed", "Percentage points"),
    "wage_observed": ("Monthly-wage outcome observed", "Percentage points"),
    "worked_past_week": ("Worked in the past week", "Percentage points"),
}

MODIFIER_DISPLAY = {
    "female": "Female minus male heat slope",
    "older_worker": "Age 50-64 minus age 15-49 heat slope",
    "rural": "Rural minus urban heat slope",
    "electricity_spending_positive": "Positive versus no electricity-spending heat slope",
    "high_crowding": "High- versus lower-crowding heat slope",
    "agriculture_occupation_candidate": "Agriculture-candidate versus other-worker heat slope",
    "manual_outdoor_occupation_candidate": "Manual/outdoor-candidate versus other-worker heat slope",
}

CONSTRUCT_JUDGMENT = {
    "inclusive_hours": "Adverse direction under both definitions; imprecise",
    "worker_hours": "Adverse direction under both definitions; imprecise",
    "log_wage": "Opposing directions across exposure definitions",
    "hours_observed": "No consistent outcome-observation evidence",
    "wage_observed": "No consistent outcome-observation evidence",
    "female": "No replicated multiplicity-adjusted support",
    "older_worker": "No replicated multiplicity-adjusted support",
    "rural": "No replicated multiplicity-adjusted support",
    "electricity_spending_positive": "No electricity-buffering pattern",
    "high_crowding": "No replicated multiplicity-adjusted support",
    "agriculture_occupation_candidate": "Suggestive only; not globally robust",
    "manual_outdoor_occupation_candidate": "No adverse occupation interaction pattern",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--review-output", type=Path, default=DEFAULT_REVIEW_OUTPUT)
    parser.add_argument(
        "--modifier-review-output", type=Path, default=DEFAULT_MODIFIER_REVIEW_OUTPUT
    )
    parser.add_argument(
        "--multiplicity-audit", type=Path, default=DEFAULT_MULTIPLICITY_AUDIT
    )
    return parser.parse_args()


def model_plan() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    outcomes = [
        ("weekly_hours_including_zero", "inclusive_hours"),
        ("weekly_hours_worked", "worker_hours"),
        ("log_monthly_salary_wages", "log_wage"),
    ]
    for family, exposure in EXPOSURES.items():
        term = exposure["cumulative"]
        for outcome, construct in outcomes:
            rows.append(
                {
                    "block": "Secondary labor outcome",
                    "construct": construct,
                    "model_id": f"outcome_{construct}_{family}_cumulative",
                    "exposure_family": family,
                    "window": "Current + previous",
                    "outcome": outcome,
                    "terms": [term],
                    "target_term": term,
                    "controls": BASE_CONTROLS,
                    "mask_type": "outcome",
                    "estimand": "Cumulative heat association",
                }
            )
        current = exposure["current"]
        for outcome in ["hours_observed", "wage_observed"]:
            rows.append(
                {
                    "block": "Outcome coverage",
                    "construct": outcome,
                    "model_id": f"selection_{outcome}_{family}",
                    "exposure_family": family,
                    "window": "Current month",
                    "outcome": outcome,
                    "terms": [current],
                    "target_term": current,
                    "controls": BASE_CONTROLS,
                    "mask_type": "selection",
                    "estimand": "Heat association with outcome observation",
                }
            )

    modifiers = [
        "female",
        "older_worker",
        "rural",
        "electricity_spending_positive",
        "high_crowding",
    ]
    for modifier in modifiers:
        for family, exposure in EXPOSURES.items():
            current = exposure["current"]
            interaction = f"{current}_x_{modifier}"
            rows.append(
                {
                    "block": "Unsupported modifier",
                    "construct": modifier,
                    "model_id": f"heterogeneity_{family}_{modifier}",
                    "exposure_family": family,
                    "window": "Current month",
                    "outcome": "worked_past_week",
                    "terms": [current, modifier, interaction],
                    "target_term": interaction,
                    "controls": [item for item in BASE_CONTROLS if item != modifier],
                    "mask_type": "modifier",
                    "modifier": modifier,
                    "base_term": current,
                    "estimand": MODIFIER_DISPLAY[modifier],
                }
            )

    occupation_modifiers = [
        "agriculture_occupation_candidate",
        "manual_outdoor_occupation_candidate",
    ]
    for modifier in occupation_modifiers:
        for family, exposure in EXPOSURES.items():
            current = exposure["current"]
            interaction = f"{current}_x_{modifier}"
            rows.append(
                {
                    "block": "Occupation modifier",
                    "construct": modifier,
                    "model_id": f"occupation_{family}_{modifier}",
                    "exposure_family": family,
                    "window": "Current month",
                    "outcome": "weekly_hours_worked",
                    "terms": [current, modifier, interaction],
                    "target_term": interaction,
                    "controls": BASE_CONTROLS,
                    "mask_type": "occupation",
                    "modifier": modifier,
                    "base_term": current,
                    "estimand": MODIFIER_DISPLAY[modifier],
                }
            )
    return rows


def fit_plan_row(
    frame: pd.DataFrame,
    plan: dict[str, object],
    support: pd.Series,
) -> dict[str, object]:
    working = frame.copy()
    target_term = str(plan["target_term"])
    terms = list(plan["terms"])
    mask_type = str(plan["mask_type"])
    if mask_type in {"modifier", "occupation"}:
        modifier = str(plan["modifier"])
        base_term = str(plan["base_term"])
        working[target_term] = working[base_term] * working[modifier]
    if mask_type == "outcome":
        mask = working[str(plan["outcome"])].notna() & working[target_term].notna()
        standardization_term = target_term
    elif mask_type == "selection":
        mask = working[target_term].notna()
        standardization_term = target_term
    elif mask_type == "modifier":
        mask = (
            working["worked_past_week"].notna()
            & working[str(plan["modifier"])].notna()
            & working[str(plan["base_term"])].notna()
        )
        standardization_term = str(plan["base_term"])
    elif mask_type == "occupation":
        mask = (
            working["worked_past_week"].eq(1)
            & working["weekly_hours_worked"].notna()
            & working[str(plan["modifier"])].notna()
            & working[str(plan["base_term"])].notna()
        )
        standardization_term = str(plan["base_term"])
    else:
        raise ValueError(f"Unknown mask type: {mask_type}")

    records, model = fit_hdfe(
        working,
        str(plan["model_id"]),
        str(plan["outcome"]),
        terms,
        list(plan["controls"]),
        sample_label=str(plan["block"]),
        mask=mask,
    )
    record = next(row for row in records if row["term"] == target_term)
    scale = result_scale(str(plan["outcome"]))
    within_sd = float(support[standardization_term])
    return {
        **plan,
        "coefficient": float(record["coefficient"]),
        "standard_error": float(record["standard_error"]),
        "ci_lower_95": float(record["ci_lower_95"]),
        "ci_upper_95": float(record["ci_upper_95"]),
        "p_value": float(record["p_value"]),
        "effect_display": scale * float(record["coefficient"]),
        "ci_lower_display": scale * float(record["ci_lower_95"]),
        "ci_upper_display": scale * float(record["ci_upper_95"]),
        "within_fe_sd": within_sd,
        "standardized_effect_display": scale * float(record["coefficient"]) * within_sd,
        "n_obs": int(model.nobs),
        "n_admin2": int(record["n_admin2"]),
        "n_provinces": int(record["n_provinces"]),
        "fixed_effects": str(record["fixed_effects"]),
        "cluster_level": str(record["cluster_level"]),
        "controls_display": str(record["controls"]),
    }


def attach_multiplicity(
    estimates: pd.DataFrame,
    audit_path: Path,
) -> pd.DataFrame:
    audit = pd.read_csv(audit_path)
    audit = audit.loc[audit["is_target_test"].eq(True)].copy()
    frozen = audit.set_index(["model_id", "term"])
    family_q: list[float] = []
    global_q: list[float] = []
    for row in estimates.itertuples(index=False):
        key = (row.model_id, row.target_term)
        if key not in frozen.index:
            raise AssertionError(f"Model is absent from the frozen audit: {key}")
        match = frozen.loc[key]
        if isinstance(match, pd.DataFrame):
            raise AssertionError(f"Frozen audit key is duplicated: {key}")
        if not np.isclose(float(row.effect_display), float(match["effect_display"]), atol=2e-6):
            raise AssertionError(f"Frozen effect mismatch: {key}")
        if not np.isclose(float(row.p_value), float(match["p_value"]), atol=2e-6):
            raise AssertionError(f"Frozen p-value mismatch: {key}")
        family_q.append(float(match["family_bh_q_value"]))
        global_q.append(float(match["global_bh_q_value"]))
    result = estimates.copy()
    result["family_bh_q_value"] = family_q
    result["global_bh_q_value"] = global_q
    result["outcome_display"] = result["outcome"].map(
        {key: value[0] for key, value in OUTCOME_DISPLAY.items()}
    )
    result["effect_unit"] = result["outcome"].map(
        {key: value[1] for key, value in OUTCOME_DISPLAY.items()}
    )
    result["exposure_display"] = result["exposure_family"].map(
        {"wb26": "WB26 threshold days", "wbmean": "Mean maximum wet-bulb temperature"}
    )
    result["evidence_judgment"] = result["construct"].map(CONSTRUCT_JUDGMENT)
    return result


def estimate_all(frame: pd.DataFrame, audit_path: Path) -> pd.DataFrame:
    cells = build_cell_frame(frame)
    support = calculate_term_support(cells).set_index("term")["residual_sd"]
    rows = [fit_plan_row(frame, plan, support) for plan in model_plan()]
    estimates = pd.DataFrame(rows)
    return attach_multiplicity(estimates, audit_path)


def format_ci(lower: float, upper: float) -> str:
    return f"[{lower:.3f}, {upper:.3f}]"


def format_effect_ci(row: pd.Series) -> str:
    return (
        f"{float(row['effect_display']):.3f}\n"
        f"{format_ci(float(row['ci_lower_display']), float(row['ci_upper_display']))}"
    )


def build_compact_tables(
    estimates: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    secondary_order = [
        "inclusive_hours",
        "worker_hours",
        "log_wage",
        "hours_observed",
        "wage_observed",
    ]
    modifier_order = [
        "female",
        "older_worker",
        "rural",
        "electricity_spending_positive",
        "high_crowding",
        "agriculture_occupation_candidate",
        "manual_outdoor_occupation_candidate",
    ]
    modifier_labels = {
        "female": "Female vs male",
        "older_worker": "Age 50-64 vs 15-49",
        "rural": "Rural vs urban",
        "electricity_spending_positive": "Positive electricity spending",
        "high_crowding": "High household crowding",
        "agriculture_occupation_candidate": "Agriculture candidate",
        "manual_outdoor_occupation_candidate": "Manual/outdoor candidate",
    }

    def pair(construct: str) -> tuple[pd.Series, pd.Series]:
        subset = estimates.loc[estimates["construct"].eq(construct)].set_index(
            "exposure_family"
        )
        if set(subset.index) != {"wb26", "wbmean"} or len(subset) != 2:
            raise AssertionError(f"Expected one WB26 and one WBmean row for {construct}")
        return subset.loc["wb26"], subset.loc["wbmean"]

    secondary_rows: list[dict[str, object]] = []
    for construct in secondary_order:
        wb26, wbmean = pair(construct)
        secondary_rows.append(
            {
                "Outcome": wb26["outcome_display"],
                "Effect Unit": wb26["effect_unit"],
                "WB26 Estimate [95% CI]": format_effect_ci(wb26),
                "WB26 Global q": float(wb26["global_bh_q_value"]),
                "Mean WBmax Estimate [95% CI]": format_effect_ci(wbmean),
                "Mean WBmax Global q": float(wbmean["global_bh_q_value"]),
                "Boundary": wb26["evidence_judgment"],
            }
        )
    modifier_rows: list[dict[str, object]] = []
    for construct in modifier_order:
        wb26, wbmean = pair(construct)
        modifier_rows.append(
            {
                "Modifier": modifier_labels[construct],
                "Outcome": wb26["outcome_display"],
                "WB26 Estimate [95% CI]": format_effect_ci(wb26),
                "WB26 Global q": float(wb26["global_bh_q_value"]),
                "Mean WBmax Estimate [95% CI]": format_effect_ci(wbmean),
                "Mean WBmax Global q": float(wbmean["global_bh_q_value"]),
                "Boundary": wb26["evidence_judgment"],
            }
        )
    secondary = pd.DataFrame(secondary_rows)
    modifiers = pd.DataFrame(modifier_rows)
    if secondary.shape != (5, 7) or modifiers.shape != (7, 7):
        raise AssertionError("Unexpected compact boundary-table dimensions")
    return secondary, modifiers


def write_compact_sheet(
    sheet: object,
    table: pd.DataFrame,
    *,
    sheet_title: str,
    title: str,
    notes: list[str],
) -> None:
    sheet.title = sheet_title
    sheet.sheet_view.showGridLines = False
    columns = list(table.columns)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title_cell = sheet.cell(1, 1, title)
    set_cell_style(title_cell, color=NAVY, bold=True, size=14, align="left")
    sheet.row_dimensions[1].height = 31
    medium = Side(style="medium", color=MEDIUM_BORDER)
    thin = Side(style="thin", color=LIGHT_BORDER)
    for column, header in enumerate(columns, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, bold=True, size=9, wrap=True)
        cell.border = Border(top=medium, bottom=medium)
    sheet.row_dimensions[2].height = 38

    for offset, row in enumerate(table.itertuples(index=False, name=None)):
        row_number = 3 + offset
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, column, value)
            bold_boundary = str(row[-1]).startswith("Suggestive only") and column == 7
            set_cell_style(
                cell,
                bold=column == 1 or bold_boundary,
                size=9,
                align="left" if column in {1, 2, 7} else "center",
                wrap=True,
            )
            cell.border = Border(
                bottom=medium if offset == len(table) - 1 else thin
            )
        for column in [4, 6]:
            sheet.cell(row_number, column).number_format = "0.000"
        sheet.row_dimensions[row_number].height = 42

    notes_start = 4 + len(table)
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=len(columns)
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 24
    widths = [31, 20, 27, 14, 29, 16, 38]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.print_area = f"A1:G{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25


def build_workbook(
    estimates: pd.DataFrame,
    secondary: pd.DataFrame,
    modifiers: pd.DataFrame,
    output: Path,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    write_compact_sheet(
        workbook.active,
        secondary,
        sheet_title="Secondary Outcomes",
        title="Boundary Outcomes and Unsupported Modifiers — Secondary Outcomes",
        notes=[
            "Notes: Effects are hours, log points, or percentage points as labeled. Each cell reports the estimate and 95% confidence interval.",
            "Global BH q-values come from the 64-test eligible universe. No secondary outcome survives global 10% false-discovery control.",
            "Results delimit the supported narrative; imprecision or inconsistency does not prove absence of an effect. Full estimates and model fields remain in Estimate Audit.",
        ],
    )
    modifier_sheet = workbook.create_sheet("Modifier Boundaries")
    write_compact_sheet(
        modifier_sheet,
        modifiers,
        sheet_title="Modifier Boundaries",
        title="Boundary Outcomes and Unsupported Modifiers — Modifier Boundaries",
        notes=[
            "Notes: Values are heat-slope differences in percentage points, except occupation interactions, which use weekly hours among current workers.",
            "The agriculture signal is suggestive at the family level but does not survive global 10% false-discovery control; other modifier patterns are null, mixed, or incomplete.",
            "Education remains the only vulnerability contrast carried into the allocation layer. Complete model fields and family q-values remain in Estimate Audit.",
        ],
    )

    audit = workbook.create_sheet("Estimate Audit")
    audit.sheet_view.showGridLines = False
    for column, header in enumerate(estimates.columns, start=1):
        cell = audit.cell(1, column, header)
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=8, wrap=True)
    for row_number, row in enumerate(estimates.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            if isinstance(value, (list, tuple)):
                value = "+".join(str(item) for item in value)
            if pd.isna(value):
                value = None
            cell = audit.cell(row_number, column, value)
            set_cell_style(
                cell,
                size=7.5,
                align="left" if isinstance(value, str) else "right",
                wrap=isinstance(value, str),
            )
        audit.row_dimensions[row_number].height = 24
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit.print_area = audit.dimensions
    audit.page_setup.orientation = "landscape"
    audit.page_setup.paperSize = audit.PAPERSIZE_A3
    audit.page_setup.fitToWidth = 1
    audit.page_setup.fitToHeight = 0
    audit.sheet_properties.pageSetUpPr.fitToPage = True
    for column in range(1, audit.max_column + 1):
        audit.column_dimensions[get_column_letter(column)].width = 18

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    block_counts = estimates["block"].value_counts()
    metadata_rows = [
        ("Field", "Value"),
        ("Table", "Boundary Outcomes and Unsupported Modifiers"),
        ("Rows", f"{len(estimates)} frozen audit rows used in this boundary table"),
        ("Secondary labor outcomes", f"{int(block_counts['Secondary labor outcome'])}"),
        ("Outcome coverage diagnostics", f"{int(block_counts['Outcome coverage'])}"),
        ("Unsupported modifiers", f"{int(block_counts['Unsupported modifier'])}"),
        ("Occupation modifiers", f"{int(block_counts['Occupation modifier'])}"),
        ("Fixed effects", "District Calendar Month and Survey Year"),
        ("Weights", "CSES analysis weights"),
        ("Inference", "District-clustered, debiased standard errors"),
        ("Multiplicity universe", "82-row audit: 64 BH-eligible tests + 18 leave-one-wave-out diagnostics"),
        ("Database session", "PostgreSQL default_transaction_read_only=on"),
        ("Database writes", "None"),
    ]
    for row_number, values in enumerate(metadata_rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = metadata.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=NAVY if row_number == 1 else None,
                color=WHITE if row_number == 1 else TEXT,
                bold=row_number == 1 or column == 1,
                align="left",
                wrap=True,
            )
    metadata.column_dimensions["A"].width = 30
    metadata.column_dimensions["B"].width = 78
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = metadata.dimensions
    metadata.print_area = metadata.dimensions
    workbook.save(output)


def render_sheet_review(output: Path, sheet_name: str, review: Path) -> None:
    with tempfile.TemporaryDirectory(
        prefix="mj02b-boundary-table-review-", dir="/private/tmp"
    ) as temporary:
        temporary_path = Path(temporary) / f"{sheet_name.replace(' ', '_')}.xlsx"
        workbook = load_workbook(output, data_only=False, read_only=False)
        for sheet in list(workbook.worksheets):
            if sheet.title != sheet_name:
                workbook.remove(sheet)
        workbook.active = 0
        workbook.save(temporary_path)
        workbook.close()
        render_review_png(temporary_path, review)


def validate_outputs(
    estimates: pd.DataFrame,
    secondary: pd.DataFrame,
    modifiers: pd.DataFrame,
    workbook_path: Path,
    png_path: Path,
    modifier_png_path: Path,
) -> None:
    if len(estimates) != 24:
        raise AssertionError(f"Expected 24 frozen boundary tests, observed {len(estimates)}")
    expected_blocks = {
        "Secondary labor outcome": 6,
        "Outcome coverage": 4,
        "Unsupported modifier": 10,
        "Occupation modifier": 4,
    }
    if estimates["block"].value_counts().to_dict() != expected_blocks:
        raise AssertionError("Unexpected boundary-test block counts")
    if estimates[["family_bh_q_value", "global_bh_q_value"]].isna().any().any():
        raise AssertionError("Multiplicity values are incomplete")
    if float(estimates["global_bh_q_value"].min()) < 0.10:
        raise AssertionError("A boundary result unexpectedly survives global 10% FDR")
    if secondary.shape != (5, 7) or modifiers.shape != (7, 7):
        raise AssertionError("Compact boundary displays have unexpected dimensions")
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    expected = [
        "Secondary Outcomes",
        "Modifier Boundaries",
        "Estimate Audit",
        "Metadata",
    ]
    if workbook.sheetnames != expected:
        raise AssertionError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Secondary Outcomes"]
    if not str(sheet["A1"].value).startswith("Boundary Outcomes and Unsupported Modifiers"):
        raise AssertionError("Workbook title is missing")
    if sheet["A2"].value != "Outcome" or sheet.freeze_panes != "A3":
        raise AssertionError("Title must be followed immediately by column headers")
    if sheet.max_row != 11 or sheet.max_column != 7:
        raise AssertionError("Unexpected secondary-outcome display dimensions")
    modifier_sheet = workbook["Modifier Boundaries"]
    if modifier_sheet["A2"].value != "Modifier" or modifier_sheet.max_row != 13:
        raise AssertionError("Unexpected modifier-boundary display dimensions")
    if any(item.print_area is None for item in workbook.worksheets):
        raise AssertionError("Every worksheet must have a print area")
    workbook.close()
    for path in [png_path, modifier_png_path]:
        if not path.exists() or path.stat().st_size == 0:
            raise AssertionError(f"PNG review copy is missing: {path}")


def main() -> None:
    args = parse_args()
    output = resolve_path(args.output)
    review_output = resolve_path(args.review_output)
    modifier_review_output = resolve_path(args.modifier_review_output)
    audit_path = resolve_path(args.multiplicity_audit)
    with psycopg.connect(**connection_args(args)) as connection:
        frame = prepare_frame(load_data(connection, args.schema))
    estimates = estimate_all(frame, audit_path)
    secondary, modifiers = build_compact_tables(estimates)
    build_workbook(estimates, secondary, modifiers, output)
    render_sheet_review(output, "Secondary Outcomes", review_output)
    render_sheet_review(output, "Modifier Boundaries", modifier_review_output)
    validate_outputs(
        estimates,
        secondary,
        modifiers,
        output,
        review_output,
        modifier_review_output,
    )
    print(f"table={output}")
    print(f"review={review_output}")
    print(f"modifier_review={modifier_review_output}")
    print(f"display_secondary={secondary.shape} display_modifiers={modifiers.shape}")
    print(f"rows={len(estimates)} min_global_q={estimates['global_bh_q_value'].min():.4f}")
    for block, group in estimates.groupby("block", sort=False):
        print(
            f"block={block} rows={len(group)} "
            f"nominal_p_lt_05={int(group['p_value'].lt(0.05).sum())}"
        )


if __name__ == "__main__":
    main()
