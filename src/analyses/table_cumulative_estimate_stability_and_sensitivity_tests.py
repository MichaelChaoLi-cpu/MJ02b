#!/usr/bin/env python3
"""Generate cumulative estimate stability and sensitivity tests.

Plan: Report native and common samples, frozen control/geography/inference
sensitivities, and leave-one-wave-out estimates for both cumulative humid-heat
definitions.

Framework: AnaSOP Section 5 identifies influential waves, temporal mismatch,
geographic coverage, sample definition, and clustered inference as threats.
Section 6 freezes two cumulative exposure definitions and their sensitivity
set. Section 7 Step 7 evaluates whether either result is driven by one wave,
sample, geography, or inference choice. All analytical inputs are read from
PostgreSQL under a forced read-only session; database writes are never made.
"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from figure_cumulative_estimate_credibility_checks import (
    BASE_CONTROLS,
    HUMID_EXPOSURES,
    connection_args,
    estimate_leave_one_wave_out,
    estimate_robustness,
    fit_one,
    load_data,
    prepare_frame,
    reference_support,
    required_mask,
)
from table_main_cumulative_humid_heat_and_work_participation_estimates import (
    LIGHT_BORDER,
    MEDIUM_BORDER,
    MUTED,
    NAVY,
    PALE_BLUE,
    PALE_GOLD,
    PALE_TEAL,
    TEXT,
    WHITE,
    formatted_p,
    render_review_png,
    resolve_path,
    set_cell_style,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_cumulative_estimate_stability_and_sensitivity_tests.xlsx"
)
DEFAULT_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_cumulative_estimate_stability_and_sensitivity_tests.png"
)

EXPOSURE_DISPLAY = {
    "threshold": {
        "label": "Cumulative WB26 threshold days",
        "unit": "5 additional days across 2 months",
    },
    "continuous": {
        "label": "Two-month mean maximum wet-bulb temperature",
        "unit": "1 C higher 2-month average",
    },
}

ROBUSTNESS_LABELS = {
    "no_rainfall": ("Controls", "Omit monthly rainfall"),
    "no_demographics": ("Controls", "Omit demographic controls"),
    "balanced_districts": ("Geographic coverage", "Districts observed in all nine waves"),
    "exclude_fallback": ("Geographic coverage", "Exclude coastal climate fallback"),
    "province_clustering": ("Inference", "Province-clustered standard errors"),
    "exclude_timing_risk": ("Timing risk", "Exclude 2011-12, 2019, and 2021"),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--review-output", type=Path, default=DEFAULT_REVIEW_OUTPUT)
    return parser.parse_args()


def estimate_native_samples(frame: pd.DataFrame) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for family, metadata in HUMID_EXPOSURES.items():
        term = metadata["term"]
        mask = required_mask(frame, [term], BASE_CONTROLS)
        support = reference_support(frame, mask, [term])
        rows.append(
            fit_one(
                frame,
                panel="sample_contract",
                specification="native_sample",
                specification_label="Native complete-case sample",
                exposure_family=family,
                term=term,
                controls=BASE_CONTROLS,
                mask=mask,
                within_fe_sd=float(support[term]),
            )
        )
    return rows


def evidence_judgment(effect: float, upper: float) -> str:
    if effect < 0 and upper < 0:
        return "Adverse; CI excludes zero"
    if effect < 0:
        return "Adverse; imprecise"
    return "No adverse association"


def organize_estimates(frame: pd.DataFrame) -> pd.DataFrame:
    native = estimate_native_samples(frame)
    robustness = estimate_robustness(frame)
    leave_one_out = estimate_leave_one_wave_out(frame)
    rows: list[dict[str, object]] = []

    for record in native:
        rows.append(
            {
                **record,
                "diagnostic_block": "Sample contract",
                "display_specification": "Native exposure sample",
                "restriction_or_omission": "One exposure definition required",
                "sort_block": 1,
                "sort_specification": 1,
            }
        )

    for record in robustness:
        specification = str(record["specification"])
        if specification == "frozen_primary":
            block = "Sample contract"
            label = "Common exposure sample"
            restriction = "Both cumulative definitions required"
            sort_block = 1
            sort_specification = 2
        else:
            block, label = ROBUSTNESS_LABELS[specification]
            restriction = str(record["specification_label"])
            sort_block = {
                "Controls": 2,
                "Geographic coverage": 3,
                "Inference": 4,
                "Timing risk": 5,
            }[block]
            sort_specification = list(ROBUSTNESS_LABELS).index(specification) + 1
        rows.append(
            {
                **record,
                "diagnostic_block": block,
                "display_specification": label,
                "restriction_or_omission": restriction,
                "sort_block": sort_block,
                "sort_specification": sort_specification,
            }
        )

    wave_order = {
        wave: order
        for order, wave in enumerate(
            ["2007", "2009", "2011-12", "2013", "2014", "2016", "2017", "2019", "2021"],
            start=1,
        )
    }
    for record in leave_one_out:
        wave = str(record["omitted_wave"])
        rows.append(
            {
                **record,
                "diagnostic_block": "Leave one wave out",
                "display_specification": f"Exclude {wave}",
                "restriction_or_omission": f"Omitted survey wave: {wave}",
                "sort_block": 6,
                "sort_specification": wave_order[wave],
            }
        )

    estimates = pd.DataFrame(rows)
    family_order = {"threshold": 1, "continuous": 2}
    estimates["sort_family"] = estimates["exposure_family"].map(family_order)
    estimates = estimates.sort_values(
        ["sort_block", "sort_specification", "sort_family"]
    ).reset_index(drop=True)
    estimates["exposure_display"] = estimates["exposure_family"].map(
        {key: value["label"] for key, value in EXPOSURE_DISPLAY.items()}
    )
    estimates["effect_unit"] = estimates["exposure_family"].map(
        {key: value["unit"] for key, value in EXPOSURE_DISPLAY.items()}
    )
    estimates["clusters_display"] = np.where(
        estimates["cluster_level"].eq("admin1_code"),
        estimates["n_provinces"].astype(int).astype(str) + " provinces",
        estimates["n_admin2"].astype(int).astype(str) + " districts",
    )
    estimates["evidence_judgment"] = [
        evidence_judgment(float(effect), float(upper))
        for effect, upper in zip(
            estimates["raw_effect_pp"], estimates["raw_ci_upper_pp"], strict=True
        )
    ]
    return estimates


def format_ci(lower: float, upper: float) -> str:
    return f"[{lower:.3f}, {upper:.3f}]"


def build_workbook(estimates: pd.DataFrame, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stability Tests"
    sheet.sheet_view.showGridLines = False

    headers = [
        "Diagnostic",
        "Specification",
        "Exposure",
        "Restriction / Omission",
        "Effect Unit",
        "Effect (pp)",
        "95% CI",
        "p threshold",
        "Standardized Effect (pp)",
        "N",
        "Clusters",
        "Evidence Judgment",
    ]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = sheet.cell(1, 1, "Cumulative Estimate Stability and Sensitivity Tests")
    set_cell_style(title, fill=NAVY, color=WHITE, bold=True, size=15, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=8, wrap=True)
    sheet.row_dimensions[2].height = 42

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    previous_block = None
    previous_specification = None
    for offset, row in enumerate(estimates.itertuples(index=False), start=0):
        row_number = 3 + offset
        values = [
            row.diagnostic_block,
            row.display_specification,
            row.exposure_display,
            row.restriction_or_omission,
            row.effect_unit,
            float(row.raw_effect_pp),
            format_ci(float(row.raw_ci_lower_pp), float(row.raw_ci_upper_pp)),
            formatted_p(float(row.p_value)),
            float(row.standardized_effect_pp),
            int(row.n_obs),
            row.clusters_display,
            row.evidence_judgment,
        ]
        starts_block = previous_block is not None and row.diagnostic_block != previous_block
        is_common = row.display_specification == "Common exposure sample"
        fill = PALE_GOLD if is_common else (PALE_BLUE if offset % 2 else WHITE)
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=fill,
                bold=is_common and column in [2, 12],
                size=7.5,
                align="left" if column in [1, 2, 3, 4, 5, 11, 12] else "right",
                wrap=column in [1, 2, 3, 4, 5, 11, 12],
            )
            cell.border = Border(top=group_top if starts_block else Side(), bottom=thin)
        for column in [6, 9]:
            sheet.cell(row_number, column).number_format = "0.000"
        sheet.cell(row_number, 10).number_format = "#,##0"
        sheet.row_dimensions[row_number].height = 27
        previous_block = row.diagnostic_block
        previous_specification = row.display_specification

    last_data_row = 2 + len(estimates)
    notes = [
        "Notes: Effects are percentage-point differences in working in the past week. Threshold effects are per five additional WB26 days; continuous effects are per 1 C higher two-month mean maximum wet-bulb temperature.",
        "Except where explicitly varied, models include age, age squared, female, rural, monthly rainfall, District Calendar Month and Survey Year fixed effects, CSES analysis weights, and District-clustered standard errors.",
        "Standardized effects use the fixed-effect-residual exposure SD from the relevant frozen reference sample. Leave-one-wave-out models keep the full frozen model and omit one released survey wave at a time.",
        "Timing-risk exclusion omits 2011-12, 2019, and 2021 without claiming corrected interview dates. Sensitivity p-values are descriptive and do not create new confirmatory tests.",
        "Database access: read-only SELECT from mda.public; database writes: none. Generated "
        + datetime.now().astimezone().isoformat(timespec="seconds"),
    ]
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=len(headers)
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=7.5, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 24

    widths = [19, 27, 29, 31, 24, 12, 19, 11, 20, 12, 16, 25]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:L{last_data_row}"
    sheet.print_area = f"A1:L{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.15
    sheet.page_margins.right = 0.15
    sheet.page_margins.top = 0.18
    sheet.page_margins.bottom = 0.18

    audit = workbook.create_sheet("Estimate Audit")
    audit.sheet_view.showGridLines = False
    for column, header in enumerate(estimates.columns, start=1):
        display_header = "p_threshold" if header == "p_value" else header
        cell = audit.cell(1, column, display_header)
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=8, wrap=True)
    for row_number, row in enumerate(estimates.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            if estimates.columns[column - 1] == "p_value":
                value = formatted_p(float(value))
            if pd.isna(value):
                value = None
            cell = audit.cell(row_number, column, value)
            set_cell_style(
                cell,
                size=7.5,
                align="left" if isinstance(value, str) else "right",
                wrap=isinstance(value, str),
            )
        audit.row_dimensions[row_number].height = 24
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit.print_area = audit.dimensions
    audit.page_setup.orientation = "landscape"
    audit.page_setup.paperSize = audit.PAPERSIZE_A3
    audit.page_setup.fitToWidth = 1
    audit.page_setup.fitToHeight = 0
    audit.sheet_properties.pageSetUpPr.fitToPage = True
    for column in range(1, audit.max_column + 1):
        audit.column_dimensions[get_column_letter(column)].width = 18

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    metadata_rows = [
        ("Field", "Value"),
        ("Table", "Cumulative Estimate Stability and Sensitivity Tests"),
        ("Outcome", "Worked in the past week"),
        ("Rows", f"{len(estimates)} unique model estimates"),
        ("Sample-contract rows", f"{int(estimates['diagnostic_block'].eq('Sample contract').sum())}"),
        (
            "Frozen sensitivity rows",
            f"{int((~estimates['diagnostic_block'].isin(['Sample contract', 'Leave one wave out'])).sum())}",
        ),
        ("Leave-one-wave-out rows", f"{int(estimates['diagnostic_block'].eq('Leave one wave out').sum())}"),
        ("Fixed effects", "District Calendar Month and Survey Year"),
        ("Weights", "CSES analysis weights"),
        ("Primary inference", "District-clustered, debiased standard errors"),
        ("Database session", "PostgreSQL default_transaction_read_only=on"),
        ("Database writes", "None"),
    ]
    for row_number, values in enumerate(metadata_rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = metadata.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=NAVY if row_number == 1 else None,
                color=WHITE if row_number == 1 else TEXT,
                bold=row_number == 1 or column == 1,
                align="left",
                wrap=True,
            )
    metadata.column_dimensions["A"].width = 30
    metadata.column_dimensions["B"].width = 72
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = metadata.dimensions
    metadata.print_area = metadata.dimensions
    workbook.save(output)


def validate_outputs(estimates: pd.DataFrame, workbook_path: Path, png_path: Path) -> None:
    if len(estimates) != 34:
        raise AssertionError(f"Expected 34 unique sensitivity estimates, observed {len(estimates)}")
    expected_blocks = {
        "Sample contract": 4,
        "Controls": 4,
        "Geographic coverage": 4,
        "Inference": 2,
        "Timing risk": 2,
        "Leave one wave out": 18,
    }
    if estimates["diagnostic_block"].value_counts().to_dict() != expected_blocks:
        raise AssertionError("Unexpected sensitivity block counts")
    if not (estimates["raw_effect_pp"] < 0).all():
        raise AssertionError("At least one frozen sensitivity estimate changed direction")
    if estimates.loc[
        estimates["diagnostic_block"].eq("Leave one wave out"), "specification"
    ].nunique() != 9:
        raise AssertionError("Leave-one-wave-out coverage is incomplete")
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    if workbook.sheetnames != ["Stability Tests", "Estimate Audit", "Metadata"]:
        raise AssertionError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Stability Tests"]
    if sheet["A1"].value != "Cumulative Estimate Stability and Sensitivity Tests":
        raise AssertionError("Workbook title is missing")
    if sheet["A2"].value != "Diagnostic" or sheet.freeze_panes != "A3":
        raise AssertionError("Title must be followed immediately by column headers")
    if any(item.print_area is None for item in workbook.worksheets):
        raise AssertionError("Every worksheet must have a print area")
    workbook.close()
    if not png_path.exists() or png_path.stat().st_size == 0:
        raise AssertionError("PNG review copy is missing")


def main() -> None:
    args = parse_args()
    output = resolve_path(args.output)
    review_output = resolve_path(args.review_output)
    with psycopg.connect(**connection_args(args)) as connection:
        frame = prepare_frame(load_data(connection, args.schema))
    estimates = organize_estimates(frame)
    build_workbook(estimates, output)
    render_review_png(output, review_output)
    validate_outputs(estimates, output, review_output)
    print(f"table={output}")
    print(f"review={review_output}")
    print(f"rows={len(estimates)} all_adverse={(estimates['raw_effect_pp'] < 0).all()}")
    print(
        "ci_excludes_zero="
        f"{int(((estimates['raw_effect_pp'] < 0) & (estimates['raw_ci_upper_pp'] < 0)).sum())}"
    )
    for block, group in estimates.groupby("diagnostic_block", sort=False):
        print(
            f"block={block} rows={len(group)} "
            f"effect_range=[{group['raw_effect_pp'].min():.4f}, {group['raw_effect_pp'].max():.4f}]"
        )


if __name__ == "__main__":
    main()
