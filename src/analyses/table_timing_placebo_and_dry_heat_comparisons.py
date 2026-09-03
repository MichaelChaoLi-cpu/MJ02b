#!/usr/bin/env python3
"""Generate the timing, placebo, and dry-heat comparison table.

Plan: Put current, preceding, cumulative, and following-month humid-heat
associations beside matched-window dry-heat specificity estimates on one
complete-case sample.

Framework: AnaSOP Section 5 uses timing, future-exposure placebo, and dry-heat
specificity diagnostics. Section 6 defines separate survey-weighted models with
District Calendar Month and Survey Year fixed effects, full controls, and
District-clustered standard errors. Section 7 Steps 3 and 6 use this output to
assess cumulative timing and humid-versus-dry specificity. All analytical data
are queried from PostgreSQL in a forced read-only session; database writes are
not permitted or performed.
"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from figure_cumulative_estimate_credibility_checks import (
    BASE_CONTROLS,
    connection_args,
    load_data,
    prepare_frame,
    reference_support,
    required_mask,
)
from run_heat_labor_credibility_audit import fit_hdfe
from table_main_cumulative_humid_heat_and_work_participation_estimates import (
    LIGHT_BORDER,
    MEDIUM_BORDER,
    MUTED,
    NAVY,
    PALE_BLUE,
    PALE_GOLD,
    PALE_TEAL,
    TEXT,
    WHITE,
    render_review_png,
    resolve_path,
    set_cell_style,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_timing_placebo_and_dry_heat_comparisons.xlsx"
)
DEFAULT_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_timing_placebo_and_dry_heat_comparisons.png"
)

ROWS = [
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Threshold humid heat",
        "definition": "WBmax >= 26 C days",
        "window": "Current month",
        "term": "wb26_5days",
        "unit": "5 additional days",
        "evidence_role": "Timing",
    },
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Threshold humid heat",
        "definition": "WBmax >= 26 C days",
        "window": "Previous month",
        "term": "lag_wb26_5days",
        "unit": "5 additional days",
        "evidence_role": "Timing",
    },
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Threshold humid heat",
        "definition": "WBmax >= 26 C days",
        "window": "Current + previous",
        "term": "wb26_current_lag_total_5days",
        "unit": "5 additional days",
        "evidence_role": "Frozen cumulative",
    },
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Threshold humid heat",
        "definition": "WBmax >= 26 C days",
        "window": "Future month",
        "term": "lead_wb26_5days",
        "unit": "5 additional days",
        "evidence_role": "Placebo",
    },
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Continuous humid heat",
        "definition": "Mean daily maximum WB",
        "window": "Current month",
        "term": "wbmean_1c",
        "unit": "1 C higher",
        "evidence_role": "Timing",
    },
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Continuous humid heat",
        "definition": "Mean daily maximum WB",
        "window": "Previous month",
        "term": "lag_wbmean_1c",
        "unit": "1 C higher",
        "evidence_role": "Timing",
    },
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Continuous humid heat",
        "definition": "Mean daily maximum WB",
        "window": "Current + previous average",
        "term": "wbmean_current_lag_average_1c",
        "unit": "1 C higher",
        "evidence_role": "Frozen cumulative",
    },
    {
        "comparison": "Humid-heat timing",
        "exposure_type": "Continuous humid heat",
        "definition": "Mean daily maximum WB",
        "window": "Future month",
        "term": "lead_wbmean_1c",
        "unit": "1 C higher",
        "evidence_role": "Placebo",
    },
    {
        "comparison": "Dry-heat specificity",
        "exposure_type": "Threshold dry heat",
        "definition": "Tmax >= 35 C days",
        "window": "Current + previous",
        "term": "tmax35_current_lag_total_5days",
        "unit": "5 additional days",
        "evidence_role": "Specificity",
    },
    {
        "comparison": "Dry-heat specificity",
        "exposure_type": "Continuous dry heat",
        "definition": "Mean daily maximum Tmax",
        "window": "Current + previous average",
        "term": "tmaxmean_current_lag_average_1c",
        "unit": "1 C higher",
        "evidence_role": "Specificity",
    },
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--review-output", type=Path, default=DEFAULT_REVIEW_OUTPUT)
    return parser.parse_args()


def prepare_timing_frame(frame: pd.DataFrame) -> pd.DataFrame:
    frame = prepare_frame(frame)
    frame["lag_wb26_5days"] = frame["lag_days_wbmax_ge_26c"] / 5.0
    frame["wbmean_1c"] = frame["wet_bulb_max_mean_c"]
    frame["lag_wbmean_1c"] = frame["lag_wet_bulb_max_mean_c"]
    return frame


def evidence_interpretation(row: dict[str, object], estimate: float, upper: float) -> str:
    role = str(row["evidence_role"])
    if role == "Placebo":
        if estimate < 0 and upper < 0:
            return "Adverse placebo; weakens separation"
        if estimate < 0:
            return "Adverse but imprecise placebo"
        return "Null/opposing placebo"
    if role == "Specificity":
        if estimate < 0 and upper < 0:
            return "Adverse dry-heat comparison"
        if estimate < 0:
            return "Adverse but imprecise dry heat"
        return "No adverse dry-heat pattern"
    if estimate < 0 and upper < 0:
        return "Adverse; CI excludes zero"
    if estimate < 0:
        return "Adverse; imprecise"
    return "No adverse association"


def estimate_models(frame: pd.DataFrame) -> pd.DataFrame:
    terms = [str(row["term"]) for row in ROWS]
    mask = required_mask(frame, terms, BASE_CONTROLS)
    support = reference_support(frame, mask, terms)
    estimates: list[dict[str, object]] = []
    for order, metadata in enumerate(ROWS, start=1):
        term = str(metadata["term"])
        records, model = fit_hdfe(
            frame,
            f"timing_placebo_dry_{order:02d}_{term}",
            "worked_past_week",
            [term],
            BASE_CONTROLS,
            sample_label="Common timing-placebo-specificity sample",
            mask=mask,
        )
        record = records[0]
        within_sd = float(support[term])
        effect = 100.0 * float(record["coefficient"])
        lower = 100.0 * float(record["ci_lower_95"])
        upper = 100.0 * float(record["ci_upper_95"])
        estimates.append(
            {
                "row_order": order,
                **metadata,
                "effect_pp": effect,
                "ci_lower_pp": lower,
                "ci_upper_pp": upper,
                "standard_error_pp": 100.0 * float(record["standard_error"]),
                "p_value": float(record["p_value"]),
                "within_fe_sd": within_sd,
                "standardized_effect_pp": effect * within_sd,
                "standardized_ci_lower_pp": lower * within_sd,
                "standardized_ci_upper_pp": upper * within_sd,
                "n_obs": int(model.nobs),
                "n_admin2": int(record["n_admin2"]),
                "n_provinces": int(record["n_provinces"]),
                "controls": str(record["controls"]),
                "fixed_effects": str(record["fixed_effects"]),
                "cluster_level": str(record["cluster_level"]),
                "weights": "CSES analysis weights",
                "interpretation": evidence_interpretation(metadata, effect, upper),
            }
        )
    result = pd.DataFrame(estimates)
    if result["n_obs"].nunique() != 1 or result["n_admin2"].nunique() != 1:
        raise AssertionError("Every comparison must use the same sample and District set")
    if int(result["n_obs"].iloc[0]) != int(mask.sum()):
        raise AssertionError("Model observation count does not equal the common mask")
    return result


def format_ci(lower: float, upper: float) -> str:
    return f"[{lower:.3f}, {upper:.3f}]"


def format_p(value: float) -> str:
    """Return the strongest manuscript-approved p-value threshold."""
    if pd.isna(value) or value >= 0.10:
        return "—"
    if value < 0.01:
        return "p < 0.01"
    if value < 0.05:
        return "p < 0.05"
    return "p < 0.10"


def build_workbook(estimates: pd.DataFrame, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Comparisons"
    sheet.sheet_view.showGridLines = False

    headers = [
        "Comparison",
        "Exposure Type",
        "Exposure Window",
        "Evidence Role",
        "Effect Unit",
        "Effect (pp)",
        "95% CI",
        "p threshold",
        "Standardized Effect (pp)",
        "Within-FE SD",
        "N",
        "Interpretation",
    ]
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title = sheet.cell(1, 1, "Timing, Placebo, and Dry-Heat Comparisons")
    set_cell_style(title, fill=NAVY, color=WHITE, bold=True, size=15, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=8, wrap=True)
    sheet.row_dimensions[2].height = 42

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    previous_type = None
    for offset, row in enumerate(estimates.itertuples(index=False), start=0):
        row_number = 3 + offset
        values = [
            row.comparison,
            row.exposure_type,
            row.window,
            row.evidence_role,
            row.unit,
            float(row.effect_pp),
            format_ci(float(row.ci_lower_pp), float(row.ci_upper_pp)),
            format_p(float(row.p_value)),
            float(row.standardized_effect_pp),
            float(row.within_fe_sd),
            int(row.n_obs),
            row.interpretation,
        ]
        starts_group = previous_type is not None and row.exposure_type != previous_type
        fill = PALE_GOLD if row.evidence_role == "Frozen cumulative" else (
            PALE_BLUE if offset % 2 else WHITE
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=fill,
                bold=row.evidence_role == "Frozen cumulative" and column in [3, 4],
                size=8,
                align="left" if column in [1, 2, 3, 4, 5, 12] else "right",
                wrap=column in [1, 2, 3, 4, 5, 12],
            )
            cell.border = Border(top=group_top if starts_group else Side(), bottom=thin)
        for column in [6, 9, 10]:
            sheet.cell(row_number, column).number_format = "0.000"
        sheet.cell(row_number, 11).number_format = "#,##0"
        sheet.row_dimensions[row_number].height = 31
        previous_type = row.exposure_type

    last_data_row = 2 + len(estimates)
    notes = [
        "Notes: Effects are percentage-point differences in working in the past week. All rows use one common complete-case sample so timing and exposure definitions are directly comparable.",
        "Every model includes age, age squared, female, rural, monthly rainfall, District Calendar Month and Survey Year fixed effects, CSES analysis weights, and District-clustered standard errors.",
        "Standardized effects multiply raw effects by one fixed-effect-residual standard deviation of each exposure across unique district-year-month cells. Future-month exposures are placebos, not treatments.",
        "Interpretation is descriptive: an adverse placebo weakens timing separation; a weak dry-heat result is consistent with, but does not prove, humid-heat specificity. The study remains exploratory after global multiplicity adjustment.",
        "Database access: read-only SELECT from mda.public; database writes: none. Generated "
        + datetime.now().astimezone().isoformat(timespec="seconds"),
    ]
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number, start_column=1, end_row=row_number, end_column=len(headers)
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 25

    widths = [20, 22, 24, 19, 20, 12, 19, 11, 20, 14, 12, 29]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:L{last_data_row}"
    sheet.print_area = f"A1:L{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.18
    sheet.page_margins.right = 0.18
    sheet.page_margins.top = 0.22
    sheet.page_margins.bottom = 0.22

    audit = workbook.create_sheet("Estimate Audit")
    audit.sheet_view.showGridLines = False
    for column, header in enumerate(estimates.columns, start=1):
        display_header = "p_threshold" if header == "p_value" else header
        cell = audit.cell(1, column, display_header)
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=8, wrap=True)
    for row_number, row in enumerate(estimates.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(row, start=1):
            if estimates.columns[column - 1] == "p_value":
                value = format_p(float(value))
            if pd.isna(value):
                value = None
            cell = audit.cell(row_number, column, value)
            set_cell_style(
                cell,
                size=8,
                align="left" if isinstance(value, str) else "right",
                wrap=isinstance(value, str),
            )
        audit.row_dimensions[row_number].height = 25
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit.print_area = audit.dimensions
    audit.page_setup.orientation = "landscape"
    audit.page_setup.paperSize = audit.PAPERSIZE_A3
    audit.page_setup.fitToWidth = 1
    audit.page_setup.fitToHeight = 0
    audit.sheet_properties.pageSetUpPr.fitToPage = True
    for column in range(1, audit.max_column + 1):
        audit.column_dimensions[get_column_letter(column)].width = 18

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    metadata_rows = [
        ("Field", "Value"),
        ("Table", "Timing, Placebo, and Dry-Heat Comparisons"),
        ("Outcome", "Worked in the past week"),
        ("Common sample", f"{int(estimates['n_obs'].iloc[0]):,} person-wave records"),
        ("Districts", f"{int(estimates['n_admin2'].iloc[0])}"),
        ("Fixed effects", "District Calendar Month and Survey Year"),
        ("Controls", "Age, age squared, female, rural, and monthly rainfall"),
        ("Weights", "CSES analysis weights"),
        ("Inference", "District-clustered, debiased standard errors"),
        ("Database session", "PostgreSQL default_transaction_read_only=on"),
        ("Database writes", "None"),
    ]
    for row_number, values in enumerate(metadata_rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = metadata.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=NAVY if row_number == 1 else None,
                color=WHITE if row_number == 1 else TEXT,
                bold=row_number == 1 or column == 1,
                align="left",
                wrap=True,
            )
    metadata.column_dimensions["A"].width = 25
    metadata.column_dimensions["B"].width = 72
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = metadata.dimensions
    metadata.print_area = metadata.dimensions
    workbook.save(output)


def validate_outputs(estimates: pd.DataFrame, workbook_path: Path, png_path: Path) -> None:
    if len(estimates) != 10:
        raise AssertionError(f"Expected 10 unique comparisons, observed {len(estimates)}")
    if estimates["term"].nunique() != 10:
        raise AssertionError("Comparison rows must not duplicate a model term")
    if estimates["n_obs"].nunique() != 1 or estimates["n_admin2"].nunique() != 1:
        raise AssertionError("Comparison rows do not share a common sample")
    expected_roles = {"Timing": 4, "Frozen cumulative": 2, "Placebo": 2, "Specificity": 2}
    if estimates["evidence_role"].value_counts().to_dict() != expected_roles:
        raise AssertionError("Unexpected timing/placebo/specificity role counts")
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    if workbook.sheetnames != ["Comparisons", "Estimate Audit", "Metadata"]:
        raise AssertionError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Comparisons"]
    if sheet["A1"].value != "Timing, Placebo, and Dry-Heat Comparisons":
        raise AssertionError("Workbook title is missing")
    if sheet["A2"].value != "Comparison" or sheet.freeze_panes != "A3":
        raise AssertionError("Title must be followed immediately by column headers")
    if any(item.print_area is None for item in workbook.worksheets):
        raise AssertionError("Every worksheet must have a print area")
    workbook.close()
    if not png_path.exists() or png_path.stat().st_size == 0:
        raise AssertionError("PNG review copy is missing")


def main() -> None:
    args = parse_args()
    output = resolve_path(args.output)
    review_output = resolve_path(args.review_output)
    with psycopg.connect(**connection_args(args)) as connection:
        frame = prepare_timing_frame(load_data(connection, args.schema))
    estimates = estimate_models(frame)
    build_workbook(estimates, output)
    render_review_png(output, review_output)
    validate_outputs(estimates, output, review_output)
    print(f"table={output}")
    print(f"review={review_output}")
    print(
        f"rows={len(estimates)} n={int(estimates['n_obs'].iloc[0])} "
        f"districts={int(estimates['n_admin2'].iloc[0])}"
    )
    for row in estimates.itertuples(index=False):
        print(
            f"term={row.term} effect_pp={row.effect_pp:.4f} "
            f"std_effect_pp={row.standardized_effect_pp:.4f} p={row.p_value:.4f}"
        )


if __name__ == "__main__":
    main()
