#!/usr/bin/env python3
"""Education Difference Stability and Sensitivity Tests.

Plan: Report both frozen cumulative humid-heat-by-low-education interactions
under leave-one-wave-out and prespecified sample, geography, timing, and
clustering checks.

Framework: AnaSOP Sections 5-7 define the low-minus-higher education
interaction as an exploratory vulnerability contrast. This table reports all
30 frozen stability estimates without selecting by statistical significance.
All PostgreSQL inputs are queried under a forced read-only session.
"""

from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter

from figure_cumulative_humid_heat_by_education import (
    connection_args,
    estimate_stability,
    load_data,
    prepare_frame,
)
from table_main_cumulative_humid_heat_and_work_participation_estimates import (
    LIGHT_BORDER,
    MEDIUM_BORDER,
    MUTED,
    NAVY,
    PALE_BLUE,
    PALE_GOLD,
    PALE_TEAL,
    TEXT,
    WHITE,
    formatted_p,
    render_review_png,
    resolve_path,
    set_cell_style,
)


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_education_difference_stability_and_sensitivity_tests.xlsx"
)
DEFAULT_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_education_difference_stability_and_sensitivity_tests.png"
)

DISPLAY = {
    "threshold": ("Cumulative WB26 threshold days", "5 additional days / 2 months"),
    "continuous": ("Two-month mean maximum WB temperature", "1 C higher 2-month mean"),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--review-output", type=Path, default=DEFAULT_REVIEW_OUTPUT)
    return parser.parse_args()


def format_ci(lower: float, upper: float) -> str:
    return f"[{lower:.3f}, {upper:.3f}]"


def evidence_judgment(effect: float, upper: float) -> str:
    if effect < 0 and upper < 0:
        return "More adverse for low education; CI excludes zero"
    if effect < 0:
        return "More adverse for low education; imprecise"
    return "Direction does not support greater low-education vulnerability"


def prepare_table(estimates: pd.DataFrame) -> pd.DataFrame:
    if len(estimates) != 30:
        raise AssertionError(f"Expected 30 stability estimates, observed {len(estimates)}")
    rows: list[dict[str, object]] = []
    for row in estimates.itertuples(index=False):
        exposure, unit = DISPLAY[str(row.exposure_family)]
        if row.panel == "leave_one_wave_out":
            check = "Leave one wave out"
            restriction = f"Exclude {row.omitted_wave}"
        else:
            check = "Frozen sensitivity"
            restriction = str(row.specification_label)
        if row.cluster_level == "admin1_code":
            cluster = f"{int(row.n_provinces)} provinces / Province"
        else:
            cluster = f"{int(row.n_admin2)} districts / District"
        rows.append(
            {
                "Exposure": exposure,
                "Check": check,
                "Specification": str(row.specification_label),
                "Omitted Wave / Restriction": restriction,
                "Effect Unit": unit,
                "Interaction Effect (pp)": float(row.raw_effect_pp),
                "95% CI": format_ci(float(row.raw_ci_lower_pp), float(row.raw_ci_upper_pp)),
                "Standardized Effect (pp)": float(row.standardized_effect_pp),
                "p threshold": formatted_p(float(row.p_value)),
                "N": int(row.n_obs),
                "Clusters / Level": cluster,
                "Evidence Judgment": evidence_judgment(
                    float(row.raw_effect_pp), float(row.raw_ci_upper_pp)
                ),
            }
        )
    table = pd.DataFrame(rows)
    if not table["Interaction Effect (pp)"].lt(0).all():
        raise AssertionError("At least one frozen interaction changed direction")
    return table


def build_workbook(
    table: pd.DataFrame,
    audit_frame: pd.DataFrame,
    output: Path,
    args: argparse.Namespace,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Education Stability"
    sheet.sheet_view.showGridLines = False
    headers = list(table.columns)
    last_column = len(headers)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title = sheet.cell(1, 1, "Education Difference Stability and Sensitivity Tests")
    set_cell_style(title, fill=NAVY, color=WHITE, bold=True, size=15, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=8, wrap=True)
    sheet.row_dimensions[2].height = 43

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    previous_check = None
    for offset, row in enumerate(table.itertuples(index=False, name=None), start=0):
        row_number = 3 + offset
        check = str(row[1])
        starts_group = previous_check is not None and check != previous_check
        fill = PALE_GOLD if check == "Frozen sensitivity" else (PALE_BLUE if offset % 2 else WHITE)
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=fill,
                bold=column == 12 and "excludes zero" in str(value),
                size=7.5,
                align="left" if column in [1, 2, 3, 4, 5, 7, 11, 12] else "right",
                wrap=column in [1, 2, 3, 4, 5, 7, 11, 12],
            )
            cell.border = Border(top=group_top if starts_group else Side(), bottom=thin)
        for column in [6, 8]:
            sheet.cell(row_number, column).number_format = "0.000"
        sheet.cell(row_number, 10).number_format = "#,##0"
        sheet.row_dimensions[row_number].height = 31
        previous_check = check

    last_data_row = 2 + len(table)
    notes = [
        "Notes: The interaction is low education minus higher education; negative values indicate a more adverse association for the low-education group.",
        "All models include age, age squared, female, rural, monthly rainfall, District Calendar Month and Survey Year fixed effects, and CSES analysis weights.",
        "Leave-one-wave-out and sensitivity estimates are frozen stability diagnostics and are not added to the Benjamini-Hochberg testing universe.",
        "The education contrast is associational and exploratory; it is not the causal effect of schooling. Database access was read-only SELECT; writes: none.",
    ]
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(
            start_row=row_number,
            start_column=1,
            end_row=row_number,
            end_column=last_column,
        )
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 23

    widths = [27, 19, 26, 27, 24, 16, 20, 20, 11, 13, 23, 42]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:L{last_data_row}"
    sheet.print_area = f"A1:L{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.16
    sheet.page_margins.right = 0.16
    sheet.page_margins.top = 0.20
    sheet.page_margins.bottom = 0.20

    audit = workbook.create_sheet("Estimate Audit")
    audit.sheet_view.showGridLines = False
    for column, header in enumerate(audit_frame.columns, start=1):
        display_header = "p_threshold" if header == "p_value" else header
        cell = audit.cell(1, column, display_header)
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=7, wrap=True)
    for row_number, row in enumerate(
        audit_frame.itertuples(index=False, name=None), start=2
    ):
        for column, value in enumerate(row, start=1):
            if audit_frame.columns[column - 1] == "p_value":
                value = formatted_p(float(value))
            if pd.isna(value):
                value = None
            cell = audit.cell(row_number, column, value)
            set_cell_style(
                cell,
                size=7,
                align="left" if isinstance(value, str) else "right",
                wrap=isinstance(value, str),
            )
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit.print_area = audit.dimensions
    audit.page_setup.orientation = "landscape"
    audit.page_setup.paperSize = audit.PAPERSIZE_A3
    audit.page_setup.fitToWidth = 1
    audit.page_setup.fitToHeight = 0
    audit.sheet_properties.pageSetUpPr.fitToPage = True
    for column in range(1, audit.max_column + 1):
        audit.column_dimensions[get_column_letter(column)].width = 17

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    metadata_rows = [
        ("Field", "Value"),
        ("Table", "Education Difference Stability and Sensitivity Tests"),
        ("Generated At", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Database", args.dbname),
        ("Schema", args.schema),
        ("Database Access", "read-only SELECT; writes: none"),
        ("Rows", len(table)),
        ("Leave-One-Wave-Out Rows", int(table["Check"].eq("Leave one wave out").sum())),
        ("Frozen Sensitivity Rows", int(table["Check"].eq("Frozen sensitivity").sum())),
        ("Direction", f"{int(table['Interaction Effect (pp)'].lt(0).sum())}/{len(table)} negative"),
    ]
    for row_number, values in enumerate(metadata_rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = metadata.cell(row_number, column, value)
            set_cell_style(
                cell,
                fill=NAVY if row_number == 1 else None,
                color=WHITE if row_number == 1 else TEXT,
                bold=row_number == 1 or column == 1,
                align="left",
                wrap=True,
            )
    metadata.column_dimensions["A"].width = 31
    metadata.column_dimensions["B"].width = 76
    metadata.freeze_panes = "A2"
    metadata.auto_filter.ref = metadata.dimensions
    metadata.print_area = metadata.dimensions
    workbook.save(output)


def validate_workbook(table: pd.DataFrame, output: Path, review: Path) -> None:
    workbook = load_workbook(output, data_only=False, read_only=False)
    if workbook.sheetnames != ["Education Stability", "Estimate Audit", "Metadata"]:
        raise AssertionError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["Education Stability"]
    if sheet["A1"].value != "Education Difference Stability and Sensitivity Tests":
        raise AssertionError("Workbook title is missing")
    if sheet["A2"].value != "Exposure" or sheet.freeze_panes != "A3":
        raise AssertionError("Title must be followed immediately by column headers")
    formulas = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if formulas:
        raise AssertionError("Formal result sheet must contain cached values, not formulas")
    workbook.close()
    cached = load_workbook(output, data_only=True, read_only=True)
    result = cached["Education Stability"]
    values = [
        result.cell(row, column).value
        for row in range(3, 33)
        for column in range(1, 13)
    ]
    if any(value is None for value in values):
        raise AssertionError("Cached-value read found a blank result cell")
    cached.close()
    if len(table) != 30 or not review.exists() or review.stat().st_size == 0:
        raise AssertionError("Education stability outputs are incomplete")


def main() -> None:
    args = parse_args()
    output = resolve_path(args.output)
    review = resolve_path(args.review_output)
    with psycopg.connect(**connection_args(args)) as connection:
        frame = prepare_frame(load_data(connection, args.schema))
        connection.rollback()
    estimates = estimate_stability(frame)
    table = prepare_table(estimates)
    build_workbook(table, estimates, output, args)
    render_review_png(output, review)
    validate_workbook(table, output, review)
    print(f"saved_workbook={output.relative_to(ROOT)}")
    print(f"saved_review_png={review.relative_to(ROOT)}")
    print(f"rows={len(table)} negative={int(table['Interaction Effect (pp)'].lt(0).sum())}")


if __name__ == "__main__":
    main()
