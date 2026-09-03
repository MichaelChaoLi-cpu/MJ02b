#!/usr/bin/env python3
"""Generate the two descriptive and identifying-variation tables.

Plan: Summarize weighted labor outcomes, person characteristics, climate
exposures, household resources, and fixed-effect-residual heat variation.

Framework: AnaSOP Section 5 identifying-variation diagnostic, Section 6 frozen
person- and household-sample contracts, and Section 7 Step 2. All database
inputs are queried under a forced read-only PostgreSQL session.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import psycopg
from linearmodels.iv import AbsorbingLS
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from psycopg import sql


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DESCRIPTIVE_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_descriptive_statistics_of_the_analytical_sample.xlsx"
)
DEFAULT_DESCRIPTIVE_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_descriptive_statistics_of_the_analytical_sample.png"
)
DEFAULT_IDENTIFICATION_OUTPUT = (
    ROOT
    / "data/results/tables"
    / "Table_identifying_variation_in_humid_heat_exposure.xlsx"
)
DEFAULT_IDENTIFICATION_REVIEW_OUTPUT = (
    ROOT
    / "data/exp/figure-table-review"
    / "Table_identifying_variation_in_humid_heat_exposure.png"
)

NAVY = "17324D"
PALE_TEAL = "DCEBEA"
PALE_BLUE = "EEF4F7"
PALE_GOLD = "FFF2CC"
WHITE = "FFFFFF"
TEXT = "23313D"
MUTED = "5C6B73"
LIGHT_BORDER = "B8C7CE"
MEDIUM_BORDER = "6F8791"

PERSON_SAMPLE = "Person rows / Person Weight"
WORKER_SAMPLE = "Current workers / Person Weight"
CELL_SAMPLE = "District-month cells / Unweighted"
HOUSEHOLD_SAMPLE = "Household-wave / Household Weight"


def parse_args(default_table: str = "both") -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument(
        "--table",
        choices=["descriptive", "identification", "both"],
        default=default_table,
    )
    parser.add_argument(
        "--descriptive-output", type=Path, default=DEFAULT_DESCRIPTIVE_OUTPUT
    )
    parser.add_argument(
        "--descriptive-review-output",
        type=Path,
        default=DEFAULT_DESCRIPTIVE_REVIEW_OUTPUT,
    )
    parser.add_argument(
        "--identification-output", type=Path, default=DEFAULT_IDENTIFICATION_OUTPUT
    )
    parser.add_argument(
        "--identification-review-output",
        type=Path,
        default=DEFAULT_IDENTIFICATION_REVIEW_OUTPUT,
    )
    parser.add_argument("--skip-render", action="store_true")
    return parser.parse_args()


def connection_args(args: argparse.Namespace) -> dict[str, object]:
    values: dict[str, object] = {
        "host": args.host,
        "port": args.port,
        "dbname": args.dbname,
        "options": "-c default_transaction_read_only=on",
    }
    if args.user:
        values["user"] = args.user
    if args.password:
        values["password"] = args.password
    return values


def load_analysis_data(connection: psycopg.Connection, schema: str) -> pd.DataFrame:
    query = sql.SQL(
        """
        SELECT
            a.survey_wave,
            a.survey_year,
            a.survey_month,
            a.person_id,
            a.household_id,
            a.admin2_code,
            a.sex,
            a.age,
            a.urban_rural,
            a.analysis_weight,
            a.household_weight,
            a.worked_past_week,
            a.weekly_hours_including_zero,
            a.log_monthly_salary_wages,
            a.days_wbmax_ge_26c,
            a.lag_days_wbmax_ge_26c,
            a.days_wbmax_ge_28c,
            a.wet_bulb_max_mean_c,
            a.lag_wet_bulb_max_mean_c,
            a.days_tmax_ge_35c,
            a.apparent_temperature_max_mean_c,
            a.precipitation_month_sum_mm,
            a.household_member_count,
            a.rooms_used,
            a.floor_area_square_meters,
            a.electricity_spending_positive,
            a.monthly_electricity_expense_riel,
            a.monthly_water_charges_riel,
            ec.main_occupation_source_code,
            ed.education_level_harmonized,
            climate.days_observed
        FROM {}.{} AS a
        LEFT JOIN {}.{} AS ec
          ON ec.survey_wave = a.survey_wave
         AND ec.person_id = a.person_id
        LEFT JOIN {}.{} AS ed
          ON ed.survey_wave = a.survey_wave
         AND ed.person_id = a.person_id
        LEFT JOIN {}.{} AS climate
          ON climate.admin2_code = a.admin2_code
         AND climate.reference_year = a.survey_year
         AND climate.calendar_month = a.survey_month
        WHERE a.analysis_weight > 0
        ORDER BY a.survey_year, a.person_id
        """
    ).format(
        sql.Identifier(schema),
        sql.Identifier("final_HEAT_LABOR_ANALYTIC"),
        sql.Identifier(schema),
        sql.Identifier("final_EC_CSES"),
        sql.Identifier(schema),
        sql.Identifier("final_ED_CSES"),
        sql.Identifier(schema),
        sql.Identifier("final_CLIMATE_MONTHLY_ADMIN2"),
    )
    with connection.cursor() as cursor:
        cursor.execute(query)
        columns = [column.name for column in cursor.description]
        rows = cursor.fetchall()
    frame = pd.DataFrame(rows, columns=columns)
    numeric = [
        "survey_year",
        "survey_month",
        "sex",
        "age",
        "urban_rural",
        "analysis_weight",
        "household_weight",
        "worked_past_week",
        "weekly_hours_including_zero",
        "log_monthly_salary_wages",
        "days_wbmax_ge_26c",
        "lag_days_wbmax_ge_26c",
        "days_wbmax_ge_28c",
        "wet_bulb_max_mean_c",
        "lag_wet_bulb_max_mean_c",
        "days_tmax_ge_35c",
        "apparent_temperature_max_mean_c",
        "precipitation_month_sum_mm",
        "household_member_count",
        "rooms_used",
        "floor_area_square_meters",
        "electricity_spending_positive",
        "monthly_electricity_expense_riel",
        "monthly_water_charges_riel",
        "education_level_harmonized",
        "days_observed",
    ]
    for column in numeric:
        frame[column] = pd.to_numeric(frame[column], errors="coerce")
    return frame


def construct_variables(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    result["female"] = result["sex"].eq(2).where(result["sex"].isin([1, 2])).astype(float)
    result["rural"] = result["urban_rural"].eq(2).where(
        result["urban_rural"].isin([1, 2])
    ).astype(float)
    result["low_education"] = result["education_level_harmonized"].le(2).where(
        result["education_level_harmonized"].between(0, 7)
    ).astype(float)
    result["older_worker"] = result["age"].ge(50).where(result["age"].notna()).astype(float)

    occupation = result["main_occupation_source_code"].astype("string")
    valid_occupation = occupation.str.fullmatch(r"\d{3}", na=False)
    result["agriculture_occupation_candidate"] = (
        occupation.str.startswith("6").where(valid_occupation).astype(float)
    )

    result["five_day_humid_heat_exposure"] = result["days_wbmax_ge_26c"] / 5.0
    result["two_month_cumulative_extreme_humid_heat_days"] = (
        result["days_wbmax_ge_26c"] + result["lag_days_wbmax_ge_26c"]
    )
    result["two_month_average_max_wet_bulb_c"] = (
        result["wet_bulb_max_mean_c"] + result["lag_wet_bulb_max_mean_c"]
    ) / 2.0
    result["monthly_humid_heat_share"] = (
        result["days_wbmax_ge_26c"] / result["days_observed"]
    ).where(result["days_observed"].gt(0))

    positive_rooms = result["rooms_used"].gt(0) & result["household_member_count"].gt(0)
    result["persons_per_room"] = (
        result["household_member_count"] / result["rooms_used"]
    ).where(positive_rooms)
    positive_members = result["household_member_count"].gt(0)
    result["floor_area_per_capita"] = (
        result["floor_area_square_meters"] / result["household_member_count"]
    ).where(positive_members & result["floor_area_square_meters"].ge(0))
    result["log_electricity_expenditure"] = np.log1p(
        result["monthly_electricity_expense_riel"].where(
            result["monthly_electricity_expense_riel"].ge(0)
        )
    )
    result["log_water_expenditure"] = np.log1p(
        result["monthly_water_charges_riel"].where(
            result["monthly_water_charges_riel"].ge(0)
        )
    )
    return result


def weighted_quantiles(
    values: np.ndarray, weights: np.ndarray
) -> tuple[float, float, float, float, float]:
    order = np.argsort(values, kind="mergesort")
    sorted_values = values[order]
    sorted_weights = weights[order]
    positions = (np.cumsum(sorted_weights) - 0.5 * sorted_weights) / sorted_weights.sum()
    quantiles = np.interp([0.05, 0.25, 0.50, 0.75, 0.95], positions, sorted_values)
    return tuple(float(value) for value in quantiles)


def summarize_series(
    values: pd.Series,
    weights: pd.Series,
    base_mask: pd.Series,
) -> dict[str, float | int]:
    base = base_mask.fillna(False) & weights.notna() & weights.gt(0)
    observed = base & values.notna()
    sample_n = int(base.sum())
    observed_n = int(observed.sum())
    if sample_n == 0 or observed_n == 0:
        return {
            "sample_n": sample_n,
            "observed_n": observed_n,
            "mean": np.nan,
            "standard_deviation": np.nan,
            "p05": np.nan,
            "p25": np.nan,
            "median": np.nan,
            "p75": np.nan,
            "p95": np.nan,
            "missing_rate": 1.0 if sample_n else np.nan,
        }
    numeric_values = values.loc[observed].astype(float).to_numpy()
    numeric_weights = weights.loc[observed].astype(float).to_numpy()
    mean = float(np.average(numeric_values, weights=numeric_weights))
    variance = float(np.average((numeric_values - mean) ** 2, weights=numeric_weights))
    p05, p25, median, p75, p95 = weighted_quantiles(numeric_values, numeric_weights)
    return {
        "sample_n": sample_n,
        "observed_n": observed_n,
        "mean": mean,
        "standard_deviation": float(np.sqrt(max(variance, 0.0))),
        "p05": p05,
        "p25": p25,
        "median": median,
        "p75": p75,
        "p95": p95,
        "missing_rate": 1.0 - observed_n / sample_n,
    }


def residualize_heat(cells: pd.DataFrame, exposure: str) -> pd.Series:
    working = cells.dropna(subset=[exposure]).copy()
    working["admin2_calendar_month"] = (
        working["admin2_code"].astype("string")
        + "-m"
        + working["survey_month"].astype("Int64").astype("string").str.zfill(2)
    )
    absorb = pd.DataFrame(
        {
            "admin2_calendar_month": working["admin2_calendar_month"].astype("category"),
            "survey_year": working["survey_year"].astype("Int64").astype("category"),
        },
        index=working.index,
    )
    model = AbsorbingLS(
        working[exposure].astype(float),
        pd.DataFrame({"constant": 1.0}, index=working.index),
        absorb=absorb,
        drop_absorbed=True,
    ).fit()
    residuals = pd.Series(model.resids, index=working.index, dtype=float)
    return residuals.reindex(cells.index)


def build_statistics(frame: pd.DataFrame) -> pd.DataFrame:
    people = frame.copy()
    person_weight = people["analysis_weight"]
    person_mask = pd.Series(True, index=people.index)
    worker_mask = people["worked_past_week"].eq(1)

    cells = people.drop_duplicates(
        ["admin2_code", "survey_year", "survey_month"]
    ).reset_index(drop=True)
    cells["residual_two_month_cumulative_extreme_humid_heat_days"] = residualize_heat(
        cells, "two_month_cumulative_extreme_humid_heat_days"
    )
    cell_weight = pd.Series(1.0, index=cells.index)
    cell_mask = pd.Series(True, index=cells.index)

    households = people.drop_duplicates(["survey_wave", "household_id"]).reset_index(drop=True)
    household_weight = households["household_weight"]
    household_mask = pd.Series(True, index=households.index)

    specifications = [
        ("Labor outcomes", "Work Participation", PERSON_SAMPLE, people, "worked_past_week", person_weight, person_mask, "binary"),
        ("Labor outcomes", "Weekly Hours Including Nonworkers", PERSON_SAMPLE, people, "weekly_hours_including_zero", person_weight, person_mask, "continuous"),
        ("Labor outcomes", "Log Monthly Wage", PERSON_SAMPLE, people, "log_monthly_salary_wages", person_weight, person_mask, "log"),
        ("Individual characteristics", "Age (years)", PERSON_SAMPLE, people, "age", person_weight, person_mask, "continuous"),
        ("Individual characteristics", "Female", PERSON_SAMPLE, people, "female", person_weight, person_mask, "binary"),
        ("Individual characteristics", "Rural", PERSON_SAMPLE, people, "rural", person_weight, person_mask, "binary"),
        ("Individual characteristics", "Low Education", PERSON_SAMPLE, people, "low_education", person_weight, person_mask, "binary"),
        ("Individual characteristics", "Older Worker", PERSON_SAMPLE, people, "older_worker", person_weight, person_mask, "binary"),
        ("Individual characteristics", "Agricultural Occupation Candidate", WORKER_SAMPLE, people, "agriculture_occupation_candidate", person_weight, worker_mask, "binary"),
        ("Climate exposure", "Extreme Humid-Heat Days", CELL_SAMPLE, cells, "days_wbmax_ge_26c", cell_weight, cell_mask, "continuous"),
        ("Climate exposure", "Two-Month Cumulative Extreme Humid-Heat Days", CELL_SAMPLE, cells, "two_month_cumulative_extreme_humid_heat_days", cell_weight, cell_mask, "continuous"),
        ("Climate exposure", "Monthly Humid-Heat Share", CELL_SAMPLE, cells, "monthly_humid_heat_share", cell_weight, cell_mask, "binary"),
        ("Climate exposure", "Two-Month Average Maximum Wet-Bulb Temperature (C)", CELL_SAMPLE, cells, "two_month_average_max_wet_bulb_c", cell_weight, cell_mask, "continuous"),
        ("Climate exposure", "Severe Humid-Heat Days", CELL_SAMPLE, cells, "days_wbmax_ge_28c", cell_weight, cell_mask, "continuous"),
        ("Climate exposure", "Extreme Dry-Heat Days", CELL_SAMPLE, cells, "days_tmax_ge_35c", cell_weight, cell_mask, "continuous"),
        ("Climate exposure", "Mean Maximum Apparent Temperature (C)", CELL_SAMPLE, cells, "apparent_temperature_max_mean_c", cell_weight, cell_mask, "continuous"),
        ("Climate exposure", "Monthly Rainfall (mm)", CELL_SAMPLE, cells, "precipitation_month_sum_mm", cell_weight, cell_mask, "continuous"),
        ("Identification diagnostic", "Residual Two-Month Cumulative Extreme Humid-Heat Days", CELL_SAMPLE, cells, "residual_two_month_cumulative_extreme_humid_heat_days", cell_weight, cell_mask, "continuous"),
        ("Household resources", "Household Member Count", HOUSEHOLD_SAMPLE, households, "household_member_count", household_weight, household_mask, "continuous"),
        ("Household resources", "Persons per Room", HOUSEHOLD_SAMPLE, households, "persons_per_room", household_weight, household_mask, "continuous"),
        ("Household resources", "Floor Area per Capita (sq m)", HOUSEHOLD_SAMPLE, households, "floor_area_per_capita", household_weight, household_mask, "continuous"),
        ("Household resources", "Positive Electricity Spending", HOUSEHOLD_SAMPLE, households, "electricity_spending_positive", household_weight, household_mask, "binary"),
        ("Household resources", "Log Electricity Expenditure", HOUSEHOLD_SAMPLE, households, "log_electricity_expenditure", household_weight, household_mask, "log"),
        ("Household resources", "Log Water Expenditure", HOUSEHOLD_SAMPLE, households, "log_water_expenditure", household_weight, household_mask, "log"),
    ]

    rows: list[dict[str, object]] = []
    for domain, label, sample, source, column, weights, mask, display_type in specifications:
        row = {
            "domain": domain,
            "variable": label,
            "sample_weighting": sample,
            "display_type": display_type,
        }
        row.update(summarize_series(source[column], weights, mask))
        rows.append(row)
    return pd.DataFrame(rows)


def set_cell_style(
    cell: object,
    *,
    fill: str | None = None,
    color: str = TEXT,
    bold: bool = False,
    size: int = 10,
    align: str = "center",
    wrap: bool = False,
) -> None:
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)


def build_workbook(
    statistics: pd.DataFrame,
    output: Path,
    args: argparse.Namespace,
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Descriptive Statistics"
    sheet.sheet_view.showGridLines = False

    headers = [
        "Domain",
        "Variable",
        "Sample / Weighting",
        "N",
        "Mean",
        "Standard Deviation",
        "P25",
        "Median",
        "P75",
        "Missing Rate",
    ]
    sheet.merge_cells("A1:J1")
    title = sheet["A1"]
    title.value = "Descriptive Statistics and Identifying Variation"
    set_cell_style(title, fill=NAVY, color=WHITE, bold=True, size=15, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=9, wrap=True)
    sheet.row_dimensions[2].height = 40

    audit = workbook.create_sheet("Statistics Audit")
    audit.sheet_view.showGridLines = False
    audit_headers = [
        "domain",
        "variable",
        "sample_weighting",
        "display_type",
        "sample_n",
        "observed_n",
        "mean",
        "standard_deviation",
        "p25",
        "median",
        "p75",
        "missing_rate",
    ]
    audit.append(audit_headers)
    for row_number, row in enumerate(statistics.itertuples(index=False), start=2):
        audit.append(
            [
                row.domain,
                row.variable,
                row.sample_weighting,
                row.display_type,
                int(row.sample_n),
                int(row.observed_n),
                float(row.mean) if pd.notna(row.mean) else None,
                float(row.standard_deviation) if pd.notna(row.standard_deviation) else None,
                float(row.p25) if pd.notna(row.p25) else None,
                float(row.median) if pd.notna(row.median) else None,
                float(row.p75) if pd.notna(row.p75) else None,
                float(row.missing_rate),
            ]
        )
    for cell in audit[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=9, wrap=True)
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit_widths = [20, 32, 28, 14, 12, 12, 14, 18, 12, 12, 12, 14]
    for column, width in enumerate(audit_widths, start=1):
        audit.column_dimensions[get_column_letter(column)].width = width
    for row_number in range(2, audit.max_row + 1):
        for column in range(1, audit.max_column + 1):
            cell = audit.cell(row_number, column)
            set_cell_style(
                cell,
                size=8,
                align="left" if column <= 4 else "right",
                wrap=column <= 4,
            )
        audit.cell(row_number, 5).number_format = "#,##0"
        audit.cell(row_number, 6).number_format = "#,##0"
        for column in range(7, 12):
            audit.cell(row_number, column).number_format = "0.000"
        audit.cell(row_number, 12).number_format = "0.0%"
        audit.row_dimensions[row_number].height = 20
    audit.print_area = f"A1:L{audit.max_row}"
    audit.page_setup.orientation = "landscape"
    audit.page_setup.paperSize = audit.PAPERSIZE_LETTER
    audit.page_setup.fitToWidth = 1
    audit.page_setup.fitToHeight = 1
    audit.sheet_properties.pageSetUpPr.fitToPage = True
    audit.page_margins.left = 0.20
    audit.page_margins.right = 0.20
    audit.page_margins.top = 0.25
    audit.page_margins.bottom = 0.25

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    first_data_row = 3
    last_data_row = first_data_row + len(statistics) - 1
    previous_domain = None
    for offset, row in enumerate(statistics.itertuples(index=False), start=0):
        target_row = first_data_row + offset
        values = [
            row.domain,
            row.variable,
            row.sample_weighting,
            int(row.observed_n),
            float(row.mean),
            float(row.standard_deviation),
            float(row.p25),
            float(row.median),
            float(row.p75),
            float(row.missing_rate),
        ]
        fill = PALE_BLUE if offset % 2 else WHITE
        domain_changed = previous_domain is not None and row.domain != previous_domain
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(target_row, column, value)
            set_cell_style(
                cell,
                fill=fill,
                size=8,
                align="left" if column <= 3 else "right",
                wrap=column <= 3,
            )
            cell.border = Border(top=group_top if domain_changed else Side(), bottom=thin)
        sheet.cell(target_row, 4).number_format = "#,##0"
        number_format = "0.0%" if row.display_type == "binary" else "0.000" if row.display_type == "log" else "0.00"
        for column in range(5, 10):
            sheet.cell(target_row, column).number_format = number_format
        sheet.cell(target_row, 10).number_format = "0.0%"
        sheet.row_dimensions[target_row].height = 21
        previous_domain = row.domain

    notes = [
        "Notes: Means, standard deviations, and quantiles use the weighting stated in each row; Missing Rate is the unweighted missing share within that row's stated sample.",
        "Climate rows use 3,325 unique district-year-month cells; 3,114 cells have valid current-plus-previous-month exposure. Residual Two-Month Cumulative Extreme Humid-Heat Days absorb District Calendar Month and Survey Year fixed effects.",
        "Household rows use one record per household-wave. Agricultural Occupation Candidate is restricted to current workers and remains a conservative classification.",
        "Log expenditure variables use log(1 + nonnegative nominal riel). Database access was read-only SELECT; database writes: none.",
    ]
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=10)
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 23

    widths = [18, 32, 28, 12, 12, 17, 12, 12, 12, 14]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:J{last_data_row}"
    sheet.print_area = f"A1:J{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.20
    sheet.page_margins.right = 0.20
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25
    sheet.oddFooter.center.text = "Generated from read-only PostgreSQL input"

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    residual_row = statistics.loc[
        statistics["variable"].eq(
            "Residual Two-Month Cumulative Extreme Humid-Heat Days"
        )
    ].iloc[0]
    metadata_rows = [
        ("Generated At", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Database", args.dbname),
        ("Schema", args.schema),
        ("Database Access", "read-only SELECT"),
        ("Person-Wave Rows", int(statistics.iloc[0]["sample_n"])),
        ("District-Month Cells", int(residual_row["sample_n"])),
        ("Cumulative Exposure Cells", int(residual_row["observed_n"])),
        ("Residual Cumulative Heat SD (days)", float(residual_row["standard_deviation"])),
        ("Person Weighting", "positive analysis weight"),
        ("Household Weighting", "positive household weight on unique household-wave records"),
        ("Weighted SD", "population-form weighted second moment around weighted mean"),
        ("Weighted Quantiles", "linear interpolation on midpoint cumulative-weight positions"),
    ]
    metadata.append(["Field", "Value"])
    for row in metadata_rows:
        metadata.append(list(row))
    for cell in metadata[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True)
    metadata.column_dimensions["A"].width = 27
    metadata.column_dimensions["B"].width = 82
    metadata.print_area = f"A1:B{metadata.max_row}"
    metadata.page_setup.orientation = "landscape"
    metadata.page_setup.paperSize = metadata.PAPERSIZE_LETTER
    metadata.page_setup.fitToWidth = 1
    metadata.page_setup.fitToHeight = 1
    metadata.sheet_properties.pageSetUpPr.fitToPage = True
    metadata.page_margins.left = 0.25
    metadata.page_margins.right = 0.25
    metadata.page_margins.top = 0.30
    metadata.page_margins.bottom = 0.30

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)


def configure_single_page(sheet: object, print_area: str) -> None:
    """Apply the project's one-sheet/one-page review contract."""
    sheet.print_area = print_area
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.20
    sheet.page_margins.right = 0.20
    sheet.page_margins.top = 0.25
    sheet.page_margins.bottom = 0.25


def build_split_workbook(
    statistics: pd.DataFrame,
    output: Path,
    args: argparse.Namespace,
    table_kind: str,
) -> None:
    """Build one single-purpose result workbook from the shared statistics."""
    if table_kind == "descriptive":
        selected = statistics.loc[
            statistics["domain"].isin(
                ["Labor outcomes", "Individual characteristics", "Household resources"]
            )
        ].reset_index(drop=True)
        title_text = "Descriptive Statistics of the Analytical Sample"
        result_sheet_name = "Analytical Sample"
        headers = [
            "Domain",
            "Variable",
            "Sample / Weighting",
            "N",
            "Mean",
            "Standard Deviation",
            "P25",
            "Median",
            "P75",
            "Missing Rate",
        ]
        notes = [
            "Notes: Means, standard deviations, and quantiles use the weighting stated in each row; Missing Rate is the unweighted missing share within the stated sample.",
            "Person rows use positive analysis weights. Household rows use one record per household-wave with positive household weights.",
            "Agricultural Occupation Candidate is restricted to current workers and remains a conservative classification.",
            "Log expenditure variables use log(1 + nonnegative nominal riel). Database access was read-only SELECT; database writes: none.",
        ]
        widths = [18, 32, 28, 12, 12, 17, 12, 12, 12, 14]
    elif table_kind == "identification":
        selected = statistics.loc[
            statistics["domain"].isin(
                ["Climate exposure", "Identification diagnostic"]
            )
        ].reset_index(drop=True)
        title_text = "Identifying Variation in Humid-Heat Exposure"
        result_sheet_name = "Heat Identification"
        headers = [
            "Exposure Diagnostic",
            "N Cells",
            "Mean",
            "Standard Deviation",
            "P05",
            "P25",
            "Median",
            "P75",
            "P95",
            "Missing Rate",
        ]
        notes = [
            "Notes: Statistics use 3,325 unique district-year-month cells and are unweighted; Missing Rate is the unweighted missing share.",
            "Residual Two-Month Cumulative Extreme Humid-Heat Days absorb District Calendar Month and Survey Year fixed effects and isolate the variation available to the main design.",
            "The severe threshold is sparse and is retained as a support diagnostic rather than the preferred treatment measure.",
            "Exposure is aligned at the survey-wave year-month level. Database access was read-only SELECT; database writes: none.",
        ]
        widths = [38, 13, 13, 18, 12, 12, 12, 12, 12, 14]
    else:
        raise ValueError(f"Unknown table kind: {table_kind}")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = result_sheet_name
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:J1")
    sheet["A1"] = title_text
    set_cell_style(sheet["A1"], fill=NAVY, color=WHITE, bold=True, size=15, align="left")
    sheet.row_dimensions[1].height = 31
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=9, wrap=True)
    sheet.row_dimensions[2].height = 40

    audit = workbook.create_sheet("Statistics Audit")
    audit.sheet_view.showGridLines = False
    audit_headers = [
        "domain",
        "variable",
        "sample_weighting",
        "display_type",
        "sample_n",
        "observed_n",
        "mean",
        "standard_deviation",
        "p05",
        "p25",
        "median",
        "p75",
        "p95",
        "missing_rate",
    ]
    audit.append(audit_headers)
    for row_number, row in enumerate(selected.itertuples(index=False), start=2):
        audit.append(
            [
                row.domain,
                row.variable,
                row.sample_weighting,
                row.display_type,
                int(row.sample_n),
                int(row.observed_n),
                float(row.mean) if pd.notna(row.mean) else None,
                float(row.standard_deviation) if pd.notna(row.standard_deviation) else None,
                float(row.p05) if pd.notna(row.p05) else None,
                float(row.p25) if pd.notna(row.p25) else None,
                float(row.median) if pd.notna(row.median) else None,
                float(row.p75) if pd.notna(row.p75) else None,
                float(row.p95) if pd.notna(row.p95) else None,
                float(row.missing_rate),
            ]
        )
    for cell in audit[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=8, wrap=True)
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit_widths = [20, 34, 28, 14, 12, 12, 14, 18, 12, 12, 12, 12, 12, 14]
    for column, width in enumerate(audit_widths, start=1):
        audit.column_dimensions[get_column_letter(column)].width = width
    for row_number in range(2, audit.max_row + 1):
        for column in range(1, audit.max_column + 1):
            cell = audit.cell(row_number, column)
            set_cell_style(
                cell,
                size=8,
                align="left" if column <= 4 else "right",
                wrap=column <= 4,
            )
        audit.cell(row_number, 5).number_format = "#,##0"
        audit.cell(row_number, 6).number_format = "#,##0"
        for column in range(7, 14):
            audit.cell(row_number, column).number_format = "0.000"
        audit.cell(row_number, 14).number_format = "0.0%"
        audit.row_dimensions[row_number].height = 20
    configure_single_page(audit, f"A1:N{audit.max_row}")

    thin = Side(style="thin", color=LIGHT_BORDER)
    group_top = Side(style="medium", color=MEDIUM_BORDER)
    first_data_row = 3
    previous_domain = None
    for offset, row in enumerate(selected.itertuples(index=False)):
        target_row = first_data_row + offset
        if table_kind == "descriptive":
            values = [
                row.domain,
                row.variable,
                row.sample_weighting,
                int(row.observed_n),
                float(row.mean),
                float(row.standard_deviation),
                float(row.p25),
                float(row.median),
                float(row.p75),
                float(row.missing_rate),
            ]
        else:
            values = [
                row.variable,
                int(row.observed_n),
                float(row.mean),
                float(row.standard_deviation),
                float(row.p05),
                float(row.p25),
                float(row.median),
                float(row.p75),
                float(row.p95),
                float(row.missing_rate),
            ]
        fill = PALE_BLUE if offset % 2 else WHITE
        domain_changed = previous_domain is not None and row.domain != previous_domain
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(target_row, column, value)
            left_columns = 3 if table_kind == "descriptive" else 1
            set_cell_style(
                cell,
                fill=fill,
                size=8,
                align="left" if column <= left_columns else "right",
                wrap=column <= left_columns,
            )
            cell.border = Border(top=group_top if domain_changed else Side(), bottom=thin)
        if table_kind == "descriptive":
            sheet.cell(target_row, 4).number_format = "#,##0"
            value_format = (
                "0.0%"
                if row.display_type == "binary"
                else "0.000"
                if row.display_type == "log"
                else "0.00"
            )
            for column in range(5, 10):
                sheet.cell(target_row, column).number_format = value_format
        else:
            sheet.cell(target_row, 2).number_format = "#,##0"
            value_format = "0.0%" if row.variable == "Monthly Humid-Heat Share" else "0.00"
            for column in range(3, 10):
                sheet.cell(target_row, column).number_format = value_format
        sheet.cell(target_row, 10).number_format = "0.0%"
        sheet.row_dimensions[target_row].height = 21
        previous_domain = row.domain

    last_data_row = first_data_row + len(selected) - 1
    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row_number = notes_start + offset
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=10)
        cell = sheet.cell(row_number, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row_number].height = 23
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:J{last_data_row}"
    configure_single_page(sheet, f"A1:J{notes_start + len(notes) - 1}")
    sheet.oddFooter.center.text = "Generated from read-only PostgreSQL input"

    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    residual_row = statistics.loc[
        statistics["variable"].eq(
            "Residual Two-Month Cumulative Extreme Humid-Heat Days"
        )
    ].iloc[0]
    metadata_rows = [
        ("Table", title_text),
        ("Generated At", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Database", args.dbname),
        ("Schema", args.schema),
        ("Database Access", "read-only SELECT; writes: none"),
        ("Person-Wave Rows", int(statistics.iloc[0]["sample_n"])),
        ("District-Month Cells", int(residual_row["sample_n"])),
        ("Cumulative Exposure Cells", int(residual_row["observed_n"])),
        ("Residual Cumulative Heat SD (days)", float(residual_row["standard_deviation"])),
        ("Person Weighting", "positive analysis weight"),
        ("Household Weighting", "positive household weight on unique household-wave records"),
        ("Weighted SD", "population-form weighted second moment around weighted mean"),
        ("Weighted Quantiles", "linear interpolation on midpoint cumulative-weight positions"),
    ]
    metadata.append(["Field", "Value"])
    for row in metadata_rows:
        metadata.append(list(row))
    for cell in metadata[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True)
    for row in metadata.iter_rows(min_row=2):
        set_cell_style(row[0], size=9, align="left", bold=True)
        set_cell_style(row[1], size=9, align="left", wrap=True)
    metadata.column_dimensions["A"].width = 27
    metadata.column_dimensions["B"].width = 82
    configure_single_page(metadata, f"A1:B{metadata.max_row}")

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)


def render_review_png(workbook_path: Path, png_path: Path) -> None:
    soffice = Path("/opt/homebrew/bin/soffice")
    if not soffice.exists():
        located = shutil.which("soffice")
        if not located:
            raise RuntimeError("LibreOffice/soffice is required for workbook rendering")
        soffice = Path(located)
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("pdftoppm is required for PNG rendering")

    png_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mj02b-table-review-") as temporary:
        temporary_path = Path(temporary)
        profile_uri = (temporary_path / "lo-profile").as_uri()
        subprocess.run(
            [
                str(soffice),
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temporary_path),
                str(workbook_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        pdf_path = temporary_path / f"{workbook_path.stem}.pdf"
        if not pdf_path.exists():
            raise RuntimeError("LibreOffice did not create the expected PDF")
        prefix = png_path.with_suffix("")
        subprocess.run(
            [pdftoppm, "-png", "-r", "180", "-singlefile", str(pdf_path), str(prefix)],
            check=True,
            capture_output=True,
            text=True,
        )
    if not png_path.exists() or png_path.stat().st_size == 0:
        raise RuntimeError("PNG review render is missing or empty")


def validate_statistics(statistics: pd.DataFrame) -> None:
    if len(statistics) != 24:
        raise AssertionError(f"Expected 24 rows, observed {len(statistics)}")
    if int(statistics.iloc[0]["sample_n"]) != 179_217:
        raise AssertionError("Unexpected person-wave sample size")
    residual = statistics.loc[
        statistics["variable"].eq(
            "Residual Two-Month Cumulative Extreme Humid-Heat Days"
        )
    ].iloc[0]
    if int(residual["sample_n"]) != 3_325:
        raise AssertionError("Unexpected district-month cell count")
    residual_sd = float(residual["standard_deviation"])
    if not 3.9 < residual_sd < 4.3:
        raise AssertionError(f"Unexpected residual cumulative heat SD: {residual_sd}")
    if statistics[
        ["mean", "standard_deviation", "p05", "p25", "median", "p75", "p95"]
    ].isna().all(axis=1).any():
        raise AssertionError("At least one table row lacks all descriptive statistics")


def validate_workbook(path: Path, table_kind: str) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    if table_kind == "descriptive":
        expected_title = "Descriptive Statistics of the Analytical Sample"
        expected_sheets = ["Analytical Sample", "Statistics Audit", "Metadata"]
        expected_first_header = "Domain"
        expected_formulas = 0
    elif table_kind == "identification":
        expected_title = "Identifying Variation in Humid-Heat Exposure"
        expected_sheets = ["Heat Identification", "Statistics Audit", "Metadata"]
        expected_first_header = "Exposure Diagnostic"
        expected_formulas = 0
    else:
        raise ValueError(f"Unknown table kind: {table_kind}")
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(f"Unexpected sheets: {workbook.sheetnames}")
    sheet = workbook[expected_sheets[0]]
    if sheet["A1"].value != expected_title:
        raise AssertionError("Workbook title is missing")
    if sheet["A2"].value != expected_first_header or sheet.freeze_panes != "A3":
        raise AssertionError("Title must be followed immediately by column headers")
    formulas = [
        cell.value
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=10)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if len(formulas) != expected_formulas:
        raise AssertionError(
            f"Expected {expected_formulas} main-table formulas, observed {len(formulas)}"
        )
    if any("#REF!" in formula for formula in formulas):
        raise AssertionError("Formula reference error found")
    if any(sheet.print_area is None for sheet in workbook.worksheets):
        raise AssertionError("Every worksheet must have a defined print area")
    workbook.close()
    values = load_workbook(path, data_only=True, read_only=True)
    result_sheet = values[expected_sheets[0]]
    last_data_row = 17 if table_kind == "descriptive" else 11
    data_values = [
        result_sheet.cell(row, column).value
        for row in range(3, last_data_row + 1)
        for column in range(1, 11)
    ]
    if any(value is None for value in data_values):
        raise AssertionError("Cached-value read found a blank result cell")
    values.close()


def resolve_path(path: Path) -> Path:
    return path if path.is_absolute() else ROOT / path


def display_path(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def main(default_table: str = "both") -> None:
    args = parse_args(default_table)
    with psycopg.connect(**connection_args(args)) as connection:
        frame = load_analysis_data(connection, args.schema)
        connection.rollback()
    frame = construct_variables(frame)
    statistics = build_statistics(frame)
    validate_statistics(statistics)
    jobs: list[tuple[str, Path, Path]] = []
    if args.table in {"descriptive", "both"}:
        jobs.append(
            (
                "descriptive",
                resolve_path(args.descriptive_output),
                resolve_path(args.descriptive_review_output),
            )
        )
    if args.table in {"identification", "both"}:
        jobs.append(
            (
                "identification",
                resolve_path(args.identification_output),
                resolve_path(args.identification_review_output),
            )
        )
    for table_kind, output, review_output in jobs:
        build_split_workbook(statistics, output, args, table_kind)
        validate_workbook(output, table_kind)
        if not args.skip_render:
            render_review_png(output, review_output)
        print(f"saved_{table_kind}_workbook={display_path(output)}")
        if not args.skip_render:
            print(f"saved_{table_kind}_review_png={display_path(review_output)}")
    residual_sd = statistics.loc[
        statistics["variable"].eq(
            "Residual Two-Month Cumulative Extreme Humid-Heat Days"
        ),
        "standard_deviation",
    ].iloc[0]
    print(f"shared_statistics_rows={len(statistics)}")
    print(f"person_wave_rows={int(statistics.iloc[0]['sample_n'])}")
    print(f"district_month_cells=3325")
    print(f"residual_cumulative_heat_sd_days={float(residual_sd):.6f}")


if __name__ == "__main__":
    main()
