#!/usr/bin/env python3
"""Generate the two wave-level sample and coverage audit tables.

Plan: Keep sample/linkage diagnostics separate from variable-coverage
diagnostics. Each formal table places its column names immediately below its
title and is rendered to a PNG review copy.

Framework: AnaSOP Section 5 sample/linkage audit, Section 6 primary-sample
contract, and Section 7 Step 1. PostgreSQL access is forced to read-only.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
import psycopg
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from psycopg import sql


ROOT = Path(__file__).resolve().parents[2]
RESULTS_DIR = ROOT / "data/results/tables"
REVIEW_DIR = ROOT / "data/exp/figure-table-review"

DEFAULT_SAMPLE_OUTPUT = (
    RESULTS_DIR / "Table_analytical_sample_and_year_month_linkage_by_survey_wave.xlsx"
)
DEFAULT_SAMPLE_REVIEW = (
    REVIEW_DIR / "Table_analytical_sample_and_year_month_linkage_by_survey_wave.png"
)
DEFAULT_COVERAGE_OUTPUT = RESULTS_DIR / "Table_key_variable_coverage_by_survey_wave.xlsx"
DEFAULT_COVERAGE_REVIEW = REVIEW_DIR / "Table_key_variable_coverage_by_survey_wave.png"

NAVY = "17324D"
PALE_TEAL = "DCEBEA"
PALE_BLUE = "EEF4F7"
PALE_GOLD = "FFF2CC"
WHITE = "FFFFFF"
TEXT = "23313D"
MUTED = "5C6B73"
LIGHT_BORDER = "B8C7CE"


def parse_args(default_table: str) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--table", choices=("sample", "coverage", "both"), default=default_table)
    parser.add_argument("--host", default="127.0.0.1")
    parser.add_argument("--port", type=int, default=5432)
    parser.add_argument("--dbname", default="mda")
    parser.add_argument("--schema", default="public")
    parser.add_argument("--user")
    parser.add_argument("--password")
    parser.add_argument("--sample-output", type=Path, default=DEFAULT_SAMPLE_OUTPUT)
    parser.add_argument("--sample-review-output", type=Path, default=DEFAULT_SAMPLE_REVIEW)
    parser.add_argument("--coverage-output", type=Path, default=DEFAULT_COVERAGE_OUTPUT)
    parser.add_argument("--coverage-review-output", type=Path, default=DEFAULT_COVERAGE_REVIEW)
    parser.add_argument("--skip-render", action="store_true")
    return parser.parse_args()


def connection_args(args: argparse.Namespace) -> dict[str, object]:
    values: dict[str, object] = {
        "host": args.host,
        "port": args.port,
        "dbname": args.dbname,
        "options": "-c default_transaction_read_only=on",
    }
    if args.user:
        values["user"] = args.user
    if args.password:
        values["password"] = args.password
    return values


def fetch_frame(connection: psycopg.Connection, query: object) -> pd.DataFrame:
    with connection.cursor() as cursor:
        cursor.execute(query)
        columns = [column.name for column in cursor.description]
        rows = cursor.fetchall()
    return pd.DataFrame(rows, columns=columns)


def load_audit_counts(connection: psycopg.Connection, schema: str) -> pd.DataFrame:
    eligible_query = sql.SQL(
        """
        SELECT
            survey_wave,
            survey_year,
            count(*)::bigint AS eligible_person_rows
        FROM {}.{}
        WHERE age BETWEEN 15 AND 64
          AND survey_month BETWEEN 1 AND 12
          AND coalesce(nullif(person_weight, 0), nullif(household_weight, 0)) > 0
        GROUP BY survey_wave, survey_year
        ORDER BY survey_year
        """
    ).format(sql.Identifier(schema), sql.Identifier("final_EC_CSES"))

    matched_query = sql.SQL(
        """
        SELECT
            a.survey_wave,
            a.survey_year,
            count(*)::bigint AS matched_person_rows,
            count(DISTINCT (a.survey_wave, a.household_id))::bigint
                AS household_wave_rows,
            count(DISTINCT a.admin2_code)::bigint AS districts,
            count(DISTINCT (a.admin2_code, a.survey_month))::bigint
                AS district_month_cells,
            count(*) FILTER (WHERE a.analysis_weight > 0)::bigint AS weight_observed_n,
            count(a.worked_past_week)::bigint AS work_observed_n,
            count(a.weekly_hours_including_zero)::bigint AS hours_observed_n,
            count(a.log_monthly_salary_wages)::bigint AS wage_observed_n,
            count(ed.education_level_harmonized)::bigint AS education_observed_n,
            count(*) FILTER (WHERE a.worked_past_week = 1)::bigint AS current_workers_n,
            count(e.main_occupation_source_code) FILTER
                (WHERE a.worked_past_week = 1)::bigint AS occupation_observed_n,
            count(*) FILTER (
                WHERE a.household_member_count IS NOT NULL
                  AND a.rooms_used IS NOT NULL
                  AND a.floor_area_square_meters IS NOT NULL
                  AND a.monthly_electricity_expense_riel IS NOT NULL
                  AND a.monthly_water_charges_riel IS NOT NULL
                  AND a.has_toilet_facility IS NOT NULL
                  AND a.treats_drinking_water IS NOT NULL
            )::bigint AS household_resource_core_n,
            count(*) FILTER (
                WHERE a.days_wbmax_ge_26c IS NOT NULL
                  AND a.precipitation_month_sum_mm IS NOT NULL
            )::bigint AS climate_observed_n
        FROM {}.{} AS a
        LEFT JOIN {}.{} AS e
          ON e.survey_wave = a.survey_wave
         AND e.person_id = a.person_id
        LEFT JOIN {}.{} AS ed
          ON ed.survey_wave = a.survey_wave
         AND ed.person_id = a.person_id
        GROUP BY a.survey_wave, a.survey_year
        ORDER BY a.survey_year
        """
    ).format(
        sql.Identifier(schema),
        sql.Identifier("final_HEAT_LABOR_ANALYTIC"),
        sql.Identifier(schema),
        sql.Identifier("final_EC_CSES"),
        sql.Identifier(schema),
        sql.Identifier("final_ED_CSES"),
    )

    eligible = fetch_frame(connection, eligible_query)
    matched = fetch_frame(connection, matched_query)
    frame = eligible.merge(
        matched,
        on=["survey_wave", "survey_year"],
        how="outer",
        validate="one_to_one",
    ).sort_values("survey_year", ignore_index=True)
    integer_columns = [column for column in frame.columns if column.endswith("_n")]
    integer_columns += [
        "survey_year",
        "eligible_person_rows",
        "matched_person_rows",
        "household_wave_rows",
        "districts",
        "district_month_cells",
    ]
    for column in integer_columns:
        frame[column] = pd.to_numeric(frame[column], errors="raise").astype("int64")
    return frame


def add_total_row(frame: pd.DataFrame) -> pd.DataFrame:
    total = {
        "survey_wave": "Total",
        "survey_year": None,
        "eligible_person_rows": int(frame["eligible_person_rows"].sum()),
        "matched_person_rows": int(frame["matched_person_rows"].sum()),
        "household_wave_rows": int(frame["household_wave_rows"].sum()),
        "districts": int(frame["districts"].max()),
        "district_month_cells": int(frame["district_month_cells"].sum()),
    }
    for column in [column for column in frame.columns if column.endswith("_n")]:
        total[column] = int(frame[column].sum())
    return pd.concat([frame, pd.DataFrame([total])], ignore_index=True)


def validate_counts(frame: pd.DataFrame) -> None:
    wave_rows = frame.loc[frame["survey_wave"].ne("Total")]
    if len(wave_rows) != 9:
        raise AssertionError(f"Expected nine waves, observed {len(wave_rows)}")
    total = frame.loc[frame["survey_wave"].eq("Total")].iloc[0]
    expected = {
        "eligible_person_rows": 184_886,
        "matched_person_rows": 179_217,
        "household_wave_rows": 59_390,
        "districts": 197,
        "district_month_cells": 3_325,
        "work_observed_n": 179_215,
        "hours_observed_n": 136_364,
        "wage_observed_n": 74_509,
        "education_observed_n": 154_208,
    }
    for column, value in expected.items():
        observed = int(total[column])
        if observed != value:
            raise AssertionError(f"{column}: expected {value}, observed {observed}")
    if (wave_rows["matched_person_rows"] > wave_rows["eligible_person_rows"]).any():
        raise AssertionError("Matched rows exceed eligible rows in at least one wave")
    if (wave_rows["climate_observed_n"] != wave_rows["matched_person_rows"]).any():
        raise AssertionError("Analytic rows without complete primary climate exposure")


def set_cell_style(
    cell: object,
    *,
    fill: str | None = None,
    color: str = TEXT,
    bold: bool = False,
    size: int = 10,
    align: str = "center",
    wrap: bool = False,
) -> None:
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Aptos", size=size, bold=bold, color=color)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)


def add_audit_sheet(workbook: Workbook, frame: pd.DataFrame) -> None:
    audit = workbook.create_sheet("Audit Counts")
    audit.sheet_view.showGridLines = False
    audit.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        audit.append(list(row))
    for cell in audit[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True, size=9, wrap=True)
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    for column in range(1, audit.max_column + 1):
        audit.column_dimensions[get_column_letter(column)].width = 22


def sample_table_rows(frame: pd.DataFrame) -> list[list[object]]:
    rows: list[list[object]] = []
    for _index, row in frame.iterrows():
        eligible = int(row["eligible_person_rows"])
        matched = int(row["matched_person_rows"])
        rows.append(
            [
                row["survey_wave"],
                row["survey_year"] if pd.notna(row["survey_year"]) else "-",
                eligible,
                matched,
                matched / eligible if eligible else 0.0,
                int(row["household_wave_rows"]),
                int(row["districts"]),
                int(row["district_month_cells"]),
                int(row["weight_observed_n"]) / matched if matched else 0.0,
            ]
        )
    return rows


def coverage_table_rows(frame: pd.DataFrame) -> list[list[object]]:
    rows: list[list[object]] = []
    for _index, row in frame.iterrows():
        matched = int(row["matched_person_rows"])
        workers = int(row["current_workers_n"])
        rows.append(
            [
                row["survey_wave"],
                int(row["work_observed_n"]) / matched if matched else 0.0,
                int(row["hours_observed_n"]) / matched if matched else 0.0,
                int(row["wage_observed_n"]) / matched if matched else 0.0,
                int(row["education_observed_n"]) / matched if matched else 0.0,
                int(row["occupation_observed_n"]) / workers if workers else 0.0,
                int(row["household_resource_core_n"]) / matched if matched else 0.0,
                int(row["climate_observed_n"]) / matched if matched else 0.0,
            ]
        )
    return rows


def build_workbook(
    frame: pd.DataFrame,
    output: Path,
    args: argparse.Namespace,
    table_kind: str,
) -> None:
    if table_kind == "sample":
        title = "Analytical Sample and Year-Month Linkage by Survey Wave"
        sheet_name = "Sample and Linkage"
        headers = [
            "Survey Wave",
            "Survey Year",
            "Eligible Person Rows",
            "Matched Person Rows",
            "Retention Rate",
            "Household-Wave Rows",
            "Districts",
            "District-Month Cells",
            "Weight Coverage",
        ]
        rows = sample_table_rows(frame)
        formats = ["@", "0", "#,##0", "#,##0", "0.0%", "#,##0", "#,##0", "#,##0", "0.0%"]
        widths = [14, 12, 18, 18, 14, 19, 11, 19, 16]
        notes = [
            "Notes: Eligible rows are ages 15-64 with a valid released survey month and a positive analysis weight.",
            "Matched rows additionally require a validated current district and district-year-month climate exposure.",
            "District-Month Cells are summed across waves; Districts reports the maximum distinct count. Database access was read-only SELECT.",
        ]
    elif table_kind == "coverage":
        title = "Key Variable Coverage by Survey Wave"
        sheet_name = "Variable Coverage"
        headers = [
            "Survey Wave",
            "Work Participation",
            "Weekly Hours",
            "Monthly Wage",
            "Education",
            "Occupation Among Workers",
            "Household Resource Core",
            "Climate Exposure",
        ]
        rows = coverage_table_rows(frame)
        formats = ["@"] + ["0.0%"] * 7
        widths = [14, 16, 15, 15, 14, 20, 21, 16]
        notes = [
            "Notes: Coverage is the observed share among matched person-wave rows unless otherwise stated.",
            "Occupation coverage uses current workers as the denominator.",
            "Household Resource Core requires household size, rooms, floor area, electricity and water expenditure, toilet, and water treatment.",
            "Climate Exposure requires both primary humid-heat and monthly rainfall measures. Database access was read-only SELECT.",
        ]
    else:
        raise ValueError(f"Unknown table kind: {table_kind}")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.sheet_view.showGridLines = False
    last_column = len(headers)
    last_column_letter = get_column_letter(last_column)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = sheet.cell(1, 1, title)
    set_cell_style(title_cell, fill=NAVY, color=WHITE, bold=True, size=15, align="left")
    sheet.row_dimensions[1].height = 31

    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(2, column, header)
        set_cell_style(cell, fill=PALE_TEAL, bold=True, size=9, wrap=True)
    sheet.row_dimensions[2].height = 40

    thin = Side(style="thin", color=LIGHT_BORDER)
    first_data_row = 3
    last_data_row = first_data_row + len(rows) - 1
    for row_index, values in enumerate(rows, start=first_data_row):
        is_total = values[0] == "Total"
        fill = PALE_GOLD if is_total else (PALE_BLUE if (row_index - first_data_row) % 2 else WHITE)
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            set_cell_style(
                cell,
                fill=fill,
                bold=is_total,
                size=9,
                align="left" if column == 1 else "right",
            )
            cell.border = Border(bottom=thin)
            cell.number_format = formats[column - 1]
        sheet.row_dimensions[row_index].height = 20

    notes_start = last_data_row + 2
    for offset, note in enumerate(notes):
        row = notes_start + offset
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
        cell = sheet.cell(row, 1, note)
        set_cell_style(cell, color=MUTED, size=8, align="left", wrap=True)
        sheet.row_dimensions[row].height = 20 if len(note) < 145 else 27

    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{last_column_letter}{last_data_row}"
    sheet.print_area = f"A1:{last_column_letter}{notes_start + len(notes) - 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.3
    sheet.page_margins.bottom = 0.3
    sheet.oddFooter.center.text = "Generated from read-only PostgreSQL input"

    add_audit_sheet(workbook, frame)
    metadata = workbook.create_sheet("Metadata")
    metadata.sheet_view.showGridLines = False
    total = frame.loc[frame["survey_wave"].eq("Total")].iloc[0]
    metadata_rows = [
        ("Generated At", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Database", args.dbname),
        ("Schema", args.schema),
        ("Database Access", "read-only SELECT"),
        ("Temporal Match", "nominal survey-wave year + released survey month"),
        ("Spatial Match", "validated current district to district-month climate location"),
        ("Person-Wave Rows", int(total["matched_person_rows"])),
        ("Household-Wave Rows", int(total["household_wave_rows"])),
    ]
    metadata.append(["Field", "Value"])
    for row in metadata_rows:
        metadata.append(list(row))
    for cell in metadata[1]:
        set_cell_style(cell, fill=NAVY, color=WHITE, bold=True)
    metadata.column_dimensions["A"].width = 24
    metadata.column_dimensions["B"].width = 78

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)


def render_review_png(workbook_path: Path, png_path: Path) -> None:
    soffice = Path("/opt/homebrew/bin/soffice")
    if not soffice.exists():
        located = shutil.which("soffice")
        if not located:
            raise RuntimeError("LibreOffice/soffice is required for workbook rendering")
        soffice = Path(located)
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("pdftoppm is required for PNG rendering")

    png_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="mj02b-table-review-") as temporary:
        temporary_path = Path(temporary)
        profile_uri = (temporary_path / "lo-profile").as_uri()
        subprocess.run(
            [
                str(soffice),
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(temporary_path),
                str(workbook_path),
            ],
            check=True,
            capture_output=True,
            text=True,
        )
        pdf_path = temporary_path / f"{workbook_path.stem}.pdf"
        if not pdf_path.exists():
            raise RuntimeError("LibreOffice did not create the expected PDF")
        prefix = png_path.with_suffix("")
        subprocess.run(
            [pdftoppm, "-png", "-r", "180", "-singlefile", str(pdf_path), str(prefix)],
            check=True,
            capture_output=True,
            text=True,
        )
    if not png_path.exists() or png_path.stat().st_size == 0:
        raise RuntimeError("PNG review render is missing or empty")


def validate_workbook(path: Path, table_kind: str) -> None:
    expected = {
        "sample": (
            ["Sample and Linkage", "Audit Counts", "Metadata"],
            "Analytical Sample and Year-Month Linkage by Survey Wave",
            "Survey Wave",
            9,
            0,
        ),
        "coverage": (
            ["Variable Coverage", "Audit Counts", "Metadata"],
            "Key Variable Coverage by Survey Wave",
            "Survey Wave",
            8,
            0,
        ),
    }[table_kind]
    workbook = load_workbook(path, data_only=False, read_only=False)
    if workbook.sheetnames != expected[0]:
        raise AssertionError(f"Unexpected sheets: {workbook.sheetnames}")
    sheet = workbook[expected[0][0]]
    if sheet["A1"].value != expected[1] or sheet["A2"].value != expected[2]:
        raise AssertionError("Title-to-header layout is invalid")
    if sheet.max_column != expected[3]:
        raise AssertionError(f"Expected {expected[3]} columns, observed {sheet.max_column}")
    formulas = [
        cell.value
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if len(formulas) != expected[4]:
        raise AssertionError(f"Expected {expected[4]} formulas, observed {len(formulas)}")
    if any("#REF!" in formula for formula in formulas):
        raise AssertionError("Formula reference error found")
    if sheet.freeze_panes != "A3":
        raise AssertionError("Expected title row followed immediately by header row")
    workbook.close()
    values = load_workbook(path, data_only=True, read_only=True)
    result_sheet = values[expected[0][0]]
    data_values = [
        result_sheet.cell(row, column).value
        for row in range(3, 13)
        for column in range(1, expected[3] + 1)
    ]
    if any(value is None for value in data_values):
        raise AssertionError("Cached-value read found a blank result cell")
    values.close()


def resolve_path(path: Path) -> Path:
    return path if path.is_absolute() else ROOT / path


def display_path(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def main(default_table: str = "both") -> None:
    args = parse_args(default_table)
    with psycopg.connect(**connection_args(args)) as connection:
        frame = load_audit_counts(connection, args.schema)
        connection.rollback()
    frame = add_total_row(frame)
    validate_counts(frame)

    jobs = []
    if args.table in ("sample", "both"):
        jobs.append(("sample", resolve_path(args.sample_output), resolve_path(args.sample_review_output)))
    if args.table in ("coverage", "both"):
        jobs.append(("coverage", resolve_path(args.coverage_output), resolve_path(args.coverage_review_output)))

    for table_kind, workbook_path, review_path in jobs:
        build_workbook(frame, workbook_path, args, table_kind)
        validate_workbook(workbook_path, table_kind)
        if not args.skip_render:
            render_review_png(workbook_path, review_path)
        print(f"saved_{table_kind}_workbook={display_path(workbook_path)}")
        if not args.skip_render:
            print(f"saved_{table_kind}_review_png={display_path(review_path)}")

    print(f"validated_waves={len(frame) - 1}")
    print(f"matched_person_rows={int(frame.iloc[-1]['matched_person_rows'])}")
    print(f"household_wave_rows={int(frame.iloc[-1]['household_wave_rows'])}")


if __name__ == "__main__":
    main()
